"""
Validation as a durable job, and every run it produces as a record.

This is the product's own journey: upload a register, queue its validation,
watch it run (or close the page and come back), then read the result by run id
— paged, sorted and filtered on the server, never shipped whole to a browser.

Everything is scoped to the caller's entity. A job, run or upload belonging to
another entity answers 404 — the same answer as one that does not exist, so ids
cannot be probed.
"""
from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.config import settings
from app.database import get_db
from app.deps import get_current_entity, get_current_user, require_entity_write
from app.envelope import ok
from app.models import (
    Entity,
    FindingRecord,
    FindingState,
    PeriodSignOff,
    RegisterUpload,
    User,
    ValidationJob,
    ValidationRun,
    ValidationRunEmployee,
)
from app.services import register_uploads, run_inputs
from app.services import validation_jobs as jobs
from app.services.finding_store import current_run
from app.services.finding_taxonomy import categorise

router = APIRouter()


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _uuid(value: str, what: str) -> uuid.UUID:
    try:
        return uuid.UUID(str(value))
    except ValueError:
        raise HTTPException(status_code=404, detail=f"{what} not found") from None


def _month(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10]).replace(day=1)
    except ValueError:
        raise HTTPException(status_code=400, detail="period_month must be a date, e.g. 2026-06-01") from None


def _job(db: Session, entity: Entity, job_id: str) -> ValidationJob:
    job = db.get(ValidationJob, _uuid(job_id, "Validation job"))
    if job is None or job.entity_id != entity.id:
        raise HTTPException(status_code=404, detail="Validation job not found")
    return job


def _run(db: Session, entity: Entity, run_id: str) -> ValidationRun:
    run = db.get(ValidationRun, _uuid(run_id, "Validation run"))
    if run is None or run.entity_id != entity.id:
        raise HTTPException(status_code=404, detail="Validation run not found")
    return run


def _money(value: Any) -> float | None:
    if value is None:
        return None
    return float(value)


def _describe_job(db: Session, job: ValidationJob) -> dict[str, Any]:
    out = jobs.describe(job, db)
    out["worker_enabled"] = bool(settings.validation_worker_enabled)
    return out


def _describe_run(db: Session, run: ValidationRun, *, with_upload: bool = True) -> dict[str, Any]:
    out: dict[str, Any] = {
        "id": str(run.id),
        "period_month": run.period_month.isoformat(),
        "run_number": run.run_number or 1,
        "status": run.status or "current",
        "source": run.source,
        "run_type": run.run_type,
        "employee_count": run.employee_count,
        "total_findings": run.total_findings,
        "critical_count": run.critical_count,
        "warning_count": run.warning_count,
        "total_financial_impact": _money(run.total_financial_impact),
        "open_financial_impact": _money(run.open_financial_impact),
        "engine_version": run.engine_version,
        "job_id": str(run.job_id) if run.job_id else None,
        "superseded_by_run_id": str(run.superseded_by_run_id) if run.superseded_by_run_id else None,
        "superseded_at": run.superseded_at.isoformat() if run.superseded_at else None,
        "created_at": run.created_at.isoformat() if run.created_at else None,
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "duration_ms": run.duration_ms,
        "inputs_recorded": bool(run.input_digests),
        "has_employee_results": run.input_digests is not None or run.upload_id is not None,
    }
    if with_upload and run.upload_id:
        upload = db.get(RegisterUpload, run.upload_id)
        out["upload"] = register_uploads.describe(upload) if upload else None
    return out


def freshness(db: Session, entity: Entity, run: ValidationRun) -> dict[str, Any]:
    """Is this run still the answer for its inputs? If not, which input moved?

    Compared against the latest upload for the period — a re-upload is exactly
    the change that makes a run stale — and against today's master,
    attendance, CTC and configuration.
    """
    latest = register_uploads.latest_for_period(db, entity.id, run.period_month)
    now = run_inputs.input_digests(
        db, entity, run.period_month,
        rows_sha256=latest.rows_sha256 if latest else (run.input_digests or {}).get("register"),
    )
    changes = run_inputs.changed_inputs(run.input_digests, now)
    return {
        "is_current_for_inputs": not changes,
        "revalidation_required": bool(changes),
        "changes": changes,
        "latest_upload_id": str(latest.id) if latest else None,
    }


# ---------------------------------------------------------------------------
# Uploads
# ---------------------------------------------------------------------------
@router.get("/uploads")
def list_uploads(
    period_month: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """Every upload, newest first — the history a re-upload used to erase."""
    query = db.query(RegisterUpload).filter(RegisterUpload.entity_id == entity.id)
    month = _month(period_month)
    if month:
        query = query.filter(RegisterUpload.period_month == month)
    rows = query.order_by(RegisterUpload.created_at.desc(), RegisterUpload.revision.desc()).limit(limit).all()
    return ok([register_uploads.describe(u) for u in rows])


# ---------------------------------------------------------------------------
# Jobs
# ---------------------------------------------------------------------------
class EnqueueRequest(BaseModel):
    period_month: date
    upload_id: str | None = None
    run_type: str | None = None
    effective_month_from: date | None = None
    effective_month_to: date | None = None
    as_of_date: date | None = None


@router.post("/jobs")
def enqueue_validation(
    body: EnqueueRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(require_entity_write),
):
    """
    Queue validation of a period's register.

    Validates the named upload, or the period's latest one. A second request
    while one is live does not start a second validation: it returns the job
    already running, flagged ``already_queued``, so a double-click or a second
    tab joins the same job instead of racing it.
    """
    month = body.period_month.replace(day=1)
    if body.upload_id:
        upload = db.get(RegisterUpload, _uuid(body.upload_id, "Upload"))
        if upload is None or upload.entity_id != entity.id:
            raise HTTPException(status_code=404, detail="Upload not found")
        if upload.period_month != month:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"That upload is for {upload.period_month:%b %Y}"
                    if upload.period_month else "That upload has no payroll month"
                ) + f", not {month:%b %Y}. Choose the matching month or upload again.",
            )
    else:
        upload = register_uploads.latest_for_period(db, entity.id, month)
        if upload is None:
            raise HTTPException(
                status_code=400,
                detail=f"No register has been uploaded for {month:%b %Y}. Upload it first.",
            )
    if upload.missing_required:
        raise HTTPException(
            status_code=400,
            detail="This register is missing required columns: "
            + ", ".join(upload.missing_required[:10])
            + ". Map or add them and upload again.",
        )

    params = {
        "effective_month_from": body.effective_month_from.isoformat() if body.effective_month_from else None,
        "effective_month_to": body.effective_month_to.isoformat() if body.effective_month_to else None,
        "as_of_date": body.as_of_date.isoformat() if body.as_of_date else None,
    }
    try:
        job = jobs.enqueue(
            db,
            entity_id=entity.id,
            user_id=user.id,
            period_month=month,
            register_id=upload.register_id,
            upload_id=upload.id,
            run_type=body.run_type or upload.run_type or "regular",
            params=params,
        )
        job.employee_total = upload.row_count
        db.commit()
    except jobs.AlreadyQueued as exc:
        return ok({"job": _describe_job(db, exc.job), "already_queued": True})
    return ok({"job": _describe_job(db, job), "already_queued": False})


@router.get("/jobs")
def list_jobs(
    period_month: str | None = Query(default=None),
    active: bool = Query(default=False),
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """Recent jobs; ``active=true`` for the ones still queued or running — what a
    returning user resumes."""
    query = db.query(ValidationJob).filter(ValidationJob.entity_id == entity.id)
    month = _month(period_month)
    if month:
        query = query.filter(ValidationJob.period_month == month)
    if active:
        query = query.filter(ValidationJob.state.in_(("queued", "running")))
    rows = query.order_by(ValidationJob.queued_at.desc()).limit(limit).all()
    return ok([_describe_job(db, j) for j in rows])


@router.get("/jobs/{job_id}")
def get_job(
    job_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    return ok(_describe_job(db, _job(db, entity, job_id)))


@router.post("/jobs/{job_id}/cancel")
def cancel_job(
    job_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(require_entity_write),
):
    job = _job(db, entity, job_id)
    if job.state in ("succeeded", "failed", "cancelled"):
        raise HTTPException(status_code=409, detail=f"This validation has already {job.state}.")
    jobs.request_cancel(db, job, user.id)
    db.commit()
    return ok(_describe_job(db, job))


@router.post("/jobs/{job_id}/retry")
def retry_job(
    job_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(require_entity_write),
):
    """Queue the same validation again, as a new job.

    Safe to repeat: a failed or cancelled job wrote no run, so a retry cannot
    duplicate one — and the old job stays exactly as it failed, for the record.
    """
    job = _job(db, entity, job_id)
    if job.state not in ("failed", "cancelled"):
        raise HTTPException(status_code=409, detail="Only a failed or cancelled validation can be retried.")
    try:
        new = jobs.enqueue(
            db,
            entity_id=entity.id,
            user_id=user.id,
            period_month=job.period_month,
            register_id=job.register_id,
            upload_id=job.upload_id,
            run_type=job.run_type,
            params=dict(job.params or {}),
            retry_of_job_id=job.id,
        )
        db.commit()
    except jobs.AlreadyQueued as exc:
        return ok({"job": _describe_job(db, exc.job), "already_queued": True})
    return ok({"job": _describe_job(db, new), "already_queued": False})


# ---------------------------------------------------------------------------
# Runs
# ---------------------------------------------------------------------------
@router.get("/runs")
def list_runs(
    period_month: str | None = Query(default=None),
    include_superseded: bool = Query(default=True),
    limit: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    query = db.query(ValidationRun).filter(ValidationRun.entity_id == entity.id)
    month = _month(period_month)
    if month:
        query = query.filter(ValidationRun.period_month == month)
    if not include_superseded:
        query = query.filter(ValidationRun.status == "current")
    rows = (
        query.order_by(ValidationRun.period_month.desc(), ValidationRun.run_number.desc())
        .limit(limit)
        .all()
    )
    return ok([_describe_run(db, r) for r in rows])


@router.get("/runs/compare")
def compare_runs(
    base: str = Query(..., description="The earlier run"),
    target: str = Query(..., description="The later run"),
    show: str = Query(default="all", pattern="^(all|new|resolved|changed|unchanged)$"),
    limit: int = Query(default=200, ge=1, le=2000),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """
    What changed between two runs, finding by finding.

    Matched on fingerprint (entity + employee + rule + component), the same
    identity the worklist uses. A finding present in both is ``changed`` if
    its expected, actual, difference, impact or severity moved; otherwise
    ``unchanged``.
    """
    a = _run(db, entity, base)
    b = _run(db, entity, target)

    def _load(run: ValidationRun) -> dict[str, FindingRecord]:
        records = db.query(FindingRecord).filter(FindingRecord.run_id == run.id).all()
        return {r.fingerprint: r for r in records}

    left, right = _load(a), _load(b)

    def _brief(r: FindingRecord) -> dict[str, Any]:
        return {
            "fingerprint": r.fingerprint,
            "employee_id": r.employee_id,
            "employee_name": r.employee_name,
            "rule_id": r.rule_id,
            "rule_name": r.rule_name,
            "component": r.component,
            "severity": r.severity,
            "expected_value": r.expected_value,
            "actual_value": r.actual_value,
            "difference": r.difference,
            "financial_impact": _money(r.financial_impact),
        }

    def _same(x: FindingRecord, y: FindingRecord) -> bool:
        return (
            (x.expected_value, x.actual_value, x.difference, x.severity)
            == (y.expected_value, y.actual_value, y.difference, y.severity)
            and Decimal(x.financial_impact or 0) == Decimal(y.financial_impact or 0)
        )

    groups: dict[str, list[dict[str, Any]]] = {"new": [], "resolved": [], "changed": [], "unchanged": []}
    for fp, rec in right.items():
        if fp not in left:
            groups["new"].append(_brief(rec))
        elif _same(left[fp], rec):
            groups["unchanged"].append(_brief(rec))
        else:
            item = _brief(rec)
            item["before"] = _brief(left[fp])
            groups["changed"].append(item)
    for fp, rec in left.items():
        if fp not in right:
            groups["resolved"].append(_brief(rec))

    def _impact(items: list[dict[str, Any]]) -> float:
        return round(sum(i["financial_impact"] or 0 for i in items), 2)

    counts = {k: len(v) for k, v in groups.items()}
    impacts = {k: _impact(v) for k, v in groups.items()}
    for items in groups.values():
        items.sort(key=lambda i: (-(i["financial_impact"] or 0), i["employee_id"], i["rule_id"]))
    listed = {k: (v[:limit] if show in ("all", k) else []) for k, v in groups.items()}
    return ok({
        "base": _describe_run(db, a, with_upload=False),
        "target": _describe_run(db, b, with_upload=False),
        "counts": counts,
        "financial_impact": impacts,
        "items": listed,
        "truncated": {k: len(v) > limit for k, v in groups.items()},
    })


@router.get("/runs/{run_id}")
def get_run(
    run_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    run = _run(db, entity, run_id)
    out = _describe_run(db, run)
    out["summary"] = run.summary or {}
    out["freshness"] = freshness(db, entity, run)
    out["input_digests"] = run.input_digests
    history = (
        db.query(ValidationRun)
        .filter(ValidationRun.entity_id == entity.id, ValidationRun.period_month == run.period_month)
        .order_by(ValidationRun.run_number.desc())
        .all()
    )
    out["period_runs"] = [
        {"id": str(r.id), "run_number": r.run_number, "status": r.status,
         "created_at": r.created_at.isoformat() if r.created_at else None,
         "total_findings": r.total_findings}
        for r in history
    ]
    return ok(out)


@router.get("/runs/{run_id}/configuration")
def get_run_configuration(
    run_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """The configuration this run was validated under, as it stood then."""
    run = _run(db, entity, run_id)
    return ok({
        "run_id": str(run.id),
        "recorded": run.config_snapshot is not None,
        "configuration": run.config_snapshot,
    })


_EMPLOYEE_SORTS = {
    "employee_id": ValidationRunEmployee.employee_id,
    "employee_name": ValidationRunEmployee.employee_name,
    "risk_score": ValidationRunEmployee.risk_score,
    "failed_checks": ValidationRunEmployee.failed_checks,
    "financial_impact": ValidationRunEmployee.financial_impact,
    "gross": ValidationRunEmployee.gross,
    "department": ValidationRunEmployee.department,
    "position": ValidationRunEmployee.position,
}


#: The computed values the statutory tabs show, pulled from the stored detail.
_COMPUTED_KEYS = (
    "pf_wage", "pf_type", "pf_restricted", "pf_basis_source", "pf_amount_employee",
    "pf_amount_employer", "pf_breakup", "esic_wage", "esic_eligible", "esic_employee",
    "esic_employer", "pt_wage_base", "pt_applicable_state", "pt_due", "lwf_wage_base",
    "lwf_applicable_state", "lwf_employee", "lwf_employer", "paid_days", "lop_days",
    "days_in_month", "gross_total", "lop_check", "increment_arrear", "tds_risk_flags",
    "row_kind_label", "arrear_total", "score_breakdown",
)


def _computed(row: ValidationRunEmployee) -> dict[str, Any]:
    detail = register_uploads.gunzip_json(row.detail_gz) or {}
    return {key: detail.get(key) for key in _COMPUTED_KEYS}


def _employee_out(row: ValidationRunEmployee) -> dict[str, Any]:
    return {
        "employee_id": row.employee_id,
        "employee_name": row.employee_name,
        "department": row.department,
        "work_state": row.work_state,
        "row_kind": row.row_kind,
        "risk_score": row.risk_score,
        "risk_level": row.risk_level,
        "failed_checks": row.failed_checks,
        "critical_count": row.critical_count,
        "warning_count": row.warning_count,
        "passed_checks": row.passed_checks,
        "financial_impact": _money(row.financial_impact),
        "gross": _money(row.gross),
        "net_pay": _money(row.net_pay),
    }


@router.get("/runs/{run_id}/employees")
def list_run_employees(
    run_id: str,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=500),
    sort: str = Query(default="risk_score"),
    order: str = Query(default="desc", pattern="^(asc|desc)$"),
    q: str | None = Query(default=None, max_length=100),
    risk_level: str | None = Query(default=None),
    only_with_findings: bool = Query(default=False),
    department: str | None = Query(default=None),
    include_computed: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    run = _run(db, entity, run_id)
    query = db.query(ValidationRunEmployee).filter(ValidationRunEmployee.run_id == run.id)
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(
            ValidationRunEmployee.employee_id.ilike(like),
            ValidationRunEmployee.employee_name.ilike(like),
        ))
    if risk_level:
        query = query.filter(ValidationRunEmployee.risk_level == risk_level.upper())
    if only_with_findings:
        query = query.filter(ValidationRunEmployee.failed_checks > 0)
    if department:
        query = query.filter(ValidationRunEmployee.department == department)
    total = query.count()
    column = _EMPLOYEE_SORTS.get(sort, ValidationRunEmployee.risk_score)
    primary = column.desc() if order == "desc" else column.asc()
    rows = (
        query.order_by(primary, ValidationRunEmployee.employee_id.asc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    levels = dict(
        db.query(ValidationRunEmployee.risk_level, func.count())
        .filter(ValidationRunEmployee.run_id == run.id)
        .group_by(ValidationRunEmployee.risk_level)
        .all()
    )
    return ok({
        "run_id": str(run.id),
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": (total + page_size - 1) // page_size,
        "items": [
            {**_employee_out(r), **({"computed": _computed(r)} if include_computed else {})}
            for r in rows
        ],
        "risk_levels": levels,
        # Runs recorded before per-employee results were kept have none; say so
        # rather than showing an empty table as if nobody was on the register.
        "employee_results_recorded": total > 0 or run.employee_count == 0 or bool(levels),
    })


@router.get("/runs/{run_id}/employees/{employee_id}")
def get_run_employee(
    run_id: str,
    employee_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """One employee's result in one run, with the row exactly as uploaded."""
    run = _run(db, entity, run_id)
    row = (
        db.query(ValidationRunEmployee)
        .filter(ValidationRunEmployee.run_id == run.id, ValidationRunEmployee.employee_id == employee_id)
        .first()
    )
    if row is None:
        raise HTTPException(status_code=404, detail="That employee is not in this run")
    detail = register_uploads.gunzip_json(row.detail_gz) or {}
    source_row = None
    if run.upload_id:
        upload = db.get(RegisterUpload, run.upload_id)
        if upload is not None:
            for candidate in register_uploads.decode_rows(upload.rows_gz):
                key = str(candidate.get("employee_id") or candidate.get("emp_id")
                          or candidate.get("employee_code") or "").strip()
                if key == employee_id:
                    source_row = candidate
                    break
    return ok({
        "run": _describe_run(db, run),
        "employee": _employee_out(row),
        "result": detail,
        "source_row": source_row,
    })


_FINDING_SORTS = {
    "financial_impact": FindingRecord.financial_impact,
    "employee_id": FindingRecord.employee_id,
    "rule_id": FindingRecord.rule_id,
    "severity": FindingRecord.severity,
}


@router.get("/runs/{run_id}/findings")
def list_run_findings(
    run_id: str,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=500),
    sort: str = Query(default="financial_impact"),
    order: str = Query(default="desc", pattern="^(asc|desc)$"),
    severity: str | None = Query(default=None),
    rule_id: str | None = Query(default=None),
    employee_id: str | None = Query(default=None),
    rule_prefix: str | None = Query(default=None, max_length=100,
                                    description="Comma-separated, e.g. LOP,ATT,ARR"),
    q: str | None = Query(default=None, max_length=100),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    run = _run(db, entity, run_id)
    query = db.query(FindingRecord).filter(FindingRecord.run_id == run.id)
    if rule_prefix:
        prefixes = [p.strip() for p in rule_prefix.split(",") if p.strip()][:10]
        if prefixes:
            query = query.filter(or_(*(FindingRecord.rule_id.like(f"{p}-%") for p in prefixes)))
    if severity:
        query = query.filter(FindingRecord.severity == severity.upper())
    if rule_id:
        query = query.filter(FindingRecord.rule_id == rule_id)
    if employee_id:
        query = query.filter(FindingRecord.employee_id == employee_id)
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(
            FindingRecord.employee_id.ilike(like),
            FindingRecord.employee_name.ilike(like),
            FindingRecord.rule_name.ilike(like),
        ))
    total = query.count()
    column = _FINDING_SORTS.get(sort, FindingRecord.financial_impact)
    primary = column.desc() if order == "desc" else column.asc()
    rows = (
        query.order_by(primary, FindingRecord.employee_id.asc(), FindingRecord.rule_id.asc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    states = {
        s.fingerprint: s
        for s in db.query(FindingState).filter(
            FindingState.entity_id == entity.id,
            FindingState.fingerprint.in_([r.fingerprint for r in rows] or [""]),
        ).all()
    }
    rules = (
        db.query(FindingRecord.rule_id, FindingRecord.rule_name, FindingRecord.severity, func.count())
        .filter(FindingRecord.run_id == run.id)
        .group_by(FindingRecord.rule_id, FindingRecord.rule_name, FindingRecord.severity)
        .all()
    )
    items = []
    for r in rows:
        state = states.get(r.fingerprint)
        items.append({
            "id": str(r.id),
            "fingerprint": r.fingerprint,
            "employee_id": r.employee_id,
            "employee_name": r.employee_name,
            "rule_id": r.rule_id,
            "rule_name": r.rule_name,
            "component": r.component,
            "category": categorise(r.rule_id, r.actual_value),
            "severity": r.severity,
            "status": r.status,
            "expected_value": r.expected_value,
            "actual_value": r.actual_value,
            "difference": r.difference,
            "financial_impact": _money(r.financial_impact),
            "reason": r.reason,
            "suggested_fix": r.suggested_fix,
            "was_waived": r.was_waived,
            "state": state.state if state else None,
            "occurrence_count": state.occurrence_count if state else None,
        })
    return ok({
        "run_id": str(run.id),
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": (total + page_size - 1) // page_size,
        "items": items,
        "rules": sorted(
            ({"rule_id": rid, "rule_name": name, "severity": sev, "count": n}
             for rid, name, sev, n in rules),
            key=lambda x: -x["count"],
        ),
    })


@router.get("/runs/{run_id}/export.xlsx")
def export_run(
    run_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """The run as recorded — not a re-validation. Streams, so 20,000 rows is fine."""
    run = _run(db, entity, run_id)
    try:
        from openpyxl import Workbook
    except ImportError:  # pragma: no cover
        raise HTTPException(status_code=500, detail="Excel export is unavailable on this server.") from None

    book = Workbook(write_only=True)
    summary = book.create_sheet("Run")
    for key, value in (
        ("Company", entity.name),
        ("Payroll month", run.period_month.strftime("%b %Y")),
        ("Run", f"#{run.run_number} ({run.status})"),
        ("Run ID", str(run.id)),
        ("Validated at", run.finished_at.isoformat() if run.finished_at else ""),
        ("Engine version", run.engine_version or "not recorded"),
        ("Employees", run.employee_count),
        ("Findings", run.total_findings),
        ("Critical", run.critical_count),
        ("Warnings", run.warning_count),
        ("Gross exposure (before waivers)", float(run.total_financial_impact or 0)),
        ("Open exposure", float(run.open_financial_impact or 0)),
    ):
        summary.append([key, value])
    if run.upload_id:
        upload = db.get(RegisterUpload, run.upload_id)
        if upload is not None:
            summary.append(["Source file", upload.filename or ""])
            summary.append(["Source file SHA-256", upload.file_sha256])
            summary.append(["Upload revision", upload.revision])

    findings = book.create_sheet("Findings")
    findings.append([
        "Employee ID", "Employee", "Rule", "Rule name", "Component", "Severity",
        "Expected", "Actual", "Difference", "Financial impact", "Waived", "Reason", "Suggested fix",
    ])
    for r in (
        db.query(FindingRecord)
        .filter(FindingRecord.run_id == run.id)
        .order_by(FindingRecord.employee_id, FindingRecord.rule_id)
        .yield_per(1000)
    ):
        findings.append([
            r.employee_id, r.employee_name, r.rule_id, r.rule_name, r.component, r.severity,
            r.expected_value, r.actual_value, r.difference, float(r.financial_impact or 0),
            "yes" if r.was_waived else "no", r.reason, r.suggested_fix,
        ])

    employees = book.create_sheet("Employees")
    employees.append([
        "Employee ID", "Employee", "Department", "Risk score", "Risk level",
        "Failed checks", "Critical", "Warnings", "Passed checks", "Financial impact", "Gross",
    ])
    for e in (
        db.query(ValidationRunEmployee)
        .filter(ValidationRunEmployee.run_id == run.id)
        .order_by(ValidationRunEmployee.position)
        .yield_per(1000)
    ):
        employees.append([
            e.employee_id, e.employee_name, e.department, e.risk_score, e.risk_level,
            e.failed_checks, e.critical_count, e.warning_count, e.passed_checks,
            float(e.financial_impact or 0), float(e.gross) if e.gross is not None else None,
        ])

    buf = io.BytesIO()
    book.save(buf)
    buf.seek(0)
    name = f"validation-{run.period_month:%Y-%m}-run{run.run_number}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# ---------------------------------------------------------------------------
# The month, as one status
# ---------------------------------------------------------------------------
PERIOD_STAGES = {
    "not_uploaded": "No register uploaded",
    "uploaded": "Uploaded — not yet validated",
    "incomplete": "Incomplete — the register cannot be validated as uploaded",
    "validation_pending": "Validation in progress",
    "validation_failed": "Validation failed",
    "revalidation_required": "Revalidation required — inputs changed since the last run",
    "issues_found": "Issues found",
    "ready_for_approval": "Ready for approval",
    "pending_approval": "Submitted for approval",
    "signed_off": "Signed off",
}


@router.get("/periods/{period}/status")
def period_status(
    period: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """
    Where a month stands, in one answer.

    Never "clear" by default: a month with no run, a run older than its
    inputs, or a failed job each say so. "Ready for approval" needs a current
    run, validated against today's inputs, with nothing open.
    """
    month = _month(period)
    upload = register_uploads.latest_for_period(db, entity.id, month)
    active = jobs.active_for(db, entity.id, month)
    last_job = (
        db.query(ValidationJob)
        .filter(ValidationJob.entity_id == entity.id, ValidationJob.period_month == month)
        .order_by(ValidationJob.queued_at.desc())
        .first()
    )
    run = current_run(db, entity.id, month)
    signoff = (
        db.query(PeriodSignOff)
        .filter(PeriodSignOff.entity_id == entity.id, PeriodSignOff.period_month == month)
        .first()
    )
    fresh = freshness(db, entity, run) if run else None
    open_count = 0
    if run:
        open_count = (
            db.query(func.count(FindingRecord.id))
            .filter(FindingRecord.run_id == run.id, FindingRecord.was_waived.is_(False))
            .scalar()
            or 0
        )

    if signoff is not None and signoff.state == "signed":
        stage = "signed_off"
    elif upload is None:
        stage = "not_uploaded"
    elif upload.missing_required:
        stage = "incomplete"
    elif active is not None:
        stage = "validation_pending"
    elif run is None:
        stage = "validation_failed" if (last_job and last_job.state == "failed") else "uploaded"
    elif fresh and fresh["revalidation_required"]:
        stage = "revalidation_required"
    elif open_count:
        stage = "issues_found"
    elif signoff is not None and signoff.state == "pending_approval":
        stage = "pending_approval"
    else:
        stage = "ready_for_approval"

    return ok({
        "period_month": month.isoformat(),
        "stage": stage,
        "stage_label": PERIOD_STAGES[stage],
        "upload": register_uploads.describe(upload) if upload else None,
        "active_job": _describe_job(db, active) if active else None,
        "last_job": _describe_job(db, last_job) if last_job else None,
        "current_run": _describe_run(db, run, with_upload=False) if run else None,
        "open_findings": open_count,
        "freshness": fresh,
        "signoff_state": signoff.state if signoff else None,
    })
