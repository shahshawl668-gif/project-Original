"""
Period sign-off and the evidence pack.

The pack is the artefact that makes the rest of the product worth having on the
record: a single workbook saying what was checked, what was found, what was
accepted and by whom, and which rules and rates were in force at the time.
"""
from __future__ import annotations

import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.database import get_db
from app.deps import get_current_entity, get_current_user, require_entity_write, require_org_admin
from app.envelope import ok
from app.models import Entity, PeriodSignOff, SignOffEvent, User
from app.services import signoff as signoff_service

router = APIRouter()


class SubmitRequest(BaseModel):
    period_month: date
    notes: str | None = Field(default=None, max_length=4000)


class SignRequest(BaseModel):
    period_month: date
    notes: str | None = Field(default=None, max_length=4000)


class ReopenRequest(BaseModel):
    period_month: date
    reason: str = Field(min_length=1, max_length=4000)


def _load(db: Session, entity: Entity, period_month: date) -> PeriodSignOff:
    row = (
        db.query(PeriodSignOff)
        .filter(
            PeriodSignOff.entity_id == entity.id,
            PeriodSignOff.period_month == period_month.replace(day=1),
        )
        .first()
    )
    if row is None:
        raise HTTPException(status_code=404, detail="No sign-off for this period")
    return row


def _out(row: PeriodSignOff) -> dict:
    return {
        "id": str(row.id),
        "period_month": row.period_month.isoformat(),
        "state": row.state,
        "signed_by_email": row.signed_by_email,
        "signed_at": row.signed_at.isoformat() if row.signed_at else None,
        "prepared_at": row.prepared_at.isoformat() if row.prepared_at else None,
        "employee_count": row.employee_count,
        "open_findings": row.open_findings,
        "accepted_exposure": float(row.accepted_exposure or 0),
        "snapshot_digest": row.snapshot_digest,
        "notes": row.notes,
    }


@router.get("")
def list_signoffs(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    rows = (
        db.query(PeriodSignOff)
        .filter(PeriodSignOff.entity_id == entity.id)
        .order_by(PeriodSignOff.period_month.desc())
        .all()
    )
    return ok([_out(r) for r in rows])


@router.get("/{period}")
def get_signoff(
    period: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    row = _load(db, entity, _period(period))
    events = (
        db.query(SignOffEvent)
        .filter(SignOffEvent.signoff_id == row.id)
        .order_by(SignOffEvent.created_at.desc())
        .all()
    )
    return ok(
        {
            "signoff": _out(row),
            "snapshot": row.snapshot,
            "history": [
                {
                    "from_state": e.from_state,
                    "to_state": e.to_state,
                    "reason": e.reason,
                    "actor_email": e.actor_email,
                    "snapshot_digest": e.snapshot_digest,
                    "created_at": e.created_at.isoformat() if e.created_at else None,
                }
                for e in events
            ],
        }
    )


@router.get("/{period}/preview")
def preview(
    period: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """What signing this period would commit to, without committing to it."""
    return ok(signoff_service.build_snapshot(db, entity, _period(period)))


@router.post("/submit")
def submit(
    body: SubmitRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(require_entity_write),
):
    row = signoff_service.submit(db, entity, body.period_month, user, body.notes)
    db.commit()
    db.refresh(row)
    return ok(_out(row))


@router.post("/sign")
def sign(
    body: SignRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_org_admin),
    entity: Entity = Depends(get_current_entity),
):
    """
    Approve a period.

    Restricted to owners and managers, and separate from preparation: the
    person who ran the payroll should not be the only person who ever looked
    at it.
    """
    row = _load(db, entity, body.period_month)
    if row.state == "signed":
        raise HTTPException(status_code=400, detail="This period is already signed")
    signoff_service.sign(db, row, user, body.notes)
    db.commit()
    db.refresh(row)
    return ok(_out(row))


@router.post("/reopen")
def reopen(
    body: ReopenRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_org_admin),
    entity: Entity = Depends(get_current_entity),
):
    row = _load(db, entity, body.period_month)
    if row.state != "signed":
        raise HTTPException(status_code=400, detail="Only a signed period can be reopened")
    signoff_service.reopen(db, row, user, body.reason)
    db.commit()
    db.refresh(row)
    return ok(_out(row))


@router.get("/{period}/evidence-pack")
def evidence_pack(
    period: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    entity: Entity = Depends(get_current_entity),
):
    """
    The period as a workbook: what was checked, found, accepted and by whom.

    Built from the signed snapshot where one exists, and from live data
    otherwise — with the sheet saying plainly which of the two it is, so an
    unsigned draft can never be mistaken for an approved record.
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed.")

    period_month = _period(period)
    row = (
        db.query(PeriodSignOff)
        .filter(
            PeriodSignOff.entity_id == entity.id,
            PeriodSignOff.period_month == period_month,
        )
        .first()
    )
    if row is not None and row.state == "signed" and row.snapshot:
        snapshot = row.snapshot
        provenance = "Signed snapshot"
    else:
        snapshot = signoff_service.build_snapshot(db, entity, period_month)
        provenance = "Live data — this period is NOT signed"

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    title_font = Font(bold=True, size=12)

    def _write_table(sheet, headers: list[str], rows: list[list]) -> None:
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            cell.font = header_font
            cell.fill = header_fill
        for record in rows:
            sheet.append(record)
        for column_cells in sheet.columns:
            width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width + 3, 60)

    # --- Cover -------------------------------------------------------------
    cover = wb.active
    cover.title = "Cover"
    entity_info = snapshot.get("entity", {})
    run = snapshot.get("validation_run") or {}
    cover.append(["Payroll Compliance Evidence Pack"])
    cover["A1"].font = Font(bold=True, size=14)
    cover.append([])
    for label, value in (
        ("Entity", entity_info.get("name")),
        ("Legal name", entity_info.get("legal_name")),
        ("Entity code", entity_info.get("code")),
        ("PF establishment code", entity_info.get("pf_establishment_code")),
        ("ESIC employer code", entity_info.get("esic_employer_code")),
        ("TAN", entity_info.get("tan")),
        ("PAN", entity_info.get("pan")),
        ("Period", snapshot.get("period_month")),
        ("Source", provenance),
        ("Status", row.state if row else "not started"),
        ("Signed by", row.signed_by_email if row else None),
        ("Signed at", row.signed_at.isoformat() if row and row.signed_at else None),
        ("Snapshot digest", row.snapshot_digest if row else None),
        ("Employees", run.get("employee_count")),
        ("Findings raised", run.get("total_findings")),
        ("Critical", run.get("critical_count")),
        ("Gross financial impact", run.get("gross_financial_impact")),
        ("Pack generated at", snapshot.get("built_at")),
    ):
        cover.append([label, value])
        cover.cell(row=cover.max_row, column=1).font = title_font
    cover.column_dimensions["A"].width = 26
    cover.column_dimensions["B"].width = 52
    cover["B10"].alignment = Alignment(wrap_text=True)

    # --- Findings ----------------------------------------------------------
    finding_headers = [
        "Employee ID", "Employee", "Rule", "Rule name", "Severity", "State",
        "Months recurring", "First seen", "Financial impact", "Reason accepted", "Waived until",
    ]

    def _finding_rows(items: list[dict]) -> list[list]:
        return [
            [
                f.get("employee_id"), f.get("employee_name"), f.get("rule_id"),
                f.get("rule_name"), f.get("severity"), f.get("state"),
                f.get("occurrence_count"), f.get("first_seen_period"),
                f.get("financial_impact"), f.get("waiver_reason"), f.get("waived_until"),
            ]
            for f in items
        ]

    _write_table(
        wb.create_sheet("Outstanding"),
        finding_headers,
        _finding_rows(snapshot.get("outstanding_findings", [])),
    )
    _write_table(
        wb.create_sheet("Accepted"),
        finding_headers,
        _finding_rows(snapshot.get("accepted_findings", [])),
    )

    # --- Exposure ----------------------------------------------------------
    exposure = snapshot.get("exposure", {})
    exposure_sheet = wb.create_sheet("Exposure")
    _write_table(
        exposure_sheet,
        ["Head", "Principal", "Interest", "Damages", "Total", "Findings", "Oldest period"],
        [
            [
                h.get("head"), h.get("principal"), h.get("interest"),
                h.get("damages"), h.get("total"), h.get("finding_count"), h.get("oldest_period"),
            ]
            for h in exposure.get("by_head", [])
        ],
    )
    exposure_sheet.append([])
    exposure_sheet.append(["Total exposure", exposure.get("total_exposure")])
    exposure_sheet.append(["Of which accepted (waived)", exposure.get("waived_principal_included")])
    exposure_sheet.append(["Unclassified principal", exposure.get("unclassified_principal")])

    # --- Cost bridge -------------------------------------------------------
    bridge = snapshot.get("cost_bridge")
    if bridge:
        bridge_sheet = wb.create_sheet("Cost bridge")
        _write_table(
            bridge_sheet,
            ["Effect", "Amount", "Employees"],
            [[e["label"], e["amount"], e["employee_count"]] for e in bridge.get("effects", [])],
        )
        bridge_sheet.append([])
        bridge_sheet.append([f"Opening ({bridge.get('compare_to')})", bridge.get("opening_cost")])
        bridge_sheet.append([f"Closing ({bridge.get('period')})", bridge.get("closing_cost")])
        bridge_sheet.append(["Net change", bridge.get("net_change")])

    # --- Configuration in force -------------------------------------------
    configuration = snapshot.get("configuration", {})
    _write_table(
        wb.create_sheet("Minimum wage rates"),
        ["State", "Zone", "Scheduled employment", "Skill", "Monthly floor", "Effective from", "Source"],
        [
            [
                r.get("state"), r.get("zone"), r.get("scheduled_employment"),
                r.get("skill_category"), r.get("total_per_month"),
                r.get("effective_from"), r.get("source_reference"),
            ]
            for r in configuration.get("minimum_wage_rates", [])
        ],
    )

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"evidence-pack-{entity_info.get('code', 'entity')}-{snapshot.get('period_month')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _period(value: str) -> date:
    try:
        return date.fromisoformat(value).replace(day=1)
    except ValueError:
        raise HTTPException(status_code=400, detail="Period must be an ISO date (YYYY-MM-DD)")
