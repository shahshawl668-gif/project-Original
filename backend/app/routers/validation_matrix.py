"""
Company validation rules: drafting, testing, approval and the life of a version.

A rule is never edited in place. Every change is a new version — drafted,
submitted, published by someone entitled to (an owner for statutory rules;
someone other than the author where the organisation requires it) — and a
published version is only ever retired or superseded, so a re-run of an old
month uses the rule that applied then.
"""
from __future__ import annotations

import contextlib
import csv
import io
import json
import uuid
from datetime import UTC, date, datetime
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app.deps import get_current_entity, get_current_user, require_entity_admin
from app.envelope import ok
from app.models import ComponentConfig, Entity, SalaryRegister, SalaryRegisterRow, User, ValidationRuleVersion
from app.schemas.validation_matrix import (
    CloneRequest,
    CopyRequest,
    ImpactRequest,
    RetireRequest,
    ReturnRequest,
    RuleCreate,
    SimulationRequest,
)
from app.services import audit, rule_packs, tenancy, validation_matrix
from app.services.payroll_parse import normalize_col
from app.services.validation_catalog import BUILTIN_RULES, rule_family
from app.services.upload_safety import check_workbook

router = APIRouter()

BUILTIN_CATALOG = [
    {"family": "Payroll integrity", "examples": ["Gross and net reconciliation", "Duplicate employee", "Missing mandatory data", "Arrears and LOP"]},
    {"family": "Statutory", "examples": ["PF", "ESIC", "PT", "LWF", "TDS", "Minimum wage"]},
    {"family": "Cross-register", "examples": ["Attendance versus payroll", "Employee master", "Prior period variance"]},
]
TEMPLATES = [
    {"key": "amount_comparison", "label": "Compare two amounts", "description": "Compare a mapped component, deduction or field against a value or another field."},
    {"key": "required_value", "label": "Required value", "description": "Flag missing mapped inputs rather than silently passing."},
    {"key": "conditional_check", "label": "Conditional check", "description": "Run a comparison only when an optional condition matches."},
    {"key": "required_input", "label": "Mandatory payroll input", "description": "Flag a missing mapped field such as bank account, IFSC, work state or PAN."},
    {"key": "allowed_values", "label": "Allowed values", "description": "Check a field against a company-approved list separated by |."},
]


def _cmp(left: dict, operator: str, right: dict, tolerance: str = "0") -> dict:
    return {"left": left, "operator": operator, "right": right, "tolerance": tolerance}


def _f(key: str) -> dict:
    return {"source": "field", "key": key}


def _lit(value: str) -> dict:
    return {"source": "literal", "value": value}


#: Ready-made rules. Each is a starting point the company reviews, simulates
#: and publishes as its own policy — none is presented as law.
RULE_TEMPLATES: list[dict[str, Any]] = [
    {"key": "lop_limit", "label": "Loss of pay above a limit", "requires": [],
     "rule": {"name": "Loss of pay within 5 days", "severity": "WARNING",
              "assertion": _cmp(_f("lop_days"), "lte", _lit("5")),
              "suggested_fix": "Confirm the leave record and manager approval for extended loss of pay."}},
    {"key": "net_not_negative", "label": "Net pay not negative", "requires": [],
     "rule": {"name": "Net pay is not negative", "severity": "CRITICAL",
              "assertion": _cmp(_f("net_pay"), "gte", _lit("0")),
              "suggested_fix": "Recover the excess deduction over later months instead."}},
    {"key": "paid_days_in_month", "label": "Paid days within the month", "requires": [],
     "rule": {"name": "Paid days do not exceed days in the month", "severity": "CRITICAL",
              "assertion": _cmp(_f("paid_days"), "lte", _f("total_days"))}},
    {"key": "pan_when_tds", "label": "PAN present when TDS is deducted", "requires": [],
     "rule": {"name": "PAN on record when TDS is deducted", "severity": "CRITICAL", "editor_mode": "advanced",
              "condition_group": {"mode": "all", "negate": False, "items": [_cmp(_f("tds"), "gt", _lit("0"))]},
              "assertion": _cmp(_f("pan"), "present", _lit("1")),
              "suggested_fix": "Collect the PAN; without it TDS is due at the Sec 206AA rate."}},
    {"key": "basic_share", "label": "Basic at least a share of gross", "requires": ["basic"],
     "rule": {"name": "Basic is at least 40% of gross", "severity": "WARNING", "editor_mode": "advanced",
              "assertion": _cmp({"source": "expr", "value": "basic / gross * 100"}, "gte", _lit("40")),
              "suggested_fix": "Review the structure — a low basic shrinks PF and gratuity."}},
    {"key": "gross_swing", "label": "Gross change against last month", "requires": [],
     "rule": {"name": "Gross within 20% of last month", "severity": "WARNING", "editor_mode": "advanced",
              "on_missing": "skip",
              "assertion": _cmp({"source": "expr", "value": "abs(gross - prev_gross) / prev_gross * 100"}, "lte", _lit("20")),
              "suggested_fix": "Confirm the revision, arrear or one-off payment behind the change."}},
    {"key": "department_matches_master", "label": "Register department matches the master", "requires": [],
     "rule": {"name": "Department agrees with the employee master", "severity": "INFO", "editor_mode": "advanced",
              "assertion": _cmp(_f("department"), "eq", {"source": "master", "key": "department"}),
              "suggested_fix": "Correct the department in the payroll system or the master."}},
    {"key": "employment_types", "label": "Approved employment types", "requires": [],
     "rule": {"name": "Employment type is an approved type", "severity": "WARNING",
              "assertion": _cmp(_f("employment_type"), "in", _lit("Permanent|Contract|Trainee"))}},
]

#: Fields a version's content is made of. Two versions that agree on all of
#: them are the same rule.
CONTENT_FIELDS = (
    "name", "category", "effective_from", "effective_to", "state", "condition", "assertion",
    "applies_to", "on_missing", "editor_mode", "severity", "blocks_signoff", "responsible_team",
    "suggested_fix", "source_reference",
)
IMPORT_COLUMNS = [
    "rule_key", "name", "category", "effective_from", "effective_to", "state", "severity",
    "blocks_signoff", "on_missing", "applies_to", "condition", "assertion", "responsible_team",
    "suggested_fix", "source_reference", "change_reason",
]
IMPORT_MAX_BYTES = 2_000_000
IMPORT_MAX_ROWS = 500


def _out(rule: ValidationRuleVersion) -> dict:
    return {
        "id": str(rule.id), "rule_key": rule.rule_key, "version": rule.version,
        "name": rule.name, "category": rule.category, "status": rule.status,
        "effective_from": rule.effective_from.isoformat(),
        "effective_to": rule.effective_to.isoformat() if rule.effective_to else None,
        "state": rule.state, "condition": rule.condition, "assertion": rule.assertion,
        "applies_to": rule.applies_to, "on_missing": rule.on_missing or "cannot_validate",
        "editor_mode": rule.editor_mode or "basic",
        "severity": rule.severity, "blocks_signoff": rule.blocks_signoff,
        "responsible_team": rule.responsible_team, "suggested_fix": rule.suggested_fix,
        "source_reference": rule.source_reference, "change_reason": rule.change_reason,
        "created_by": str(rule.created_by), "approved_by": str(rule.approved_by) if rule.approved_by else None,
        "approved_at": rule.approved_at.isoformat() if rule.approved_at else None,
        "cloned_from_id": str(rule.cloned_from_id) if rule.cloned_from_id else None,
        "retired_at": rule.retired_at.isoformat() if rule.retired_at else None,
        "retire_reason": rule.retire_reason,
    }


def _load(db: Session, entity: Entity, rule_id: uuid.UUID) -> ValidationRuleVersion:
    rule = db.query(ValidationRuleVersion).filter(
        ValidationRuleVersion.id == rule_id, ValidationRuleVersion.entity_id == entity.id
    ).first()
    if rule is None:
        raise HTTPException(status_code=404, detail="Rule version not found")
    return rule


def _configured(db: Session, entity_id: Any) -> set[str]:
    return {
        normalize_col(name) for (name,) in db.query(ComponentConfig.component_name).filter(
            ComponentConfig.entity_id == entity_id
        ).all()
    }


def _check_content(content: dict[str, Any], configured: set[str]) -> None:
    validation_matrix.validate_comparison(content["assertion"], configured)
    if content.get("condition"):
        validation_matrix.validate_condition(content["condition"], configured)
    validation_matrix.validate_applies_to(content.get("applies_to"))


def _next_version(db: Session, entity_id: Any, rule_key: str) -> int:
    return (db.query(func.max(ValidationRuleVersion.version)).filter(
        ValidationRuleVersion.entity_id == entity_id,
        ValidationRuleVersion.rule_key == rule_key,
    ).scalar() or 0) + 1


def _content_of(body: RuleCreate) -> dict[str, Any]:
    return {
        "name": body.name, "category": body.category, "effective_from": body.effective_from,
        "effective_to": body.effective_to, "state": body.state, "condition": body.stored_condition(),
        "assertion": body.assertion.model_dump(), "applies_to": body.stored_applies_to(),
        "on_missing": body.on_missing, "editor_mode": body.editor_mode, "severity": body.severity,
        "blocks_signoff": body.blocks_signoff, "responsible_team": body.responsible_team,
        "suggested_fix": body.suggested_fix, "source_reference": body.source_reference,
    }


def _key_rules(rule_key: str, category: str) -> None:
    if not rule_key.startswith(("CUST-", "STATX-")):
        raise ValueError("Rule key must start with CUST- or STATX-")
    if (category == "statutory") != rule_key.startswith("STATX-"):
        raise ValueError("Statutory rules must use STATX-; custom rules use CUST-")


def _draft(db: Session, entity: Entity, user: User, rule_key: str, content: dict[str, Any],
           change_reason: str, cloned_from: uuid.UUID | None = None) -> ValidationRuleVersion:
    rule = ValidationRuleVersion(
        org_id=entity.org_id, entity_id=entity.id, rule_key=rule_key,
        version=_next_version(db, entity.id, rule_key), created_by=user.id, status="draft",
        change_reason=change_reason, cloned_from_id=cloned_from, **content,
    )
    db.add(rule)
    db.flush()
    return rule


def _in_force(db: Session, entity: Entity) -> list[ValidationRuleVersion]:
    """Published, retired-but-dated and pending versions — everything that could apply."""
    return db.query(ValidationRuleVersion).filter(
        ValidationRuleVersion.entity_id == entity.id,
        ValidationRuleVersion.status.in_(("published", "retired", "pending")),
    ).all()


def _latest_by_key(db: Session, entity_id: Any) -> dict[str, ValidationRuleVersion]:
    latest: dict[str, ValidationRuleVersion] = {}
    for row in db.query(ValidationRuleVersion).filter(ValidationRuleVersion.entity_id == entity_id):
        if row.rule_key not in latest or row.version > latest[row.rule_key].version:
            latest[row.rule_key] = row
    return latest


# ---------------------------------------------------------------------------
# Catalogue, templates and packs
# ---------------------------------------------------------------------------
@router.get("/catalog")
def catalog(db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity)):
    components = db.query(ComponentConfig.component_name).filter(ComponentConfig.entity_id == entity.id).all()
    return ok({
        "built_in": BUILTIN_CATALOG, "templates": TEMPLATES,
        "built_in_rules": [
            {"rule_id": rule_id, "name": name, "family": rule_family(rule_id)}
            for rule_id, name in BUILTIN_RULES
        ],
        "fields": sorted(validation_matrix.FIELDS),
        "deductions": sorted(validation_matrix.DEDUCTIONS),
        "components": [name for (name,) in components],
        "operators": sorted(validation_matrix.OPS),
        "master_fields": sorted(validation_matrix.MASTER_FIELDS),
        "dimensions": sorted(validation_matrix.APPLICABILITY_DIMENSIONS),
        "expression_variables": sorted(validation_matrix.expression_variables(_configured(db, entity.id))),
        "limits": {"max_depth": validation_matrix.MAX_DEPTH, "max_conditions": validation_matrix.MAX_LEAVES},
    })


@router.get("/templates")
def rule_templates(db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity)):
    """Ready-made company rules, and whether this company has what each reads."""
    configured = _configured(db, entity.id)
    return ok([
        {**t, "available": all(normalize_col(c) in configured for c in t["requires"]),
         "missing_components": [c for c in t["requires"] if normalize_col(c) not in configured]}
        for t in RULE_TEMPLATES
    ])


class PackToggle(BaseModel):
    enabled: bool
    reason: str = Field(min_length=8, max_length=1000)


@router.get("/packs")
def list_packs(db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity)):
    """Built-in checks by what they protect, what they need, and how the last run went."""
    return ok(rule_packs.describe(db, entity))


@router.put("/packs/{key}")
def toggle_pack(
    key: str, body: PackToggle,
    db: Session = Depends(get_db), user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    """Switch a whole pack on or off, with a reason. Its checks report Disabled, never Passed."""
    try:
        changed = rule_packs.set_enabled(db, entity, user, key, body.enabled, body.reason.strip())
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="No such rule pack") from exc
    db.commit()
    return ok({"key": key, "enabled": body.enabled, "changed": changed})


# ---------------------------------------------------------------------------
# Rules
# ---------------------------------------------------------------------------
@router.get("")
def list_rules(db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity)):
    rows = db.query(ValidationRuleVersion).filter(
        ValidationRuleVersion.entity_id == entity.id
    ).order_by(ValidationRuleVersion.rule_key, ValidationRuleVersion.version.desc()).all()
    return ok([_out(row) for row in rows])


@router.post("")
def create_rule(
    body: RuleCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    content = _content_of(body)
    try:
        _key_rules(body.rule_key, body.category)
        _check_content(content, _configured(db, entity.id))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    rule = _draft(db, entity, user, body.rule_key, content, body.change_reason)
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.drafted", object_type="validation_rule",
                 object_id=str(rule.id), summary=f"Drafted {rule.rule_key} v{rule.version}")
    db.commit()
    return ok(_out(rule))


@router.get("/conflicts")
def list_conflicts(db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity)):
    """Contradictions, duplicates and broken references among rules that could apply together."""
    return ok(validation_matrix.conflicts(_in_force(db, entity), _configured(db, entity.id)))


@router.get("/compare")
def compare_versions(
    a: uuid.UUID, b: uuid.UUID,
    db: Session = Depends(get_db), entity: Entity = Depends(get_current_entity),
):
    """Field by field, what differs between two versions."""
    first, second = _load(db, entity, a), _load(db, entity, b)
    changes = []
    for field in CONTENT_FIELDS:
        va, vb = getattr(first, field), getattr(second, field)
        if validation_matrix._canonical(va) != validation_matrix._canonical(vb):
            changes.append({"field": field, "a": va.isoformat() if isinstance(va, date) else va,
                            "b": vb.isoformat() if isinstance(vb, date) else vb})
    return ok({"a": _out(first), "b": _out(second), "changes": changes})


@router.get("/export.csv")
def export_rules(db: Session = Depends(get_db), user: User = Depends(get_current_user),
                 entity: Entity = Depends(get_current_entity)):
    """The latest version of every rule, in the import format."""
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(IMPORT_COLUMNS)
    for rule in sorted(_latest_by_key(db, entity.id).values(), key=lambda r: r.rule_key):
        applies = "; ".join(f"{k}={'|'.join(v)}" for k, v in (rule.applies_to or {}).items())
        values = [
            rule.rule_key, rule.name, rule.category, rule.effective_from.isoformat(),
            rule.effective_to.isoformat() if rule.effective_to else "", rule.state or "", rule.severity,
            "true" if rule.blocks_signoff else "false", rule.on_missing or "cannot_validate", applies,
            json.dumps(rule.condition) if rule.condition else "", json.dumps(rule.assertion),
            rule.responsible_team or "", rule.suggested_fix or "", rule.source_reference or "",
            rule.change_reason,
        ]
        # Keep text from being evaluated as a formula when opened in a spreadsheet.
        writer.writerow(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                         for v in values])
    return Response("﻿" + out.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="validation-rules.csv"'})


def _parse_import(raw: bytes, filename: str) -> list[dict[str, str]]:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        check_workbook(raw)
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — any unreadable workbook is the same answer
            raise ValueError("The file is not a readable .xlsx workbook") from exc
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        if not rows:
            raise ValueError("The workbook is empty")
        header = [str(c or "").strip() for c in rows[0]]
        records = [
            {h: ("" if v is None else (v.isoformat()[:10] if isinstance(v, (date, datetime)) else str(v)))
             for h, v in zip(header, r, strict=False) if h}
            for r in rows[1:] if any(v not in (None, "") for v in r)
        ]
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("The file is not UTF-8 text; save it as CSV (UTF-8) or .xlsx") from exc
        reader = csv.DictReader(io.StringIO(text, newline=""))
        header = [h.strip() for h in (reader.fieldnames or [])]
        records = [{(k or "").strip(): (v or "") for k, v in r.items()} for r in reader]
    missing = [c for c in ("rule_key", "name", "effective_from", "assertion", "change_reason") if c not in header]
    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}. Export the rules to get the format.")
    if len(records) > IMPORT_MAX_ROWS:
        raise ValueError(f"At most {IMPORT_MAX_ROWS} rules per file")
    return records


def _unescape(value: str) -> str:
    value = (value or "").strip()
    return value[1:] if value.startswith("'") and value[1:2] in ("=", "+", "-", "@") else value


def _record_to_body(rec: dict[str, str]) -> RuleCreate:
    v = {k: _unescape(val) for k, val in rec.items()}
    applies: dict[str, list[str]] = {}
    for part in (v.get("applies_to") or "").split(";"):
        if "=" in part:
            dim, values = part.split("=", 1)
            applies[dim.strip()] = [x.strip() for x in values.split("|") if x.strip()]
    try:
        condition = json.loads(v["condition"]) if v.get("condition") else None
        assertion = json.loads(v["assertion"])
    except json.JSONDecodeError as exc:
        raise ValueError(f"condition/assertion must be JSON as exported: {exc.msg}") from exc
    payload: dict[str, Any] = {
        "rule_key": v["rule_key"].upper(), "name": v["name"], "category": v.get("category") or "custom",
        "effective_from": v["effective_from"], "effective_to": v.get("effective_to") or None,
        "state": v.get("state") or None, "severity": (v.get("severity") or "WARNING").upper(),
        "blocks_signoff": (v.get("blocks_signoff") or "").strip().lower() in ("true", "yes", "1"),
        "on_missing": v.get("on_missing") or "cannot_validate", "applies_to": applies or None,
        "assertion": assertion, "responsible_team": v.get("responsible_team") or None,
        "suggested_fix": v.get("suggested_fix") or None, "source_reference": v.get("source_reference") or None,
        "change_reason": v.get("change_reason") or "",
    }
    if condition is not None:
        payload["condition_group" if validation_matrix.is_group(condition) else "condition"] = condition
    return RuleCreate(**payload)


@router.post("/import")
async def import_rules(
    file: UploadFile,
    dry_run: bool = Query(default=True),
    db: Session = Depends(get_db),
    user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    """
    Rules from a CSV or Excel file, previewed first.

    Every row is checked exactly as a drafted rule would be. Committing creates
    **drafts** only — nothing imported applies until it is submitted and
    published like any other version — and is refused while any row is wrong.
    """
    raw = await file.read(IMPORT_MAX_BYTES + 1)
    if len(raw) > IMPORT_MAX_BYTES:
        raise HTTPException(status_code=413, detail="Rule files are limited to 2 MB")
    try:
        records = _parse_import(raw, file.filename or "")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    configured = _configured(db, entity.id)
    latest = _latest_by_key(db, entity.id)
    preview: list[dict[str, Any]] = []
    plans: list[tuple[RuleCreate, dict[str, Any]]] = []
    seen_keys: set[str] = set()
    for line, rec in enumerate(records, 2):
        entry: dict[str, Any] = {"row": line, "rule_key": (rec.get("rule_key") or "").strip().upper(),
                                 "name": rec.get("name") or "", "errors": []}
        try:
            body = _record_to_body(rec)
            content = _content_of(body)
            _key_rules(body.rule_key, body.category)
            _check_content(content, configured)
            if body.rule_key in seen_keys:
                raise ValueError("This rule key appears twice in the file")
            seen_keys.add(body.rule_key)
        except ValidationError as exc:
            entry["errors"] = [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()]
            entry["status"] = "error"
        except (ValueError, KeyError) as exc:
            entry["errors"] = [str(exc)]
            entry["status"] = "error"
        else:
            current = latest.get(body.rule_key)
            if current is None:
                entry["status"] = "new"
            else:
                changed = [f for f in CONTENT_FIELDS
                           if validation_matrix._canonical(getattr(current, f)) != validation_matrix._canonical(content[f])]
                entry["status"] = "new_version" if changed else "unchanged"
                entry["changes"] = changed
                entry["current_version"] = current.version
            if entry["status"] != "unchanged":
                plans.append((body, content))
        preview.append(entry)
    errors = sum(1 for p in preview if p["status"] == "error")
    summary = {s: sum(1 for p in preview if p["status"] == s) for s in ("new", "new_version", "unchanged", "error")}
    if dry_run:
        return ok({"preview": True, "rows": preview, "summary": summary})
    if errors:
        raise HTTPException(status_code=422, detail=f"{errors} row(s) have errors; fix them and preview again")
    created = []
    for body, content in plans:
        rule = _draft(db, entity, user, body.rule_key, content, body.change_reason)
        created.append(f"{rule.rule_key} v{rule.version}")
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.imported", object_type="validation_rule", object_id=None,
                 summary=f"Imported {len(created)} rule draft(s) from {file.filename}",
                 detail={"drafts": created})
    db.commit()
    return ok({"preview": False, "created": created, "summary": summary})


@router.post("/copy")
def copy_rules(
    body: CopyRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    """
    Copy rule versions to other companies, as drafts there.

    The caller must be an owner or manager of every target; a company they
    cannot administer is reported "not available" whether or not it exists.
    Each copy is checked against the target's own components, arrives as a
    draft (it is approved there, by that company's rules), and is recorded in
    the audit trail of both companies.
    """
    sources = []
    for rid in body.rule_ids:
        try:
            sources.append(_load(db, entity, uuid.UUID(rid)))
        except ValueError as exc:
            raise HTTPException(status_code=404, detail="Rule version not found") from exc
    results = []
    for tid in dict.fromkeys(body.target_entity_ids):
        target = None
        with contextlib.suppress(ValueError):
            target = db.get(Entity, uuid.UUID(tid))
        if (target is None or target.id == entity.id or not target.is_active
                or not tenancy.role_at_least(db, user, "manager", target)):
            results.append({"entity_id": tid, "entity_name": None, "status": "not_available", "rules": []})
            continue
        configured = _configured(db, target.id)
        outcome = []
        for src in sources:
            content = {f: getattr(src, f) for f in CONTENT_FIELDS}
            try:
                _check_content(content, configured)
            except ValueError as exc:
                outcome.append({"rule_key": src.rule_key, "status": "skipped", "reason": str(exc)})
                continue
            if body.dry_run:
                outcome.append({"rule_key": src.rule_key, "status": "will_copy",
                                "version": _next_version(db, target.id, src.rule_key)})
                continue
            copy = _draft(db, target, user, src.rule_key, content,
                          f"Copied from {entity.name} {src.rule_key} v{src.version}: {body.change_reason}")
            outcome.append({"rule_key": src.rule_key, "status": "copied", "version": copy.version})
        if not body.dry_run:
            copied = [o for o in outcome if o["status"] == "copied"]
            audit.record(db, entity_id=target.id, org_id=target.org_id, user=user,
                         action="validation_rule.copied_in", object_type="validation_rule", object_id=None,
                         summary=f"{len(copied)} rule draft(s) copied in from {entity.name}",
                         detail={"from_entity_id": str(entity.id), "rules": copied, "reason": body.change_reason})
            audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                         action="validation_rule.copied_out", object_type="validation_rule", object_id=None,
                         summary=f"{len(copied)} rule(s) copied to {target.name} as drafts",
                         detail={"to_entity_id": str(target.id), "rules": copied, "reason": body.change_reason})
        results.append({"entity_id": str(target.id), "entity_name": target.name, "status": "ok", "rules": outcome})
    if not body.dry_run:
        db.commit()
    return ok({"dry_run": body.dry_run, "targets": results})


# ---------------------------------------------------------------------------
# One version
# ---------------------------------------------------------------------------
def _month_rows(db: Session, entity: Entity, period: date) -> list[dict[str, Any]]:
    """Every employee of the saved register for ``period``, with last month and the master."""
    import calendar

    from app.services.validation import _prior_register_rows
    from app.services.workforce import master_as_of

    register = db.query(SalaryRegister).filter(
        SalaryRegister.entity_id == entity.id, SalaryRegister.period_month == period
    ).order_by(SalaryRegister.created_at.desc()).first()
    if register is None:
        raise HTTPException(status_code=404, detail="No saved salary register for this month")
    rows = db.query(SalaryRegisterRow).filter(
        SalaryRegisterRow.entity_id == entity.id, SalaryRegisterRow.register_id == register.id
    ).order_by(SalaryRegisterRow.employee_id).all()
    prior = _prior_register_rows(db, entity.id, period)
    end = period.replace(day=calendar.monthrange(period.year, period.month)[1])
    master = master_as_of(db, entity.id, end)
    out = []
    for row in rows:
        data = validation_matrix.stored_row(row)
        data["_previous"] = validation_matrix.stored_row(prior[row.employee_id]) if row.employee_id in prior else None
        data["_master"] = validation_matrix.master_dict(master.get(row.employee_id))
        out.append(data)
    return out


@router.post("/{rule_id}/simulate")
def simulate(
    rule_id: uuid.UUID, body: SimulationRequest,
    db: Session = Depends(get_db), user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    rule = _load(db, entity, rule_id)
    period = body.period_month.replace(day=1)
    rows = _month_rows(db, entity, period)[: body.limit]
    findings = [result for row in rows if (result := validation_matrix.evaluate(rule, row)) is not None]
    return ok({"period_month": period.isoformat(), "sampled": len(rows),
               "findings": findings, "truncated": len(rows) == body.limit})


@router.post("/{rule_id}/impact")
def impact(
    rule_id: uuid.UUID, body: ImpactRequest,
    db: Session = Depends(get_db), user: User = Depends(require_entity_admin),
    entity: Entity = Depends(get_current_entity),
):
    """
    What publishing this version would change, on a whole saved month.

    Every employee is evaluated against this version and against the version
    of the same rule currently in force for that month, if any.
    """
    rule = _load(db, entity, rule_id)
    period = body.period_month.replace(day=1)
    rows = _month_rows(db, entity, period)
    current = next((r for r in validation_matrix.published_for(db, entity.id, period)
                    if r.rule_key == rule.rule_key and r.id != rule.id), None)
    proposed: dict[str, dict] = {}
    existing: dict[str, dict] = {}
    tally = {"passed": 0, "failed": 0, "cannot_validate": 0, "skipped": 0, "out_of_scope": 0}
    for row in rows:
        eid = str(row.get("employee_id") or "")
        verdict, f = validation_matrix.outcome(rule, row)
        tally[verdict] += 1
        if f is not None:
            proposed[eid] = f
        if current is not None and (g := validation_matrix.evaluate(current, row)) is not None:
            existing[eid] = g
    unverifiable = [e for e, f in proposed.items() if f["evidence"]["unverifiable"]]
    new = sorted(set(proposed) - set(existing))
    cleared = sorted(set(existing) - set(proposed))
    return ok({
        "period_month": period.isoformat(),
        "employees_evaluated": len(rows),
        # Every employee lands in exactly one of these; skipped and out-of-scope
        # are said, not folded into "nothing failed".
        "outcomes": tally,
        "compared_with": _out(current) if current else None,
        "would_flag": len(proposed) - len(unverifiable),
        "would_be_unverifiable": len(unverifiable),
        "currently_flagged": len(existing),
        "newly_flagged": len(new),
        "no_longer_flagged": len(cleared),
        "unchanged": len(set(proposed) & set(existing)),
        "sample": [proposed[e] for e in new[:25]] or list(proposed.values())[:25],
        "cleared_sample": cleared[:25],
    })


@router.post("/{rule_id}/submit")
def submit(rule_id: uuid.UUID, db: Session = Depends(get_db),
           user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    rule = _load(db, entity, rule_id)
    if rule.status != "draft":
        raise HTTPException(status_code=409, detail="Only drafts can be submitted")
    rule.status = "pending"
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.submitted", object_type="validation_rule",
                 object_id=str(rule.id), summary=f"Submitted {rule.rule_key} v{rule.version}")
    db.commit()
    return ok(_out(rule))


@router.post("/{rule_id}/return")
def return_to_draft(rule_id: uuid.UUID, body: ReturnRequest, db: Session = Depends(get_db),
                    user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    """Send a submitted version back to its author, with the reason."""
    rule = _load(db, entity, rule_id)
    if rule.status != "pending":
        raise HTTPException(status_code=409, detail="Only a submitted version can be returned")
    rule.status = "draft"
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.returned", object_type="validation_rule", object_id=str(rule.id),
                 summary=f"Returned {rule.rule_key} v{rule.version} to draft: {body.reason}")
    db.commit()
    return ok(_out(rule))


@router.post("/{rule_id}/publish")
def publish(rule_id: uuid.UUID, db: Session = Depends(get_db),
            user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    rule = _load(db, entity, rule_id)
    if rule.status != "pending":
        raise HTTPException(status_code=409, detail="Submit a draft before publishing")
    membership = tenancy.get_membership(db, user)
    if rule.category == "statutory" and (membership is None or membership.role != "owner"):
        raise HTTPException(status_code=403, detail="A group owner must approve statutory rules")
    from app.services import approvals

    try:
        approvals.require_independent(
            db, entity.org_id, "matrix_publish_requires_independent_approver",
            preparer_id=rule.created_by, approver_id=user.id, what="a validation rule",
        )
    except approvals.ApprovalRefused as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    clash = db.query(ValidationRuleVersion).filter(
        ValidationRuleVersion.entity_id == entity.id,
        ValidationRuleVersion.rule_key == rule.rule_key,
        ValidationRuleVersion.effective_from == rule.effective_from,
        ValidationRuleVersion.status == "published",
    ).first()
    if clash:
        raise HTTPException(status_code=409, detail="A published version already starts on this date")
    others = [r for r in _in_force(db, entity) if r.status != "pending" and r.rule_key != rule.rule_key]
    blocking = [c for c in validation_matrix.conflicts([rule, *others], _configured(db, entity.id))
                if c["kind"] in ("contradiction", "broken_reference")
                and any(ref["id"] == str(rule.id) for ref in c["rules"])]
    if blocking:
        raise HTTPException(status_code=409, detail=blocking[0]["message"])
    rule.status = "published"
    rule.approved_by = user.id
    rule.approved_at = datetime.now(UTC)
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.published", object_type="validation_rule",
                 object_id=str(rule.id), summary=f"Published {rule.rule_key} v{rule.version}")
    db.commit()
    return ok(_out(rule))


@router.post("/{rule_id}/clone")
def clone(rule_id: uuid.UUID, body: CloneRequest, db: Session = Depends(get_db),
          user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    """A new draft from any version: the next version of the same rule, or a new rule key."""
    src = _load(db, entity, rule_id)
    key = body.rule_key or src.rule_key
    content = {f: getattr(src, f) for f in CONTENT_FIELDS}
    try:
        _key_rules(key, src.category)
        _check_content(content, _configured(db, entity.id))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    rule = _draft(db, entity, user, key, content, body.change_reason, cloned_from=src.id)
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.cloned", object_type="validation_rule", object_id=str(rule.id),
                 summary=f"Drafted {rule.rule_key} v{rule.version} from {src.rule_key} v{src.version}")
    db.commit()
    return ok(_out(rule))


@router.post("/{rule_id}/rollback")
def rollback(rule_id: uuid.UUID, body: ReturnRequest, db: Session = Depends(get_db),
             user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    """
    Bring back an earlier version — as a new draft, never by rewriting history.

    The draft carries the old version's logic under the next version number,
    and goes through submit and publish like any change.
    """
    src = _load(db, entity, rule_id)
    latest = _next_version(db, entity.id, src.rule_key) - 1
    if src.version == latest:
        raise HTTPException(status_code=409, detail="This is already the latest version")
    content = {f: getattr(src, f) for f in CONTENT_FIELDS}
    try:
        _check_content(content, _configured(db, entity.id))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"Version {src.version} no longer fits this company: {exc}") from exc
    rule = _draft(db, entity, user, src.rule_key, content,
                  f"Rollback to v{src.version}: {body.reason}", cloned_from=src.id)
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.rollback_drafted", object_type="validation_rule", object_id=str(rule.id),
                 summary=f"Drafted {rule.rule_key} v{rule.version} restoring v{src.version}: {body.reason}")
    db.commit()
    return ok(_out(rule))


@router.post("/{rule_id}/retire")
def retire(rule_id: uuid.UUID, body: RetireRequest, db: Session = Depends(get_db),
           user: User = Depends(require_entity_admin), entity: Entity = Depends(get_current_entity)):
    """Stop a published version from a date. Months before it keep the rule."""
    rule = _load(db, entity, rule_id)
    if rule.status != "published":
        raise HTTPException(status_code=409, detail="Only a published version can be retired")
    membership = tenancy.get_membership(db, user)
    if rule.category == "statutory" and (membership is None or membership.role != "owner"):
        raise HTTPException(status_code=403, detail="A group owner must retire statutory rules")
    if body.effective_to < rule.effective_from:
        raise HTTPException(status_code=400, detail="A version cannot end before it starts")
    rule.status = "retired"
    rule.effective_to = body.effective_to
    rule.retired_at = datetime.now(UTC)
    rule.retired_by = user.id
    rule.retire_reason = body.reason
    audit.record(db, entity_id=entity.id, org_id=entity.org_id, user=user,
                 action="validation_rule.retired", object_type="validation_rule", object_id=str(rule.id),
                 summary=f"Retired {rule.rule_key} v{rule.version} after {body.effective_to.isoformat()}: {body.reason}")
    db.commit()
    return ok(_out(rule))
