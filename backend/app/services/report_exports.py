"""One export implementation for direct and queued Report Builder output.

Two formats from the same preview: an Excel workbook (details, the pivot when
the report has one, a native chart when it asks for one, provenance and data
basis) and a PDF of the same content for people who will read rather than
recalculate. Both state what they cover and how they were made; neither shows a
value the preview would not.
"""
from __future__ import annotations

import hashlib
import io
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.models import SalaryRegister
from app.services import report_builder, reporting
from app.services.export_safety import neutralise_workbook

FORMATS = {
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}
#: A PDF of more rows than this is not something anyone reads; Excel is.
PDF_MAX_ROWS = 3000
_FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu")


def _sources(db, entity, spec) -> list[dict]:
    registers = (db.query(SalaryRegister)
                 .filter(SalaryRegister.entity_id == entity.id,
                         SalaryRegister.period_month >= date.fromisoformat(spec["date_from"]),
                         SalaryRegister.period_month <= date.fromisoformat(spec["date_to"]))
                 .order_by(SalaryRegister.period_month, SalaryRegister.id).all())
    return [{"id": str(reg.id), "period": reg.period_month.isoformat(),
             "uploaded_at": reg.created_at.isoformat() if reg.created_at else None}
            for reg in registers]


def _about(report, entity, spec, requester, result, sources, generated_at) -> dict:
    dataset = report_builder.DATASETS[spec["dataset"]]
    return {
        "Module": "PeopleOps Reports", "Report": report.name,
        "Definition ID": str(report.id), "Definition version": report.version,
        "Dataset": dataset["label"], "Entity": entity.name,
        "Entity code": getattr(entity, "code", "") or "",
        "Period from": spec["date_from"], "Period to": spec["date_to"],
        "Breakdown": dataset["breakdowns"][spec["dimension"]],
        "Filters": ", ".join(f"{k}: {', '.join(v)}" for k, v in spec["filters"].items()) or "none",
        "Requested by": requester.email, "Generated at": generated_at.isoformat(timespec="seconds"),
        "Data basis": "Current stored data at generation; output is fixed after generation",
        "How it is counted": dataset["note"],
        "Outcome": result["status"], "Record count": result["record_count"],
        "Source register IDs": ", ".join(ref["id"] for ref in sources) or "none",
    }


def build(db, entity, report, requester, fmt: str = "xlsx") -> tuple[bytes, dict]:
    if fmt not in FORMATS:
        raise ValueError("Output must be Excel or PDF")
    spec = report_builder.validate(report.specification)
    if not spec["date_from"] or not spec["date_to"]:
        raise ValueError("Choose both period bounds before export")
    result = report_builder.preview(db, entity.id, spec, row_limit=None)
    sources = _sources(db, entity, spec)
    generated_at = datetime.now(UTC)
    about = _about(report, entity, spec, requester, result, sources, generated_at)
    data = _pdf(report, spec, result, about) if fmt == "pdf" else _xlsx(db, entity, spec, result, about)
    return data, {
        "record_count": result["record_count"],
        "control_totals": result["control_totals"],
        "source_references": sources,
        "sha256": hashlib.sha256(data).hexdigest(),
        "generated_at": generated_at,
        "format": fmt,
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
_NUMBER_FORMAT = {"inr": '"₹"#,##0.00', "number": "#,##0.##", "pct": '0.00"%"', "month": "mmm yyyy"}


def _xlsx(db, entity, spec, result, about) -> bytes:
    wb = reporting._openpyxl().Workbook()
    reporting._provenance_sheet(wb, about)
    reporting._sheet(wb, "Summary", ["Metric", "Value"], [
        ["Matched records", result["record_count"]],
        *[[key.replace("_", " ").title(), value] for key, value in (result["control_totals"] or {}).items()],
    ])
    columns = result["columns"]
    headers = [c["label"] for c in columns]
    rows = [[date.fromisoformat(item[c["key"]]) if c["key"] == "period" and item[c["key"]]
             else item[c["key"]] for c in columns] for item in result["rows"]]
    details = reporting._sheet(wb, "Details", headers, rows)
    for index, column in enumerate(columns, 1):
        fmt = _NUMBER_FORMAT.get(column["unit"])
        if fmt:
            for cells in details.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = fmt
    if result.get("pivot"):
        _pivot_sheet(wb, result["pivot"])
    if result.get("chart"):
        _chart_sheet(wb, result["chart"])
    reporting._basis_sheet(db, entity.id, {
        "date_from": date.fromisoformat(spec["date_from"]),
        "date_to": date.fromisoformat(spec["date_to"]),
    }, wb)
    payload = io.BytesIO()
    neutralise_workbook(wb)
    wb.save(payload)
    return payload.getvalue()


def _month_label(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%b %Y")


def _pivot_sheet(wb, pivot: dict) -> None:
    headers = ["Breakdown", *[_month_label(p) for p in pivot["periods"]]]
    if pivot["row_totals"]:
        headers.append("Total")
    rows = []
    for r in pivot["rows"]:
        line = [r["dimension"], *[r["cells"].get(p) for p in pivot["periods"]]]
        if pivot["row_totals"]:
            line.append(r["total"])
        rows.append(line)
    if pivot["column_totals"]:
        line = ["Total", *[pivot["column_totals"].get(p) for p in pivot["periods"]]]
        if pivot["row_totals"]:
            line.append(pivot["grand_total"])
        rows.append(line)
    ws = reporting._sheet(wb, f"Pivot - {pivot['label']}"[:31], headers, rows)
    fmt = _NUMBER_FORMAT.get(pivot["unit"])
    for cells in ws.iter_cols(min_col=2, max_col=len(headers), min_row=2):
        for cell in cells:
            cell.number_format = fmt
    if pivot.get("note"):
        ws.append([])
        ws.append([pivot["note"]])


def _chart_sheet(wb, chart: dict) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference

    ws = wb.create_sheet("Chart")
    if chart["type"] == "bar":
        ws.append(["Breakdown", f"{chart['label']} · {_month_label(chart['period'])}"])
        for bar in chart["bars"]:
            ws.append([bar["group"], bar["value"]])
        plot = BarChart()
        plot.type = "bar"
        data = Reference(ws, min_col=2, min_row=1, max_row=len(chart["bars"]) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(chart["bars"]) + 1)
    else:
        ws.append(["Month", *[s["group"] for s in chart["series"]]])
        for i, period in enumerate(chart["periods"]):
            ws.append([_month_label(period), *[s["points"][i] for s in chart["series"]]])
        plot = LineChart()
        data = Reference(ws, min_col=2, max_col=len(chart["series"]) + 1, min_row=1,
                         max_row=len(chart["periods"]) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(chart["periods"]) + 1)
    plot.add_data(data, titles_from_data=True)
    plot.set_categories(cats)
    plot.title = chart["label"]
    plot.height, plot.width = 9, 18
    ws.add_chart(plot, "E2")
    if chart.get("omitted"):
        ws.append([])
        ws.append([_omitted(chart["omitted"], "Details")])


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _omitted(n: int, where: str) -> str:
    return (f"1 smaller group is not drawn; it is in {where}." if n == 1
            else f"{n} smaller groups are not drawn; they are in {where}.")


def _font_dir() -> Path:
    for folder in _FONT_DIRS:
        if (Path(folder) / "DejaVuSans.ttf").exists():
            return Path(folder)
    raise ValueError(
        "PDF output needs the DejaVu font on the server (fonts-dejavu-core), which is not installed. "
        "Generate Excel instead, or ask the platform team to install it."
    )


def _indian(value: Decimal) -> str:
    """12,34,567.89 — lakh and crore grouping, as the product shows money."""
    sign = "-" if value < 0 else ""
    whole, frac = f"{abs(value):.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return sign + ",".join([*groups, tail]) + "." + frac


def fmt_value(value: Any, unit: str) -> str:
    if value is None:
        return "—"
    if unit == "month":
        return _month_label(value)
    if unit == "text":
        return str(value)
    number = Decimal(str(value))
    if unit == "inr":
        return "₹" + _indian(number)
    if unit == "pct":
        return f"{number:.2f}%"
    return _indian(number).removesuffix(".00")


def _pdf(report, spec, result, about) -> bytes:
    from fpdf import FPDF

    rows = result["rows"]
    if len(rows) > PDF_MAX_ROWS:
        raise ValueError(f"{len(rows):,} rows is more than a PDF can usefully hold ({PDF_MAX_ROWS:,}). "
                         "Generate Excel, or narrow the period or breakdown.")
    fonts = _font_dir()
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_font("DejaVu", "", str(fonts / "DejaVuSans.ttf"))
    bold = fonts / "DejaVuSans-Bold.ttf"
    pdf.add_font("DejaVu", "B", str(bold if bold.exists() else fonts / "DejaVuSans.ttf"))
    pdf.set_title(report.name)
    pdf.set_creator("PeopleOps Reports")
    pdf.add_page()
    pdf.set_font("DejaVu", "B", 15)
    pdf.cell(0, 9, report.name, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 8.5)
    pdf.set_text_color(71, 79, 97)
    for label in ("Dataset", "Entity", "Period from", "Period to", "Breakdown", "Filters",
                  "Requested by", "Generated at", "Definition version", "Record count"):
        pdf.cell(0, 4.6, f"{label}: {about[label]}", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 4.6, f"How it is counted: {about['How it is counted']}", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 4.6, about["Data basis"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(17, 21, 34)
    pdf.ln(2)

    if result["status"] != "ok":
        pdf.set_font("DejaVu", "B", 11)
        pdf.multi_cell(0, 6, "No figures: " + (
            "there is no data for this period." if result["status"] == "missing_data"
            else "nothing matched the filters."), new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())

    if result.get("chart"):
        _pdf_chart(pdf, result["chart"])
    if result.get("pivot"):
        p = result["pivot"]
        headers = ["Breakdown", *[_month_label(x) for x in p["periods"]]] + (["Total"] if p["row_totals"] else [])
        body = [[r["dimension"], *[fmt_value(r["cells"].get(x), p["unit"]) for x in p["periods"]]]
                + ([fmt_value(r["total"], p["unit"])] if p["row_totals"] else []) for r in p["rows"]]
        if p["column_totals"]:
            body.append(["Total", *[fmt_value(p["column_totals"].get(x), p["unit"]) for x in p["periods"]]]
                        + ([fmt_value(p["grand_total"], p["unit"])] if p["row_totals"] else []))
        _pdf_table(pdf, f"{p['label']} by breakdown and month", headers, body,
                   ["L", *["R"] * (len(headers) - 1)])
        if p.get("note"):
            pdf.set_font("DejaVu", "", 8)
            pdf.multi_cell(0, 4.4, p["note"], new_x="LMARGIN", new_y="NEXT")
    columns = result["columns"]
    _pdf_table(pdf, "Details", [c["label"] for c in columns],
               [[fmt_value(row[c["key"]], c["unit"]) for c in columns] for row in rows],
               ["L" if c["unit"] in ("text", "month") else "R" for c in columns])
    return bytes(pdf.output())


def _pdf_table(pdf, title: str, headers: list[str], body: list[list[str]], aligns: list[str]) -> None:
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    # Text columns (breakdown names) get more room than figures.
    weights = [2.2 if a == "L" and i > 0 else (1.6 if a == "L" else 1.0) for i, a in enumerate(aligns)]
    widths = [usable * w / sum(weights) for w in weights]
    size = 8 if len(headers) <= 9 else 6.5

    def header() -> None:
        pdf.set_font("DejaVu", "B", size)
        pdf.set_fill_color(238, 241, 246)
        for text, width, align in zip(headers, widths, aligns, strict=True):
            pdf.cell(width, 6, _fit(pdf, text, width), border="B", fill=True, align=align)
        pdf.ln()
        pdf.set_font("DejaVu", "", size)
        pdf.set_draw_color(230, 233, 239)

    pdf.ln(2)
    pdf.set_font("DejaVu", "B", 10.5)
    pdf.cell(0, 7, title, new_x="LMARGIN", new_y="NEXT")
    header()
    for line in body:
        if pdf.get_y() > pdf.h - pdf.b_margin - 6:
            pdf.add_page()
            header()
        for text, width, align in zip(line, widths, aligns, strict=True):
            pdf.cell(width, 5.2, _fit(pdf, text, width), border="B", align=align)
        pdf.ln()
    pdf.set_draw_color(214, 218, 227)


def _fit(pdf, text: str, width: float) -> str:
    text = str(text)
    while text and pdf.get_string_width(text) > width - 1.5:
        text = text[:-2] + "…" if len(text) > 2 else ""
    return text


_SERIES_RGB = [(2, 117, 179), (224, 124, 36), (12, 163, 12), (137, 84, 204), (208, 59, 59), (99, 108, 126)]


def _pdf_chart(pdf, chart: dict) -> None:
    pdf.set_font("DejaVu", "B", 10.5)
    title = chart["label"] + (f" · {_month_label(chart['period'])}" if chart["type"] == "bar" else " over time")
    pdf.cell(0, 7, title, new_x="LMARGIN", new_y="NEXT")
    x0, y0 = pdf.l_margin, pdf.get_y() + 1
    width, height = pdf.w - pdf.l_margin - pdf.r_margin, 62
    pdf.set_font("DejaVu", "", 7.5)
    pdf.set_draw_color(214, 218, 227)
    if chart["type"] == "bar":
        bars = [b for b in chart["bars"] if b["value"] is not None][:14]
        top = max([abs(b["value"]) for b in bars] + [1])
        label_w, row_h = 58, min(4.4, height / max(1, len(bars)))
        for i, bar in enumerate(bars):
            y = y0 + i * row_h
            pdf.set_xy(x0, y)
            pdf.cell(label_w, row_h, _fit(pdf, bar["group"], label_w))
            length = (width - label_w - 40) * abs(bar["value"]) / top
            pdf.set_fill_color(*_SERIES_RGB[0])
            pdf.rect(x0 + label_w, y + row_h * 0.18, max(length, 0.3), row_h * 0.64, style="F")
            pdf.set_xy(x0 + label_w + length + 1.5, y)
            pdf.cell(38, row_h, fmt_value(bar["value"], chart["unit"]))
        pdf.set_y(y0 + len(bars) * row_h + 3)
    else:
        periods, series = chart["periods"], chart["series"]
        values = [v for s in series for v in s["points"] if v is not None]
        low, high = min(values + [0]), max(values + [1])
        span = (high - low) or 1
        axis_w = 30
        plot_w, plot_h = width - 70 - axis_w, height - 8
        # A value scale: the axis starts at zero, as a cost axis must.
        for fraction in (0, 0.5, 1):
            y = y0 + plot_h - fraction * plot_h
            pdf.set_xy(x0, y - 2)
            pdf.cell(axis_w - 2, 4, fmt_value(low + fraction * span, chart["unit"]), align="R")
            pdf.set_draw_color(238, 240, 243)
            pdf.line(x0 + axis_w, y, x0 + axis_w + plot_w, y)
        pdf.set_draw_color(214, 218, 227)
        x0 += axis_w
        pdf.line(x0, y0 + plot_h, x0 + plot_w, y0 + plot_h)
        step = plot_w / max(1, len(periods) - 1)
        for i, period in enumerate(periods):
            if len(periods) <= 12 or i % max(1, len(periods) // 12) == 0:
                pdf.set_xy(x0 + i * step - 7, y0 + plot_h + 0.8)
                pdf.cell(14, 4, _month_label(period), align="C")
        for k, s in enumerate(series):
            pdf.set_draw_color(*_SERIES_RGB[k % len(_SERIES_RGB)])
            pdf.set_line_width(0.6)
            last = None
            for i, v in enumerate(s["points"]):
                if v is None:
                    last = None  # a gap stays a gap
                    continue
                point = (x0 + i * step, y0 + plot_h - (v - low) / span * plot_h)
                if last:
                    pdf.line(*last, *point)
                last = point
            pdf.set_xy(x0 + plot_w + 4, y0 + k * 5)
            pdf.set_fill_color(*_SERIES_RGB[k % len(_SERIES_RGB)])
            pdf.rect(x0 + plot_w + 1, y0 + k * 5 + 1.6, 2.4, 1.4, style="F")
            pdf.cell(64, 4.6, _fit(pdf, s["group"], 64))
        pdf.set_line_width(0.2)
        pdf.set_draw_color(214, 218, 227)
        pdf.set_y(y0 + height)
    if chart.get("omitted"):
        pdf.set_font("DejaVu", "", 7.5)
        pdf.cell(0, 4.5, _omitted(chart["omitted"], "the table"), new_x="LMARGIN", new_y="NEXT")
