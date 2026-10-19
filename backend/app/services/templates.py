"""Import templates, generated from this company's own configuration.

Every template is a spreadsheet, offered as CSV or Excel, and nothing else: the
columns are what the importer reads, and the salary components a company has set
up appear as columns the moment they are saved. A template carries one example row,
and the Excel version adds a sheet saying what each column means.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.models import ComponentConfig, Entity
from app.services.payroll_parse import allowed_destinations, normalize_col

KINDS = {
    "register": "salary-register",
    "ctc": "ctc",
    "attendance": "attendance",
    "master": "employee-master",
}


def _components(db: Session, entity: Entity, *, exclude_importer_fields: bool) -> list[str]:
    rows = (
        db.query(ComponentConfig).filter(ComponentConfig.entity_id == entity.id)
        .order_by(ComponentConfig.component_name).all()
    )
    reserved = allowed_destinations(set()) if exclude_importer_fields else set()
    return [r.component_name for r in rows if normalize_col(r.component_name) not in reserved]


def table(db: Session, entity: Entity, kind: str) -> dict[str, Any]:
    """``columns`` (name, meaning, example), in order. Raises KeyError for an unknown kind."""
    if kind == "register":
        components = _components(db, entity, exclude_importer_fields=True)
        columns = [
            ("Employee ID", "Unique employee code. Required.", "E001"),
            ("Employee Name", "Name as on the payslip.", "Sample Employee"),
            ("State", "Work state, used for professional tax and LWF.", entity.primary_state or "Karnataka"),
            ("Location", "Work location or branch.", ""),
            ("UAN", "12-digit PF universal account number.", ""),
            ("ESI Number", "ESIC insurance number.", ""),
            ("PAN", "Permanent account number.", ""),
            ("Aadhaar", "12-digit Aadhaar number.", ""),
            ("Date of Birth", "YYYY-MM-DD.", ""),
            ("Date of Joining", "YYYY-MM-DD.", ""),
            ("Total Days", "Days in the pay month.", "30"),
            ("Paid Days", "Days paid.", "30"),
            ("LOP Days", "Days of loss of pay.", "0"),
            *[(name, "Salary component amount for the month.", "0") for name in components],
            *[(f"{name} Arrear", "Arrear for this component paid this month (blank if none).", "0") for name in components],
            ("Increment Arrear", "Total increment arrear paid this month (blank if none).", "0"),
            ("Arrear From", "First month the arrear covers, YYYY-MM-DD (blank if none).", ""),
            ("Arrear To", "Last month the arrear covers, YYYY-MM-DD (blank if none).", ""),
            ("Gross", "Total earnings.", "0"),
            ("PF Employee", "Employee PF deducted.", "0"),
            ("ESIC Employee", "Employee ESIC deducted.", "0"),
            ("PT", "Professional tax deducted.", "0"),
            ("LWF Employee", "Labour welfare fund deducted.", "0"),
            ("TDS", "Income tax deducted.", "0"),
            ("Total Deductions", "All deductions.", "0"),
            ("Net", "Net pay.", "0"),
        ]
    elif kind == "ctc":
        components = _components(db, entity, exclude_importer_fields=False)
        columns = [
            ("Employee ID", "Unique employee code. Required.", "E001"),
            ("Employee Name", "Name.", "Sample Employee"),
            ("Effective From", "Date this CTC applies from, YYYY-MM-DD. Required.", "2026-04-01"),
            *[(name, "ANNUAL amount for this component.", "0") for name in components],
            ("Annual CTC", "Total annual CTC. Left out, it is the sum of the components.", "0"),
        ]
    elif kind == "attendance":
        columns = [
            ("Employee ID", "Unique employee code. Required.", "E001"),
            ("Employee Name", "Name.", "Sample Employee"),
            ("Total Days", "Calendar days in the month.", "30"),
            ("Present Days", "Days present.", "22"),
            ("Paid Leave Days", "Days of paid leave.", "2"),
            ("Weekly Off Days", "Weekly offs.", "4"),
            ("Holiday Days", "Public holidays.", "2"),
            ("LOP Days", "Loss of pay days.", "0"),
            ("Paid Days", "Days to be paid.", "30"),
            ("Overtime Hours", "Overtime hours worked.", "0"),
        ]
    elif kind == "master":
        columns = [
            ("Employee ID", "Unique employee code. Required.", "E001"),
            ("Employee Name", "Name.", "Sample Employee"),
            ("Date of Joining", "YYYY-MM-DD.", "2024-04-01"),
            ("Date of Exit", "YYYY-MM-DD; blank while employed.", ""),
            ("Date of Birth", "YYYY-MM-DD.", ""),
            ("Gender", "Male / Female / Other.", "Female"),
            ("State", "Work state.", entity.primary_state or "Karnataka"),
            ("Work Location", "Branch or site.", ""),
            ("Department", "Department.", ""),
            ("Designation", "Job title.", ""),
            ("Grade", "Grade or band.", ""),
            ("Cost Centre", "Cost centre code.", ""),
            ("Employment Type", "Permanent / Contract / Trainee.", "Permanent"),
            ("UAN", "12-digit PF universal account number.", ""),
            ("ESIC IP Number", "ESIC insurance number.", ""),
            ("PAN", "Permanent account number.", ""),
            ("Aadhaar", "12-digit Aadhaar number.", ""),
            ("Bank Account", "Salary account number.", ""),
            ("IFSC", "Bank IFSC code.", ""),
        ]
    else:
        raise KeyError(kind)
    return {"kind": kind, "file": KINDS[kind], "columns": columns}


def as_csv(spec: dict[str, Any]) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([c[0] for c in spec["columns"]])
    writer.writerow([c[2] for c in spec["columns"]])
    return ("﻿" + out.getvalue()).encode("utf-8")


def as_xlsx(spec: dict[str, Any]) -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Template"
    sheet.append([c[0] for c in spec["columns"]])
    sheet.append([c[2] for c in spec["columns"]])
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for index, (name, _, _) in enumerate(spec["columns"], 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(14, len(name) + 4)
    sheet.freeze_panes = "A2"
    guide = wb.create_sheet("How to fill")
    guide.append(["Column", "What to put there"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for name, meaning, _ in spec["columns"]:
        guide.append([name, meaning])
    guide.append([])
    guide.append(["Row 2 is an example; replace or delete it. Keep the headings exactly as they are."])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
