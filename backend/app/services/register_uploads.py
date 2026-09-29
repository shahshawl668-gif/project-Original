"""
Freezing an upload, and reading it back.

See ``app/models/register_upload.py`` for why every upload is kept. This module
owns the two things that must not drift: how rows are written, and how they are
digested. A run's reproducibility rests on reading back exactly what was
validated, so the encoding lives in one place.
"""
from __future__ import annotations

import gzip
import hashlib
import io
import json
import uuid
from datetime import date
from typing import Any

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models import RegisterUpload

#: The key each parsed row carries for "where in the file did this come from".
#: Leading underscore: the engine ignores such keys, so it never reads as an
#: unmapped salary column.
SOURCE_ROW_KEY = "_source_row"


def _canonical(value: Any) -> Any:
    """JSON-safe, deterministic. Rows are already plain types; this is a guard."""
    if isinstance(value, dict):
        return {str(k): _canonical(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_canonical(v) for v in value]
    if isinstance(value, float) and value != value:  # NaN
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def encode_rows(rows: list[dict[str, Any]]) -> tuple[bytes, str]:
    """Compressed rows, and the digest of their canonical form.

    The digest is over sorted-key JSON so that the same rows always hash the
    same, whatever order a dict happened to be built in. ``mtime=0`` keeps the
    compressed bytes themselves deterministic too.
    """
    canonical = [_canonical(r) for r in rows]
    text = json.dumps(canonical, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb", mtime=0) as gz:
        gz.write(text.encode("utf-8"))
    return buf.getvalue(), digest


def decode_rows(blob: bytes) -> list[dict[str, Any]]:
    return json.loads(gzip.decompress(blob).decode("utf-8"))


def gzip_json(value: Any) -> bytes:
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb", mtime=0) as gz:
        gz.write(json.dumps(_canonical(value), separators=(",", ":"), default=str).encode("utf-8"))
    return buf.getvalue()


def gunzip_json(blob: bytes | None) -> Any:
    if not blob:
        return None
    return json.loads(gzip.decompress(blob).decode("utf-8"))


def stamp_source_rows(rows: list[dict[str, Any]]) -> None:
    """Record each row's position in the file: the header is row 1."""
    for index, row in enumerate(rows):
        row[SOURCE_ROW_KEY] = index + 2


def first_sheet_name(content: bytes, filename: str | None) -> str | None:
    """The sheet an .xlsx upload was read from; None for CSV."""
    if not (filename or "").lower().endswith(".xlsx"):
        return None
    try:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content), read_only=True)
        try:
            return book.sheetnames[0] if book.sheetnames else None
        finally:
            book.close()
    except Exception:  # noqa: BLE001 — the name is a label, not a requirement
        return None


def record(
    db: Session,
    *,
    entity_id: uuid.UUID,
    user_id: uuid.UUID | None,
    period_month: date | None,
    run_type: str,
    filename: str | None,
    content: bytes,
    rows: list[dict[str, Any]],
    source_columns: list[str] | None,
    column_mapping: dict[str, Any] | None,
    missing_required: list[str] | None,
    warnings: list[str] | None,
    register_id: uuid.UUID | None = None,
) -> RegisterUpload:
    """Write one upload. Never updates an existing one."""
    revision = 1
    if period_month is not None:
        revision = (
            db.query(func.count(RegisterUpload.id))
            .filter(
                RegisterUpload.entity_id == entity_id,
                RegisterUpload.period_month == period_month,
            )
            .scalar()
            or 0
        ) + 1
    blob, digest = encode_rows(rows)
    upload = RegisterUpload(
        entity_id=entity_id,
        user_id=user_id,
        period_month=period_month,
        register_id=register_id,
        revision=revision,
        run_type=run_type or "regular",
        filename=filename,
        sheet_name=first_sheet_name(content, filename),
        file_sha256=hashlib.sha256(content).hexdigest(),
        file_size=len(content),
        source_columns=list(source_columns or []),
        column_mapping=column_mapping if isinstance(column_mapping, dict) else None,
        missing_required=list(missing_required or []),
        warnings=list(warnings or [])[:50],
        row_count=len(rows),
        rows_sha256=digest,
        rows_gz=blob,
    )
    db.add(upload)
    db.flush()
    return upload


def latest_for_period(
    db: Session, entity_id: uuid.UUID, period_month: date, run_type: str | None = "regular",
) -> RegisterUpload | None:
    """The newest upload of the month's register — by default the regular run.

    An arrears file for June is not June's register: letting it count as the
    latest upload made June's run look stale and made the arrears file the
    default for June's next validation. Pass ``run_type=None`` for any kind.
    """
    query = db.query(RegisterUpload).filter(
        RegisterUpload.entity_id == entity_id,
        RegisterUpload.period_month == period_month,
    )
    if run_type is not None:
        query = query.filter(func.coalesce(RegisterUpload.run_type, "regular") == run_type)
    return query.order_by(RegisterUpload.revision.desc()).first()


def describe(upload: RegisterUpload) -> dict[str, Any]:
    return {
        "id": str(upload.id),
        "period_month": upload.period_month.isoformat() if upload.period_month else None,
        "revision": upload.revision,
        "run_type": upload.run_type,
        "filename": upload.filename,
        "sheet_name": upload.sheet_name,
        "file_sha256": upload.file_sha256,
        "file_size": upload.file_size,
        "row_count": upload.row_count,
        "rows_sha256": upload.rows_sha256,
        "missing_required": upload.missing_required or [],
        "warnings": upload.warnings or [],
        "stored_as_register": upload.register_id is not None,
        "uploaded_at": upload.created_at.isoformat() if upload.created_at else None,
    }
