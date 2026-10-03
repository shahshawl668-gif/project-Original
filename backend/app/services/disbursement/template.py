"""
Bank payment file templates: reading a file by its layout, and writing the clean copy.

A template is data — the columns and their order, the delimiter, an optional
header record and footer record with the values they state, and how dates and
amounts are written. Adding a bank format is adding a template, not code.

The templates shipped here are generic and synthetic. None is any bank's
official specification, and none should be labelled as one: a company's bank
format is whatever its corporate banking channel was set up to accept, and a
template must be checked against a real file from that channel before use.

The clean copy
--------------
Kept payment lines are copied **byte for byte** from the input — same order,
same quoting, same line endings. Only what must change is rewritten: held lines
are left out, and the header and footer records' count and total are
recalculated. Everything else in those records is kept as it was. A workbook is
rebuilt cell by cell, with fixed timestamps so the same input always produces
the same bytes, and so the same SHA-256 fingerprint.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.services.bank_file import match_key, parse_amount, parse_bank_date
from app.services.disbursement import values
from app.services.disbursement.fields import BANK_COLUMNS, RECORD_FIELDS
from app.services.disbursement.model import BankFile, BankRow

TEMPLATE_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "disbursement" / "templates"
LAYOUTS = ("delimited", "excel")


class TemplateError(ValueError):
    """The file does not fit the template, or the template is not usable. Says where."""


def builtin_templates() -> dict[str, dict[str, Any]]:
    out = {}
    for path in sorted(TEMPLATE_DIR.glob("*.json")):
        spec = json.loads(path.read_text(encoding="utf-8"))
        out[spec["key"]] = check_template(spec)
    return out


def check_template(spec: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise TemplateError("A template is an object.")
    layout = spec.get("layout", "delimited")
    if layout not in LAYOUTS:
        raise TemplateError(f"layout must be one of {', '.join(LAYOUTS)}")
    columns = spec.get("columns")
    if not isinstance(columns, list) or not columns:
        raise TemplateError("A template lists its columns.")
    fields = [c.get("field") for c in columns if isinstance(c, dict)]
    for need in ("employee_id", "amount"):
        if need not in fields:
            raise TemplateError(f"A payment line must have a column for {need}.")
    for c in columns:
        f = c.get("field")
        if f not in BANK_COLUMNS and f != "ignore":
            raise TemplateError(f"Unknown column field {f!r}; use one of {', '.join(BANK_COLUMNS)} or ignore.")
    if spec.get("column_header") and not all(c.get("header") for c in columns):
        raise TemplateError("With a column header row, every column needs its header name.")
    for which in ("header_record", "footer_record"):
        rec = spec.get(which)
        if rec is None:
            continue
        for f in rec.get("fields") or []:
            if f.get("name") not in RECORD_FIELDS:
                raise TemplateError(f"{which}: unknown field {f.get('name')!r}")
            if not isinstance(f.get("position"), int) or f["position"] < 0:
                raise TemplateError(f"{which}: {f.get('name')} needs a position (0 = first)")
    amount = spec.get("amount") or {}
    if amount.get("unit", "rupees") not in ("rupees", "paise"):
        raise TemplateError("amount.unit is rupees or paise")
    if layout == "delimited" and len(spec.get("delimiter", ",")) != 1:
        raise TemplateError("delimiter is a single character")
    return spec


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------
@dataclass
class _Line:
    number: int          # 1-based position in the file
    raw: Any             # the text line with its ending, or the worksheet row's values
    cells: list[Any]


@dataclass
class Parsed:
    bank: BankFile
    lines: list[_Line]
    header_line: _Line | None
    column_line: _Line | None
    data_lines: list[_Line]
    footer_line: _Line | None
    eol: str
    ends_with_newline: bool
    encoding: str
    bom: bool = False


def _decode(content: bytes, encoding: str) -> tuple[str, str]:
    for enc in (encoding, "utf-8-sig", "cp1252"):
        try:
            text = content.decode(enc)
            return (text[1:] if text.startswith("﻿") else text), enc
        except (UnicodeDecodeError, LookupError):
            continue
    raise TemplateError("The file's text encoding could not be read.")


def _excel_rows(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    from app.services.upload_safety import check_workbook

    check_workbook(content)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # a damaged or non-workbook zip is refused, not a 500
        raise TemplateError("The file is not a readable .xlsx workbook.") from exc
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    while rows and all(v is None or str(v).strip() == "" for v in rows[-1]):
        rows.pop()
    return rows


def _cell_text(value: Any) -> str | None:
    if isinstance(value, float) and value == int(value) and abs(value) < 1e15:
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return values.text(value)


def _record(line: _Line, rec: dict[str, Any], where: str, fmt: str | None, unit: str) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for f in rec.get("fields") or []:
        pos, name = f["position"], f["name"]
        raw = line.cells[pos] if pos < len(line.cells) else None
        if f.get("literal") is not None and (values.text(raw) or "") != str(f["literal"]):
            raise TemplateError(f"Line {line.number} should be the {where} record (starting "
                                f"'{f['literal']}'), but reads {values.text(raw)!r}.")
        if name == "record_count":
            text = _cell_text(raw)
            out[name] = int(text) if text and text.isdigit() else None
            if text and not text.isdigit():
                out["record_count_raw"] = text
        elif name == "total_amount":
            out[name] = parse_amount(raw, unit=unit)
        elif name == "value_date":
            out[name] = raw if isinstance(raw, (date, datetime)) else parse_bank_date(_cell_text(raw), f.get("format") or fmt)
            out["value_date_raw"] = _cell_text(raw)
        else:
            out[name] = _cell_text(raw)
    return out


def read(content: bytes, filename: str, template: dict[str, Any]) -> Parsed:
    t = check_template(template)
    layout = t.get("layout", "delimited")
    encoding = t.get("encoding", "utf-8")
    eol, ends = "\n", True
    lines: list[_Line] = []
    if layout == "excel":
        for i, row in enumerate(_excel_rows(content), start=1):
            lines.append(_Line(i, row, row))
        used_encoding = encoding
    else:
        text, used_encoding = _decode(content, encoding)
        raw_lines = text.splitlines(keepends=True)
        if raw_lines:
            first = raw_lines[0]
            eol = "\r\n" if first.endswith("\r\n") else "\r" if first.endswith("\r") else "\n"
            ends = raw_lines[-1].endswith(("\n", "\r"))
        delim = t.get("delimiter", ",")
        for i, raw in enumerate(raw_lines, start=1):
            body = raw.rstrip("\r\n")
            if not body.strip():
                continue
            cells = next(csv.reader([body], delimiter=delim)) if body else []
            lines.append(_Line(i, raw, cells))
    if not lines:
        raise TemplateError("The file is empty.")

    rest = list(lines)
    header_line = rest.pop(0) if t.get("header_record") else None
    if t.get("footer_record"):
        if not rest:
            raise TemplateError("The file has a header record but no footer record.")
        footer_line = rest.pop()
    else:
        footer_line = None
    column_line = rest.pop(0) if t.get("column_header") and rest else None

    columns = t["columns"]
    if column_line is not None:
        found = {(_cell_text(c) or "").lower(): i for i, c in enumerate(column_line.cells)}
        positions = {}
        missing = []
        for c in columns:
            idx = found.get(str(c["header"]).strip().lower())
            if idx is None:
                if c["field"] in ("employee_id", "amount"):
                    missing.append(c["header"])
                continue
            positions[c["field"]] = idx
        if missing:
            raise TemplateError(f"The file's column header has no {', '.join(repr(m) for m in missing)} column.")
    else:
        positions = {c["field"]: i for i, c in enumerate(columns) if c["field"] != "ignore"}
    literals = [(i, str(c["literal"])) for i, c in enumerate(columns) if c.get("literal") is not None]

    fmt = t.get("date_format")
    unit = (t.get("amount") or {}).get("unit", "rupees")
    transform = t.get("employee_id_transform", "trim")
    rows: list[BankRow] = []
    for n, line in enumerate(rest, start=1):
        for i, lit in literals:
            if (values.text(line.cells[i]) if i < len(line.cells) else None) != lit:
                raise TemplateError(f"Line {line.number} is not a payment record (expected '{lit}' in "
                                    f"position {i + 1}).")

        def cell(f: str, cells: list[Any] = line.cells) -> Any:
            i = positions.get(f)
            return cells[i] if i is not None and i < len(cells) else None

        emp = _cell_text(cell("employee_id"))
        raw_amount = cell("amount")
        rows.append(BankRow(
            row=n, line=line.number, employee_id=emp, key=match_key(emp, transform) if emp else None,
            beneficiary_name=_cell_text(cell("beneficiary_name")),
            account_number=values.account(_cell_text(cell("account_number"))),
            ifsc=values.ifsc(cell("ifsc")),
            amount=parse_amount(raw_amount, unit=unit),
            amount_raw=_cell_text(raw_amount),
            reference=_cell_text(cell("reference")),
        ))
    header = _record(header_line, t["header_record"], "header", fmt, unit) if header_line else None
    footer = _record(footer_line, t["footer_record"], "footer", fmt, unit) if footer_line else None
    present = set(positions)
    bank = BankFile(rows=rows, columns=present, header=header, footer=footer,
                    has_header_record=header_line is not None, filename=filename)
    bom = layout != "excel" and content.startswith(b"\xef\xbb\xbf")
    return Parsed(bank, lines, header_line, column_line, rest, footer_line, eol, ends, used_encoding, bom)


# ---------------------------------------------------------------------------
# Writing the clean copy
# ---------------------------------------------------------------------------
def format_amount(value: Decimal, template: dict[str, Any]) -> str:
    amount = template.get("amount") or {}
    if amount.get("unit", "rupees") == "paise":
        return str(int((value * 100).quantize(Decimal("1"))))
    decimals = int(amount.get("decimals", 2))
    return f"{value:.{decimals}f}"


def _rewrite(line: _Line, rec: dict[str, Any], count: int, total: Decimal, template: dict[str, Any]) -> list[Any]:
    cells = list(line.cells)
    for f in rec.get("fields") or []:
        pos = f["position"]
        while len(cells) <= pos:
            cells.append("")
        if f["name"] == "record_count":
            cells[pos] = str(count)
        elif f["name"] == "total_amount":
            cells[pos] = format_amount(total, template)
    return cells


def _deterministic_zip(data: bytes) -> bytes:
    """Rewrite a workbook's zip with fixed timestamps, so equal content gives equal bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, src.read(name))
    return out.getvalue()


def write_clean(parsed: Parsed, kept_lines: set[int], template: dict[str, Any]) -> bytes:
    """The input file without the held lines, its control totals recalculated."""
    kept = [line for line in parsed.data_lines if line.number in kept_lines]
    by_line = {r.line: r for r in parsed.bank.rows}
    total = sum((by_line[line.number].amount or Decimal("0") for line in kept), Decimal("0"))
    count = len(kept)
    if template.get("layout", "delimited") == "excel":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Payments"
        if parsed.header_line is not None:
            ws.append(_rewrite(parsed.header_line, template["header_record"], count, total, template))
        if parsed.column_line is not None:
            ws.append(list(parsed.column_line.raw))
        for line in kept:
            ws.append(list(line.raw))
        if parsed.footer_line is not None:
            ws.append(_rewrite(parsed.footer_line, template["footer_record"], count, total, template))
        fixed = datetime(2000, 1, 1)
        wb.properties.created = fixed
        wb.properties.modified = fixed
        wb.properties.creator = "PeopleOpsLab"
        buf = io.BytesIO()
        wb.save(buf)
        return _deterministic_zip(buf.getvalue())

    delim = template.get("delimiter", ",")
    eol = parsed.eol

    def serialise(cells: list[Any]) -> str:
        buf = io.StringIO()
        csv.writer(buf, delimiter=delim, lineterminator="").writerow(cells)
        return buf.getvalue() + eol

    parts: list[str] = []
    if parsed.header_line is not None:
        parts.append(serialise(_rewrite(parsed.header_line, template["header_record"], count, total, template)))
    if parsed.column_line is not None:
        parts.append(parsed.column_line.raw if parsed.column_line.raw.endswith(("\n", "\r")) else parsed.column_line.raw + eol)
    for line in kept:
        parts.append(line.raw if line.raw.endswith(("\n", "\r")) else line.raw + eol)
    if parsed.footer_line is not None:
        parts.append(serialise(_rewrite(parsed.footer_line, template["footer_record"], count, total, template)))
    text = "".join(parts)
    if not parsed.ends_with_newline and text.endswith(eol):
        text = text[: -len(eol)]
    out = text.encode(parsed.encoding if parsed.encoding != "utf-8-sig" else "utf-8")
    return (b"\xef\xbb\xbf" + out) if parsed.bom else out


def fingerprint(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
