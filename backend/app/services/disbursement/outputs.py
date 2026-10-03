"""
What a disbursement check hands to people: the exception report and the
approver summary.

Both are built from the *report* — a plain dictionary the run produces and
the database stores — never from live objects, so a report downloaded a week
later is the same document as the one read on the day.

Account numbers and IFSCs are shown in full. Masking them would defeat the
point: the approver is being asked whether these are the right accounts.

Nothing is left out because it passed or because it could not be checked: a
check that did not run, or that the company switched off, is a row in the
exception report and a line in the summary, never a silence.
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.services.disbursement.config import (
    CLEAR,
    DISABLED,
    DO_NOT_RELEASE,
    FLAG,
    HOLD_ROW,
    NOT_RUN,
    STOP_FILE,
    WITH_HOLDS,
)
from app.services.disbursement.engine import Result

VERDICT_TEXT = {
    CLEAR: "Clear to release",
    WITH_HOLDS: "Release with holds",
    DO_NOT_RELEASE: "Do not release",
}
VERDICT_MEANING = {
    CLEAR: "Every check that ran passed or only raised flags. The clean file is the bank file as uploaded.",
    WITH_HOLDS: "Some employees are held back. The clean file pays everyone else; the held lines are listed "
                "with their reasons.",
    DO_NOT_RELEASE: "A problem affects the whole file. No clean file is produced: fix the cause and check again.",
}
SEVERITY_TEXT = {STOP_FILE: "Stops the file", HOLD_ROW: "Held", FLAG: "Flag", NOT_RUN: "Not run",
                 DISABLED: "Switched off"}

FINDING_COLUMNS = [
    ("rule_id", "Check"), ("severity", "Severity"), ("employee_id", "Employee ID"),
    ("employee_name", "Employee name"), ("field", "Field"), ("expected", "Expected"),
    ("actual", "Actual"), ("amount", "Amount in file"), ("rows", "Bank file rows"), ("reason", "Reason"),
]
_FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu")


# ---------------------------------------------------------------------------
# The report: what is stored, and what every output is drawn from
# ---------------------------------------------------------------------------
def build_report(result: Result, meta: dict[str, Any], bank_rows: list | None = None) -> dict[str, Any]:
    """
    Freeze a run into plain data. ``meta`` carries who, when, what and the
    fingerprints; ``bank_rows`` gives each held line's own amount.
    """
    summary = result.summary()
    findings = [f.as_dict() for f in result.findings]
    # A switched-off check is reported as such: the company chose not to look, and the
    # approver should see that choice rather than an absence.
    for rule in result.rules:
        if rule.status == DISABLED:
            findings.append({"rule_id": rule.rule_id, "severity": DISABLED, "employee_id": "",
                             "employee_name": "", "field": "", "expected": "", "actual": "",
                             "reason": rule.reason, "rows": [], "amount": ""})
    return {
        **meta,
        "verdict": result.verdict,
        "totals": summary["totals"],
        "bridge": summary["bridge"],
        "rules": summary["rules"],
        "findings": findings,
        "held_lines": [{"row": r.row, "line": r.line, "employee_id": r.employee_id or "",
                        "amount": str(r.amount) if r.amount is not None else ""}
                       for r in (bank_rows or []) if r.line in result.held_lines],
        "notes": [{"source": n.source, "message": n.message, "row": n.row, "employee_id": n.employee_id}
                  for n in result.notes],
    }


def held_employees(report: dict[str, Any]) -> list[dict[str, Any]]:
    """One line per held employee: the amount held and every reason they were held."""
    line_amount = {h["row"]: Decimal(h["amount"]) for h in report.get("held_lines", []) if h["amount"]}
    by_emp: dict[str, dict[str, Any]] = {}
    for f in report["findings"]:
        if f["severity"] != HOLD_ROW:
            continue
        key = f["employee_id"] or f"row {','.join(str(r) for r in f['rows'])}"
        entry = by_emp.setdefault(key, {"employee_id": f["employee_id"], "employee_name": f["employee_name"],
                                        "rows": set(), "reasons": []})
        entry["employee_name"] = entry["employee_name"] or f["employee_name"]
        entry["rows"].update(f["rows"])
        entry["reasons"].append(f"{f['rule_id']}: {f['reason']}")
    # The hold unit is the employee: every line of theirs is held, including any a finding did not name.
    for h in report.get("held_lines", []):
        if h["employee_id"] in by_emp:
            by_emp[h["employee_id"]]["rows"].add(h["row"])
    out = []
    for entry in by_emp.values():
        out.append({"employee_id": entry["employee_id"], "employee_name": entry["employee_name"],
                    "rows": sorted(entry["rows"]),
                    "amount": sum((line_amount.get(r, Decimal("0")) for r in entry["rows"]), Decimal("0")),
                    "reasons": entry["reasons"]})
    return sorted(out, key=lambda e: (e["employee_id"], e["rows"][:1]))


# ---------------------------------------------------------------------------
# Exception report — CSV
# ---------------------------------------------------------------------------
def _csv_text(value: Any) -> Any:
    """Free text that a spreadsheet would run as a formula is written with a leading apostrophe."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


def _cell(finding: dict[str, Any], key: str) -> Any:
    value = finding.get(key)
    if key == "rows":
        return " ".join(str(r) for r in value or [])
    if key == "amount":
        return Decimal(value) if value else ""
    return value if value is not None else ""


def exception_csv(report: dict[str, Any]) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([label for _, label in FINDING_COLUMNS])
    for f in report["findings"]:
        writer.writerow([_csv_text(_cell(f, key)) if key != "amount" else _cell(f, key)
                         for key, _ in FINDING_COLUMNS])
    # UTF-8 with a byte-order mark so Excel reads ₹ and names correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


# ---------------------------------------------------------------------------
# Exception report — Excel
# ---------------------------------------------------------------------------
def exception_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from app.services.export_safety import neutralise_workbook

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2A44")

    def table(ws, headers: list[str], rows: list[list[Any]], widths: list[int], text_cols=()) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for row in rows:
            ws.append(row)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        for col in text_cols:   # account numbers and IDs stay text, never 1.23E+15
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "@"
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Summary"
    totals = report["totals"]
    summary_rows = [
        ("Verdict", VERDICT_TEXT.get(report["verdict"], report["verdict"])),
        ("What it means", VERDICT_MEANING.get(report["verdict"], "")),
        ("Period", report.get("period", "")),
        ("Company", report.get("entity_name", "")),
        ("Checked by", report.get("run_by", "")),
        ("Checked at", report.get("generated_at", "")),
        ("Amounts checked against", totals.get("amount_basis", "")),
        ("To release — amount", _money(totals.get("release_amount"))),
        ("To release — employees", totals.get("release_employees")),
        ("Held — amount", _money(totals.get("held_amount"))),
        ("Held — employees", totals.get("held_employees")),
        ("Flags", totals.get("flags")),
        ("Clean file", report.get("clean_filename") or "None — the file must not be released"),
        ("Clean file SHA-256", report.get("clean_sha256") or ""),
    ]
    table(ws, ["Item", "Value"], [list(r) for r in summary_rows], [28, 90])
    for (cell,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    table(wb.create_sheet("Findings"), [label for _, label in FINDING_COLUMNS],
          [[_cell(f, key) for key, _ in FINDING_COLUMNS] for f in report["findings"]],
          [9, 14, 14, 24, 16, 22, 22, 14, 12, 80], text_cols=(3, 6, 7))
    table(wb.create_sheet("Checks"), ["Check", "Title", "Severity", "Status", "Findings", "Reason"],
          [[r["rule_id"], r["title"], r["severity"], r["status"], r["findings"], r["reason"]]
           for r in report["rules"]], [9, 48, 12, 16, 10, 80])
    table(wb.create_sheet("Bridge"), ["Step", "Amount", "Lines"],
          [[b["label"], Decimal(b["amount"]) if b.get("amount") is not None else None, b.get("count")]
           for b in report["bridge"]], [90, 18, 10])
    table(wb.create_sheet("Inputs"), ["Input", "File", "Rows", "SHA-256"],
          [[i.get("label", ""), i.get("filename", ""), i.get("rows"), i.get("sha256", "")]
           for i in report.get("inputs", [])], [28, 40, 10, 70])
    table(wb.create_sheet("Notes"), ["Source", "Row", "Employee ID", "Note"],
          [[n["source"], n["row"], n["employee_id"] or "", n["message"]] for n in report["notes"]],
          [18, 8, 14, 100], text_cols=(3,))
    for name in ("Bridge",):
        for (cell,) in wb[name].iter_rows(min_row=2, min_col=2, max_col=2):
            cell.number_format = "#,##0.00"
    for (cell,) in wb["Findings"].iter_rows(min_row=2, min_col=8, max_col=8):
        cell.number_format = "#,##0.00"
    neutralise_workbook(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Approver summary — PDF
# ---------------------------------------------------------------------------
def _indian(value: Decimal) -> str:
    sign = "-" if value < 0 else ""
    whole, frac = f"{abs(value):.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    groups: list[str] = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return sign + ",".join([*groups, tail]) + "." + frac


def _money(value: Any) -> str:
    if value is None or value == "":
        return "—"
    number = Decimal(str(value))
    return ("-₹" if number < 0 else "₹") + _indian(abs(number))


def _font_dir() -> Path:
    for folder in _FONT_DIRS:
        if (Path(folder) / "DejaVuSans.ttf").exists():
            return Path(folder)
    raise ValueError("The approver summary needs the DejaVu font on the server (fonts-dejavu-core), "
                     "which is not installed. Download the Excel exception report instead.")


def approver_pdf(report: dict[str, Any], approval: dict[str, Any] | None = None) -> bytes:
    from fpdf import FPDF

    fonts = _font_dir()
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_font("DejaVu", "", str(fonts / "DejaVuSans.ttf"))
    bold = fonts / "DejaVuSans-Bold.ttf"
    pdf.add_font("DejaVu", "B", str(bold if bold.exists() else fonts / "DejaVuSans.ttf"))
    pdf.set_title(f"Salary payment check {report.get('period', '')}")
    pdf.set_creator("Peopleopslab disbursement validation")
    # A fixed creation date keeps the same report producing the same bytes.
    pdf.set_creation_date(_stamp(report.get("generated_at")))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    totals = report["totals"]
    verdict = report["verdict"]

    def heading(text: str) -> None:
        pdf.ln(3)
        pdf.set_font("DejaVu", "B", 11)
        pdf.cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 9)

    def para(text: str, size: float = 9) -> None:
        pdf.set_font("DejaVu", "", size)
        pdf.multi_cell(0, 4.8, text, new_x="LMARGIN", new_y="NEXT")

    def pairs(items: list[tuple[str, str]]) -> None:
        pdf.set_font("DejaVu", "", 9)
        for label, value in items:
            pdf.set_font("DejaVu", "B", 9)
            pdf.cell(58, 5.4, label)
            pdf.set_font("DejaVu", "", 9)
            pdf.multi_cell(width - 58, 5.4, str(value), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("DejaVu", "B", 15)
    pdf.cell(0, 9, "Salary payment file — approver summary", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 8.5)
    pdf.set_text_color(71, 79, 97)
    para(f"{report.get('entity_name', '')} · period {report.get('period', '')} · checked by "
         f"{report.get('run_by', '')} at {report.get('generated_at', '')}", 8.5)
    para("This product checks the payment file; it does not pay anyone or send anything to a bank. "
         "A person approves, and a person uploads the clean file to the bank.", 8.5)
    pdf.set_text_color(17, 21, 34)

    pdf.ln(2)
    colour = {CLEAR: (12, 120, 52), WITH_HOLDS: (166, 98, 0), DO_NOT_RELEASE: (180, 35, 35)}.get(verdict, (0, 0, 0))
    pdf.set_fill_color(*colour)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 11, "  " + VERDICT_TEXT.get(verdict, verdict).upper(), fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(17, 21, 34)
    para(VERDICT_MEANING.get(verdict, ""))

    heading("What the clean file pays")
    variance = totals.get("variance_vs_previous")
    pct = totals.get("variance_pct_vs_previous")
    pairs([
        ("To release", _money(totals.get("release_amount")) if verdict != DO_NOT_RELEASE else "Nothing — no clean file"),
        ("Employees paid", str(totals.get("release_employees")) if verdict != DO_NOT_RELEASE else "—"),
        ("Would be held" if verdict == DO_NOT_RELEASE else "Held back",
         f"{_money(totals.get('held_amount'))} · {totals.get('held_employees')} employee(s)"),
        ("Flags to read", str(totals.get("flags"))),
        ("Amounts checked against", totals.get("amount_basis") or "—"),
        ("Previous period paid",
         f"{_money(totals.get('previous_total'))} ({totals.get('previous_total_basis')})"
         if totals.get("previous_total") is not None else "Not checked — no previous period was given"),
        ("Change vs previous",
         "—" if verdict == DO_NOT_RELEASE or variance is None else f"{_money(variance)} ({pct}%)"),
        ("Clean file", report.get("clean_filename") or "None"),
        ("Clean file SHA-256", report.get("clean_sha256") or "—"),
    ])

    heading("From the register to the clean file")
    pdf.set_font("DejaVu", "", 8.5)
    for b in report["bridge"]:
        style = "B" if b.get("total") else ""
        pdf.set_font("DejaVu", style, 8.5)
        if pdf.get_string_width(b["label"]) > width - 54:
            pdf.set_font("DejaVu", style, 7)
        pdf.cell(width - 52, 5, pdf_fit(pdf, b["label"], width - 52))
        pdf.set_font("DejaVu", style, 8.5)
        pdf.cell(32, 5, _money(b.get("amount")), align="R")
        pdf.cell(20, 5, str(b.get("count", "")), align="R", new_x="LMARGIN", new_y="NEXT")

    stopped = verdict == DO_NOT_RELEASE
    if stopped:
        para("The file is stopped, so nothing is released: the last two lines show what would be held "
             "and paid once the cause is fixed.", 8)

    stops = [f for f in report["findings"] if f["severity"] == STOP_FILE]
    if stops:
        heading(f"Why the file must not be released ({len(stops)})")
        for f in stops[:40]:
            para(f"• {f['rule_id']}{' · ' + f['employee_id'] if f['employee_id'] else ''}: {f['reason']}")
        if len(stops) > 40:
            para(f"… and {len(stops) - 40} more in the exception report.")

    held = held_employees(report)
    heading(("Would also be held" if stopped else "Held back")
            + f" ({len(held)} employee(s), {_money(totals.get('held_amount'))})")
    if not held:
        para("No one is held.")
    for h in held[:60]:
        pdf.set_font("DejaVu", "B", 8.5)
        pdf.multi_cell(0, 4.6, f"{h['employee_id'] or '—'} {h['employee_name']} · {_money(h['amount'])}",
                       new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 8)
        for reason in h["reasons"]:
            pdf.multi_cell(0, 4.2, "   " + reason, new_x="LMARGIN", new_y="NEXT")
    if len(held) > 60:
        para(f"… and {len(held) - 60} more in the exception report.")

    flags = [f for f in report["findings"] if f["severity"] == FLAG]
    heading(f"Flags ({len(flags)})")
    para("Flags do not hold anyone. Read them before approving." if flags else "No flags.")
    by_rule: dict[str, int] = {}
    for f in flags:
        by_rule[f["rule_id"]] = by_rule.get(f["rule_id"], 0) + 1
    titles = {r["rule_id"]: r["title"] for r in report["rules"]}
    for rule_id, n in sorted(by_rule.items()):
        para(f"• {rule_id} {titles.get(rule_id, '')}: {n}")

    unchecked = [r for r in report["rules"] if r["status"] in (NOT_RUN, DISABLED)]
    heading(f"Checks that did not run ({len(unchecked)})")
    if not unchecked:
        para("Every check ran.")
    else:
        para("These were NOT checked. Approving means accepting that, and each must be acknowledged.")
        for r in unchecked:
            para(f"• {r['rule_id']} {r['title']} — {SEVERITY_TEXT.get(r['status'], r['status'])}: {r['reason']}")

    if report.get("inputs"):
        heading("Files checked")
        pdf.set_font("DejaVu", "", 7.5)
        for i in report["inputs"]:
            pdf.multi_cell(0, 4, f"{i.get('label')}: {i.get('filename')} · {i.get('rows')} rows · "
                                 f"SHA-256 {i.get('sha256')}", new_x="LMARGIN", new_y="NEXT")

    heading("Approval")
    if approval:
        pairs([("Approved by", approval.get("approver", "")), ("Approved at", approval.get("approved_at", "")),
               ("Clean file SHA-256", approval.get("fingerprint", "")),
               ("Checks not run, acknowledged", ", ".join(approval.get("acknowledged", [])) or "None")])
        if approval.get("comment"):
            para("Comment: " + approval["comment"])
    elif verdict == DO_NOT_RELEASE:
        para("This file cannot be approved.")
    else:
        para("Not yet approved. Approve in the product, or sign below.")
        pdf.ln(8)
        pdf.cell(80, 5, "Name ______________________________")
        pdf.cell(0, 5, "Signature / date ______________________", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _stamp(iso: str | None):
    from datetime import UTC, datetime

    try:
        value = datetime.fromisoformat(iso) if iso else datetime(2000, 1, 1, tzinfo=UTC)
    except ValueError:
        value = datetime(2000, 1, 1, tzinfo=UTC)
    return value if value.tzinfo else value.replace(tzinfo=UTC)


def pdf_fit(pdf, text: str, width: float) -> str:
    text = str(text)
    while text and pdf.get_string_width(text) > width - 1.5:
        text = text[:-2] + "…" if len(text) > 2 else ""
    return text
