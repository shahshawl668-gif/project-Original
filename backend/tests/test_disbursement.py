"""
Disbursement validation, tested against synthetic data with a known answer.

Every scenario pack is generated with a fixed seed, read through the product's
own readers and mapping profiles, and checked rule by rule against the pack's
answer key: every planted error caught, nothing found on rows nobody planted.
"""
from __future__ import annotations

import io
import time
from datetime import date
from decimal import Decimal
from functools import cache

import pytest

from app.services.disbursement import config, engine, inputs, template, values
from app.services.disbursement.model import BankFile, BankRow, Inputs, RegisterRow, Table
from app.services.studio import mapping
from tools.disbursement_synth import LAYOUTS, PERIOD, generate, run_pack, score

SEED = 20260901
SCENARIOS = ("A", "B", "C", "D")


@cache
def _pack(scenario: str, layout: str, n: int = 500):
    pack = generate(n, SEED, scenario, layout)
    result, parsed, tmpl = run_pack(pack)
    return pack, result, parsed, tmpl


def _found(result) -> set[tuple[str, str]]:
    return {(f.rule_id, f.employee_id or "") for f in result.findings}


def _expected(pack) -> set[tuple[str, str]]:
    return {(e["rule_id"], e["employee_id"]) for e in pack.expected}


# ---------------------------------------------------------------------------
# The answer key
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("layout", list(LAYOUTS))
@pytest.mark.parametrize("scenario", SCENARIOS)
def test_findings_match_the_answer_key_exactly(scenario, layout):
    pack, result, _, _ = _pack(scenario, layout)
    got, want = _found(result), _expected(pack)
    assert want - got == set(), f"planted but not caught: {sorted(want - got)}"
    assert got - want == set(), f"found where nothing was planted: {sorted(got - want)}"
    severities = {(f.rule_id, f.employee_id or ""): f.severity for f in result.findings}
    for e in pack.expected:
        assert severities[(e["rule_id"], e["employee_id"])] == e["severity"], e


@pytest.mark.parametrize("layout", list(LAYOUTS))
@pytest.mark.parametrize("scenario", SCENARIOS)
def test_each_scenario_reaches_its_verdict(scenario, layout):
    pack, result, _, _ = _pack(scenario, layout)
    assert result.verdict == pack.verdict


def test_a_clean_month_produces_no_findings_at_all():
    for layout in LAYOUTS:
        pack, result, _, _ = _pack("A", layout)
        assert result.findings == []
        assert result.totals["held_rows"] == 0
        assert result.totals["release_amount"] == result.totals["file_total"]


def test_every_rule_is_planted_at_least_twice_and_all_are_caught():
    ran = [(_pack(s, lay)[0], _pack(s, lay)[1]) for s in SCENARIOS for lay in LAYOUTS]
    table = score(ran)
    for rule in config.RULE_BY_ID:
        assert table.get(rule, {}).get("planted", 0) >= 2, f"{rule} is planted fewer than twice"
    for rule, row in table.items():
        assert row["caught"] == row["planted"], (rule, row)
        assert row["false_positives"] == 0, (rule, row)


def test_missing_optional_inputs_are_not_run_with_a_reason_and_the_rest_still_runs():
    pack, result, _, _ = _pack("D", "darwinbox_style")
    status = {r.rule_id: r for r in result.rules}
    for rule in ("DSB-06", "DSB-11", "DSB-13", "DSB-04"):
        assert status[rule].status == config.NOT_RUN and status[rule].reason, rule
    not_run = [f for f in result.findings if f.severity == config.NOT_RUN]
    assert {f.rule_id for f in not_run} == {"DSB-04", "DSB-06", "DSB-11", "DSB-13"}
    assert all(f.reason for f in not_run)
    # Status columns in the register keep DSB-09 running without a hold list.
    assert status["DSB-09"].status == config.RAN
    # A format without a header record cannot have its header checked: not applicable, not a pass.
    generic = {r.rule_id: r for r in _pack("D", "generic")[1].rules}
    assert generic["DSB-04"].status == config.NOT_APPLICABLE


def test_the_bridge_walks_from_register_to_file_without_a_gap():
    for scenario in ("A", "B", "D"):
        for layout in LAYOUTS:
            _, result, _, _ = _pack(scenario, layout)
            lines = result.bridge
            walked = sum((b["amount"] for b in lines[:6]), Decimal("0"))
            assert walked == lines[6]["amount"] == result.totals["file_total"]
            assert lines[6]["amount"] + lines[7]["amount"] == lines[8]["amount"] == result.totals["release_amount"]


# ---------------------------------------------------------------------------
# The clean file
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("layout", list(LAYOUTS))
def test_the_clean_file_drops_held_rows_and_recalculates_totals(layout):
    pack, result, parsed, tmpl = _pack("B", layout)
    clean = template.write_clean(parsed, result.kept_lines, tmpl)
    reread = template.read(clean, pack.files["bank_file"][0], tmpl).bank
    held_ids = {f.employee_id for f in result.findings if f.severity == config.HOLD_ROW}
    paid_ids = {r.employee_id for r in reread.rows}
    assert not held_ids & paid_ids
    assert len(reread.rows) == result.totals["release_rows"]
    total = sum(r.amount for r in reread.rows)
    assert total == result.totals["release_amount"]
    if reread.has_header_record:
        assert reread.header["record_count"] == len(reread.rows) == reread.footer["record_count"]
        assert reread.header["total_amount"] == total == reread.footer["total_amount"]
        assert reread.header["debit_account"] == parsed.bank.header["debit_account"]
    # Kept lines are the input's bytes, unchanged.
    original = pack.files["bank_file"][1].decode()
    for line in clean.decode().splitlines(keepends=True):
        if line.startswith(("D|", "E")):
            assert line in original


@pytest.mark.parametrize("layout", list(LAYOUTS))
def test_the_fingerprint_is_identical_across_repeated_runs(layout):
    prints = set()
    for _ in range(3):
        pack = generate(500, SEED, "B", layout)
        result, parsed, tmpl = run_pack(pack)
        prints.add(template.fingerprint(template.write_clean(parsed, result.kept_lines, tmpl)))
    assert len(prints) == 1


def test_a_stopped_file_produces_no_clean_file():
    for layout in LAYOUTS:
        _, result, _, _ = _pack("C", layout)
        assert result.verdict == config.DO_NOT_RELEASE
        assert result.kept_lines == set()


def test_an_excel_payment_file_round_trips_with_fixed_bytes():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["H", "ZZ0000123456", "30092026", 3, 300.0])
    ws.append(["Employee ID", "Beneficiary Name", "Account Number", "IFSC", "Amount"])
    ws.append(["E1", "Asha Rao", "001234567890", "ZZZA0AB1234", 100.0])
    ws.append(["E2", "Ravi Iyer", "2234567890", "ZZZB0CD5678", 120.0])
    ws.append(["E3", "Meera Das", "3234567890", "ZZZC0EF9012", 80.0])
    ws.append(["T", 3, 300.0])
    buf = io.BytesIO()
    wb.save(buf)
    tmpl = {
        "key": "x", "layout": "excel", "column_header": True,
        "columns": [{"field": "employee_id", "header": "Employee ID"},
                    {"field": "beneficiary_name", "header": "Beneficiary Name"},
                    {"field": "account_number", "header": "Account Number"},
                    {"field": "ifsc", "header": "IFSC"}, {"field": "amount", "header": "Amount"}],
        "header_record": {"fields": [{"name": "record_type", "position": 0, "literal": "H"},
                                     {"name": "debit_account", "position": 1},
                                     {"name": "value_date", "position": 2, "format": "%d%m%Y"},
                                     {"name": "record_count", "position": 3},
                                     {"name": "total_amount", "position": 4}]},
        "footer_record": {"fields": [{"name": "record_type", "position": 0, "literal": "T"},
                                     {"name": "record_count", "position": 1},
                                     {"name": "total_amount", "position": 2}]},
    }
    parsed = template.read(buf.getvalue(), "pay.xlsx", tmpl)
    assert parsed.bank.rows[0].account_number == "001234567890"      # leading zeros kept
    assert parsed.bank.header["value_date"] == date(2026, 9, 30)
    keep = {parsed.bank.rows[0].line, parsed.bank.rows[2].line}
    first = template.write_clean(parsed, keep, tmpl)
    second = template.write_clean(template.read(buf.getvalue(), "pay.xlsx", tmpl), keep, tmpl)
    assert first == second
    again = template.read(first, "clean.xlsx", tmpl).bank
    assert [r.employee_id for r in again.rows] == ["E1", "E3"]
    assert again.header["record_count"] == 2 and again.header["total_amount"] == Decimal("180.00")
    assert again.footer["total_amount"] == Decimal("180.00")


def test_a_file_that_does_not_fit_its_template_says_where():
    tmpl = template.builtin_templates()["batch_pipe_hdt"]
    with pytest.raises(template.TemplateError, match="header record"):
        template.read(b"X|bad\nD|E1|A|1234567890|ZZZA0AB1234|10.00|r\nT|1|10.00\n", "f.txt", tmpl)
    with pytest.raises(template.TemplateError, match="not a payment record"):
        template.read(b"H|A|30092026|1|10.00\nQ|E1|A|1234567890|ZZZA0AB1234|10.00|r\nT|1|10.00\n", "f.txt", tmpl)
    with pytest.raises(template.TemplateError, match="Amount"):
        template.read(b"Employee ID,Beneficiary Name\nE1,A\n", "f.csv", template.builtin_templates()["generic_csv"])


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------
def _tiny(**over):
    bank = BankFile(rows=[
        BankRow(1, 2, "E1", "E1", "Asha Rao", "1234567890", "ZZZA0AB1234", Decimal("100.00"), "100.00"),
        BankRow(2, 3, "E2", "E2", "Rahul Verma", "2234567890", "ZZZA0AB1234", Decimal("120.00"), "120.00"),
    ], columns={"employee_id", "beneficiary_name", "account_number", "ifsc", "amount"})
    reg = Table([RegisterRow(2, "E1", "E1", "Asha Rao", Decimal("100.00")),
                 RegisterRow(3, "E2", "E2", "Priya Nair", Decimal("120.00"))],
                {"employee_id", "employee_name", "net_pay"})
    return engine.run(Inputs(bank_file=bank, register=reg), config.build_settings(PERIOD, over))


def test_severities_and_switches_are_per_company():
    assert _tiny().verdict == config.CLEAR                       # DSB-14 is a flag by default
    assert any(f.rule_id == "DSB-14" for f in _tiny().findings)
    # Holding one of two payments is half the file: above the 25% cap, so the file stops…
    capped = _tiny(severities={"DSB-14": "HOLD_ROW"})
    assert capped.verdict == config.DO_NOT_RELEASE
    assert {f.field for f in capped.findings if f.rule_id in ("DSB-01", "DSB-02")} == {
        "held_amount_share", "held_row_share"}
    # …and with the cap switched off, that employee alone is held.
    held = _tiny(severities={"DSB-14": "HOLD_ROW"}, max_hold_share_pct="0")
    assert held.verdict == config.WITH_HOLDS and held.totals["held_employees"] == 1
    off = _tiny(enabled={"DSB-14": False})
    assert not any(f.rule_id == "DSB-14" for f in off.findings)
    assert {r.rule_id: r.status for r in off.rules}["DSB-14"] == config.DISABLED
    assert {r.rule_id: r.status for r in off.rules}["DSB-16"] == config.NOT_RUN   # no master: said, not skipped


def test_bad_settings_are_refused_not_defaulted():
    with pytest.raises(ValueError, match="DSB-15 can be"):
        config.build_settings(PERIOD, {"severities": {"DSB-15": "HOLD_ROW"}})
    with pytest.raises(ValueError, match="Unknown setting"):
        config.build_settings(PERIOD, {"varience_pct": 10})
    with pytest.raises(ValueError, match="name_similarity"):
        config.build_settings(PERIOD, {"name_similarity": "1.5"})
    with pytest.raises(ValueError, match="period"):
        config.build_settings("Sept 2026", {})


def test_the_variance_threshold_is_a_setting():
    pack, _, _, _ = _pack("B", "generic")
    lenient = dict(pack.settings, variance_pct="100")
    loaded, _ = inputs.load(pack.files, inputs.builtin_profiles()["generic"],
                            template.builtin_templates()["generic_csv"])
    result = engine.run(loaded, config.build_settings(PERIOD, lenient, pack.value_date))
    assert not any(f.rule_id == "DSB-13" for f in result.findings)


def test_a_register_and_bank_file_are_required():
    with pytest.raises(ValueError, match="register"):
        engine.run(Inputs(bank_file=BankFile(rows=[], columns=set())), config.build_settings(PERIOD))


# ---------------------------------------------------------------------------
# Real-export mess
# ---------------------------------------------------------------------------
def test_untidy_values_are_read_for_what_they_mean():
    assert values.ifsc(" zzza0ab1234 ") == "ZZZA0AB1234"
    assert values.ifsc_problem("ZZZA0AB1234") is None
    assert "11 characters" in values.ifsc_problem("ZZZA0AB123")
    assert "4 letters, then 0" in values.ifsc_problem("ZZZA1AB1234")
    assert "scientific" in values.account_problem("1.23457E+17", 9, 18)
    assert "other than digits" in values.account_problem("1234-567890", 9, 18)
    assert values.account_problem("001234567", 9, 18) is None
    assert values.same_account("0012345678", "12345678") == (True, "the master lost leading zeros that the file keeps")
    ok, why = values.same_account("12345678", "0012345678")
    assert ok is False and "file lost leading zeros" in why
    assert values.name_similarity("Mr. A. K. Sharma", "Anil Kumar Sharma", frozenset({"mr"})) == 1.0


def test_a_spreadsheet_keeps_account_numbers_as_text():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Employee ID", "Account Number", "IFSC"])
    ws.append(["E1", "000123456789", "ZZZA0AB1234"])
    ws.append(["E2", 123456789012345, "ZZZA0AB1234"])
    buf = io.BytesIO()
    wb.save(buf)
    records, _ = inputs.read_table(buf.getvalue(), "m.xlsx")
    assert records[0]["Account Number"] == "000123456789"
    assert records[1]["Account Number"] == "123456789012345"


def test_profiles_are_ordinary_studio_mappings_and_a_new_layout_needs_no_code():
    for profile in inputs.builtin_profiles().values():
        for slot, spec in profile["inputs"].items():
            mapping.check_spec(spec, inputs.OBJECT_FOR_SLOT[slot])
    custom = {"key": "acme", "name": "Acme HRMS", "inputs": {"register": {"fields": [
        {"target": "employee_id", "source": "Staff No", "type": "id", "required": True},
        {"target": "employee_name", "source": "Full Name"},
        {"target": "net_pay", "source": "Take Home"},
    ]}}}
    inputs.check_profile(custom)
    notes: list = []
    table = inputs.read_slot("register", b"Staff No,Full Name,Take Home\nA7 ,Asha Rao,\"1,23,456.50\"\n",
                             "r.csv", custom["inputs"]["register"], notes)
    row = table.rows[0]
    assert (row.employee_id, row.net_pay) == ("A7", Decimal("123456.50"))
    assert table.columns == {"employee_id", "employee_name", "net_pay"}


def test_an_unreadable_value_is_a_note_not_a_silent_zero():
    notes: list = []
    spec = inputs.builtin_profiles()["generic"]["inputs"]["register"]
    table = inputs.read_slot("register", b"Employee ID,Employee Name,Net Pay\nE1,Asha,abc\n", "r.csv", spec, notes)
    assert table.rows[0].net_pay is None and table.rows[0].net_raw == "abc"
    assert any("could not be read as a number" in n.message for n in notes)


# ---------------------------------------------------------------------------
# Performance
# ---------------------------------------------------------------------------
def test_ten_thousand_employees_are_validated_in_under_ten_seconds():
    pack = generate(10000, SEED, "B", "generic")
    started = time.perf_counter()
    result, parsed, tmpl = run_pack(pack)
    clean = template.write_clean(parsed, result.kept_lines, tmpl)
    elapsed = time.perf_counter() - started
    assert result.verdict == config.WITH_HOLDS and clean
    assert _expected(pack) == _found(result)
    assert elapsed < 10, f"{elapsed:.2f}s"


# ---------------------------------------------------------------------------
# What the bank should pay: total salary, not net pay
# ---------------------------------------------------------------------------
def _pay(register_rows, columns, bank_amounts, **settings):
    bank = BankFile(rows=[
        BankRow(i, i + 1, emp, emp, None, f"{i}234567890", "ZZZA0AB1234", Decimal(amount), amount)
        for i, (emp, amount) in enumerate(bank_amounts, start=1)
    ], columns={"employee_id", "account_number", "ifsc", "amount"})
    reg = Table(register_rows, {"employee_id", *columns})
    return engine.run(Inputs(bank_file=bank, register=reg),
                      config.build_settings(PERIOD, {"max_hold_share_pct": "0", **settings}))


def _row(emp, net, **kw):
    return RegisterRow(2, emp, emp, None, Decimal(net), **{k: Decimal(v) for k, v in kw.items()})


def test_the_bank_amount_is_checked_against_total_salary_not_net_pay():
    # Net 40,000 + reimbursement 2,500 − 10,000 held = 32,500 payable.
    rows = [_row("E1", "40000.00", reimbursement="2500.00", salary_hold="10000.00", total_payable="32500.00")]
    cols = {"net_pay", "reimbursement", "salary_hold", "total_payable"}
    assert _pay(rows, cols, [("E1", "32500.00")]).verdict == config.CLEAR
    paid_net = _pay(rows, cols, [("E1", "40000.00")])
    finding = next(f for f in paid_net.findings if f.rule_id == "DSB-08")
    assert finding.expected == "32,500.00" and finding.actual == "40,000.00"
    assert "total salary payable" in finding.reason and "reimbursements 2,500.00" in finding.reason
    assert "held 10,000.00" in finding.reason


def test_without_a_total_column_the_amount_is_built_from_its_parts_and_said_so():
    rows = [_row("E1", "40000.00", reimbursement="2500.00", hold_release="5000.00")]
    result = _pay(rows, {"net_pay", "reimbursement", "hold_release"}, [("E1", "47500.00")])
    assert result.verdict == config.CLEAR
    rules = {r.rule_id: r for r in result.rules}
    assert "no total salary column" in rules["DSB-08"].reason
    assert rules["DSB-17"].status == config.NOT_RUN           # nothing stated to check the parts against
    assert result.totals["amount_basis"].startswith("net pay plus reimbursements")


def test_with_net_pay_only_the_comparison_says_what_it_cannot_allow_for():
    result = _pay([_row("E1", "40000.00")], {"net_pay"}, [("E1", "40000.00")])
    rules = {r.rule_id: r for r in result.rules}
    assert "cannot be allowed for" in rules["DSB-08"].reason
    assert rules["DSB-17"].status == config.NOT_RUN


def test_a_total_that_does_not_add_up_is_flagged():
    rows = [_row("E1", "40000.00", reimbursement="2500.00", total_payable="44500.00")]
    result = _pay(rows, {"net_pay", "reimbursement", "total_payable"}, [("E1", "44500.00")])
    finding = next(f for f in result.findings if f.rule_id == "DSB-17")
    assert finding.severity == config.FLAG and finding.expected == "42,500.00" and finding.actual == "44,500.00"
    assert result.verdict == config.CLEAR                      # a flag: paid, and shown to the approver
    rows.append(_row("E2", "30000.00", total_payable="30000.00"))
    held = _pay(rows, {"net_pay", "reimbursement", "total_payable"}, [("E1", "44500.00"), ("E2", "30000.00")],
                severities={"DSB-17": "HOLD_ROW"})
    assert held.verdict == config.WITH_HOLDS and held.totals["held_employees"] == 1


def test_salary_held_in_full_is_not_expected_in_the_file():
    rows = [_row("E1", "40000.00", salary_hold="40000.00", total_payable="0.00"),
            _row("E2", "30000.00", total_payable="30000.00")]
    result = _pay(rows, {"net_pay", "salary_hold", "total_payable"}, [("E2", "30000.00")])
    assert not any(f.rule_id == "DSB-15" for f in result.findings)
    assert result.verdict == config.CLEAR


# ---------------------------------------------------------------------------
# What people receive: the exception report and the approver summary
# ---------------------------------------------------------------------------
META = {"period": PERIOD, "entity_name": "Synthetic Co", "run_by": "checker@example.test",
        "generated_at": "2026-10-03T10:00:00+00:00"}


@cache
def _checked(scenario: str, layout: str, **settings_over):
    from app.services.disbursement.check import run_check

    pack = generate(500, SEED, scenario, layout)
    settings = config.build_settings(PERIOD, {**pack.settings, **dict(settings_over)}, pack.value_date)
    return pack, run_check(pack.files, inputs.builtin_profiles()[pack.profile],
                           template.builtin_templates()[pack.template], settings, META)


def _csv_rows(content: bytes) -> list[dict[str, str]]:
    import csv

    assert content.startswith(b"\xef\xbb\xbf")
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))


@pytest.mark.parametrize("layout", list(LAYOUTS))
def test_the_exception_report_lists_every_finding_including_checks_that_did_not_run(layout):
    from app.services.disbursement import outputs

    pack, checked = _checked("D", layout)
    rows = _csv_rows(outputs.exception_csv(checked.report))
    assert len(rows) == len(checked.result.findings)
    assert {(r["Check"], r["Employee ID"]) for r in rows} == _found(checked.result)
    not_run = {r["Check"] for r in rows if r["Severity"] == config.NOT_RUN}
    assert not_run == {r.rule_id for r in checked.result.rules if r.status == config.NOT_RUN} != set()
    assert all(r["Reason"] for r in rows if r["Severity"] == config.NOT_RUN)


def test_a_switched_off_check_is_reported_as_switched_off_not_left_out():
    from app.services.disbursement import outputs

    _, checked = _checked("A", "generic", enabled=(("DSB-14", False),))
    rows = _csv_rows(outputs.exception_csv(checked.report))
    assert [(r["Check"], r["Severity"]) for r in rows] == [("DSB-14", config.DISABLED)]


def test_exports_do_not_run_text_from_the_files_as_formulas_and_keep_accounts_in_full():
    from openpyxl import load_workbook

    from app.services.disbursement import outputs

    _, checked = _checked("B", "generic")
    report = {**checked.report, "findings": [
        *checked.report["findings"],
        {"rule_id": "DSB-14", "severity": config.FLAG, "employee_id": "E1",
         "employee_name": '=HYPERLINK("http://x","y")', "field": "beneficiary_name",
         "expected": "+91 Rao", "actual": "-Rao", "reason": "@cmd", "rows": [4], "amount": "-12.50"},
    ]}
    csv_rows = _csv_rows(outputs.exception_csv(report))
    last = csv_rows[-1]
    assert last["Employee name"].startswith("'=") and last["Expected"] == "'+91 Rao"
    assert last["Actual"] == "'-Rao" and last["Reason"] == "'@cmd"
    assert last["Amount in file"] == "-12.50"             # a number, not escaped text
    book = load_workbook(io.BytesIO(outputs.exception_xlsx(report)))
    assert book.sheetnames == ["Summary", "Findings", "Checks", "Bridge", "Inputs", "Notes"]
    cells = [c for row in book["Findings"].iter_rows() for c in row]
    assert not any(c.data_type == "f" for c in cells)
    # Bank accounts are shown in full: the approver is being asked whether they are right.
    accounts = [f for f in checked.report["findings"] if f["rule_id"] == "DSB-16" and f["actual"].isdigit()]
    assert accounts and any(accounts[0]["actual"] == r["Actual"] for r in csv_rows)


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_the_approver_summary_is_a_pdf_and_the_same_report_gives_the_same_bytes(scenario):
    from app.services.disbursement import outputs

    _, checked = _checked(scenario, "darwinbox_style")
    first = outputs.approver_pdf(checked.report)
    assert first.startswith(b"%PDF") and first == outputs.approver_pdf(checked.report)
    approved = outputs.approver_pdf(checked.report, {"approver": "Approver One", "approved_at": "2026-10-03T11:00",
                                                     "fingerprint": checked.report["clean_sha256"] or "",
                                                     "acknowledged": ["DSB-13"]})
    assert approved != first


def test_the_report_carries_the_clean_file_fingerprint_and_every_input_fingerprint():
    _, checked = _checked("B", "generic")
    assert checked.report["clean_sha256"] == template.fingerprint(checked.clean)
    assert checked.report["clean_filename"] == "bank_file_clean.csv"
    slots = [i["slot"] for i in checked.report["inputs"]]
    assert slots[:2] == ["bank_file", "register"] and all(len(i["sha256"]) == 64 for i in checked.report["inputs"])
    assert checked.report["settings"]["variance_pct"] == "25"
    _, stopped = _checked("C", "generic")
    assert stopped.clean is None and stopped.report["clean_sha256"] is None


def test_held_employees_are_grouped_with_every_reason():
    from app.services.disbursement import outputs

    _, checked = _checked("B", "generic")
    held = outputs.held_employees(checked.report)
    assert len(held) == checked.report["totals"]["held_employees"]
    assert sum(h["amount"] for h in held) == Decimal(checked.report["totals"]["held_amount"])
    assert all(h["reasons"] for h in held)
