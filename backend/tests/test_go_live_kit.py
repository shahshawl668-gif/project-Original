"""
The go-live kit: the D6 verification worksheet and the parallel-run pack.

Both are workbooks a person fills in and signs, so the tests check what the
person sees: every finding present with an empty verdict cell restricted to
the four verdicts, exit criteria that count those verdicts, nothing pre-filled,
and register text that cannot become a formula.
"""
from __future__ import annotations

import io
import json
import uuid

import pytest
from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import AuditEvent, FindingRecord
from tests.test_run_history import _company, _data, _drain

HEADER = "employee_id,employee_name,state,paid_days,lop_days,basic,hra,special allowance,pf_employee,pf_employer,pan"
ROWS = [
    "K1,Person K1,Karnataka,30,0,30000,12000,8000,1800,1800,ABCPK1234F",
    # PF ₹100 short, and a name that would be a live formula if written as one.
    'K2,"=HYPERLINK(""https://example.invalid"",""x"")",Karnataka,30,0,30000,12000,8000,1700,1700,ABCPK1235F',
]


@pytest.fixture()
def company(client):
    return _company(client, "golive")


def _run(client, headers) -> str:
    content = ("\n".join([HEADER, *ROWS]) + "\n").encode()
    r = client.post("/api/payroll/upload", headers=headers,
                    files={"file": ("r.csv", io.BytesIO(content), "text/csv")},
                    data={"meta": json.dumps({"period_month": "2026-09-01", "strict_header_check": False,
                                              "return_employees": False})})
    assert r.status_code == 200, r.text
    job = _data(client.post("/api/validation/jobs", headers=headers, json={"period_month": "2026-09-01"}))["job"]
    _drain()
    return _data(client.get(f"/api/validation/jobs/{job['id']}", headers=headers))["run_id"]


def test_parallel_run_pack_lists_every_finding_for_a_verdict(client, company):
    run_id = _run(client, company)
    r = client.get(f"/api/validation/runs/{run_id}/parallel-run.xlsx", headers=company)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["How to use", "Run", "Findings", "Could not check", "Exit criteria", "Sign-off"]

    db = SessionLocal()
    try:
        expected = db.query(FindingRecord).filter(FindingRecord.run_id == uuid.UUID(run_id),
                                                  FindingRecord.status == "FAIL").count()
        assert db.query(AuditEvent).filter(AuditEvent.action == "export.downloaded",
                                           AuditEvent.object_id == run_id).count() == 1
    finally:
        db.close()
    ws = wb["Findings"]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert expected > 0 and len(rows) == expected
    assert any(row[2] == "STAT-001" and row[0] == "K2" for row in rows)
    # Nothing pre-filled: every verdict is the reviewer's.
    assert all(row[10] is None for row in rows)
    rule = ws.data_validations.dataValidation[0]
    assert "False positive" in rule.formula1 and str(rule.sqref) == f"K2:K{expected + 1}"

    # The register's text stays text.
    name_cells = [c for c in ws["B"] if c.value and str(c.value).startswith("=HYPERLINK")]
    assert name_cells and all(c.data_type == "s" for c in name_cells)

    criteria = wb["Exit criteria"]
    formulas = [c.value for row in criteria.iter_rows() for c in row if isinstance(c.value, str)]
    assert any('COUNTIF(Findings!$K$2' in f and '"False positive"' in f for f in formulas)
    assert any(f.startswith("=IF(COUNTIF(C") and '"PASSED"' in f for f in formulas)


def test_parallel_run_pack_is_404_for_another_company(client, company):
    run_id = _run(client, company)
    stranger = _company(client, "golive-other")
    assert client.get(f"/api/validation/runs/{run_id}/parallel-run.xlsx", headers=stranger).status_code == 404


def test_verification_worksheet_downloads_with_every_tab(client, company):
    r = client.get("/api/config/statutory/verification-worksheet.xlsx", headers=company)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert {"How to use", "Income tax", "PF and ESIC", "Labour Codes", "PT by state", "LWF by state",
            "Sign-off"} <= set(wb.sheetnames)
    # Reviewer columns are empty: a pre-filled "Y" is not a review.
    ws = wb["PF and ESIC"]
    header = [c.value for c in ws[4]]
    col = header.index("Correct?") + 1
    assert all(ws.cell(row=r, column=col).value is None for r in range(5, ws.max_row + 1))
