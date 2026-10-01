"""
Exports: data from an upload never runs in the spreadsheet it is exported to,
and every download of payroll data is in the client's audit trail.
"""
from __future__ import annotations

import io
import uuid
import zipfile

import openpyxl

from app.database import SessionLocal
from app.models import AuditEvent
from app.services import export_safety
from tests.test_coverage import BASE_ROWS, _register, _run
from tests.test_run_history import PERIOD, _company

HOSTILE = "=2+5"   # stands for =HYPERLINK(...) or a DDE payload; no comma so the CSV stays one cell


def _formulas(xlsx: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        return sum(z.read(n).count(b"<f>") + z.read(n).count(b"<f ") for n in z.namelist() if n.startswith("xl/worksheets/"))


def _cells(xlsx: bytes) -> list:
    book = openpyxl.load_workbook(io.BytesIO(xlsx))
    return [c for ws in book.worksheets for row in ws.iter_rows() for c in row if c.value == HOSTILE]


def _exports(entity_id: str) -> list[str]:
    with SessionLocal() as db:
        return [e.object_type for e in db.query(AuditEvent).filter(
            AuditEvent.entity_id == uuid.UUID(entity_id), AuditEvent.action == "export.downloaded")]


def test_a_formula_in_an_employee_name_is_exported_as_text(client):
    company = _company(client, "xl")
    rows = [BASE_ROWS[0], BASE_ROWS[1].replace("Person E002", HOSTILE), BASE_ROWS[2]]
    run_id = _run(client, company, _register(rows))

    run_xlsx = client.get(f"/api/validation/runs/{run_id}/export.xlsx", headers=company)
    assert run_xlsx.status_code == 200, run_xlsx.text
    pack = client.get(f"/api/signoff/{PERIOD}/evidence-pack", headers=company)
    assert pack.status_code == 200, pack.text
    for payload in (run_xlsx.content, pack.content):
        found = _cells(payload)
        assert found, "the name should still be there, exactly as uploaded"
        assert all(c.data_type == "s" for c in found)
        assert _formulas(payload) == 0

    assert {"validation_run", "evidence_pack"} <= set(_exports(company["X-Entity-Id"]))


def test_the_guard_leaves_ordinary_values_alone():
    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["E1", "-", "+91 98450 00000", "@home", 1200.5, HOSTILE])
    assert export_safety.neutralise_workbook(book) == 1
    buf = io.BytesIO()
    book.save(buf)
    again = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active
    assert [c.value for c in again[1]] == ["E1", "-", "+91 98450 00000", "@home", 1200.5, HOSTILE]
    assert _formulas(buf.getvalue()) == 0
