"""
Report downloads and the audit trail.

A report that circulates by email without saying what it covered becomes an
orphan figure within a week, so every workbook opens with a provenance sheet and
this suite checks it is there and correct. And an audit trail that can be edited
is not one, so this checks the writes happen and that there is no way to unmake
them.
"""
from __future__ import annotations

import io
import uuid
import zipfile
from datetime import date
from decimal import Decimal

import pytest

from app.database import SessionLocal
from app.models import AuditEvent, Entity, SalaryRegister, SalaryRegisterRow, User
from app.services import reporting
from app.services.dimensions import UNASSIGNED

PASSWORD = "Passw0rd!x"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def workspace(client, request):
    email = f"{request.node.name[:40]}@report-example.com".replace("_", "-")
    r = client.post("/api/auth/signup",
                    json={"email": email, "password": PASSWORD, "company_name": "Report Tests"})
    if r.status_code != 200:
        r = client.post("/api/auth/login", json={"email": email, "password": PASSWORD})
    headers = {"Authorization": f"Bearer {r.json()['data']['access_token']}"}
    entity_id = client.post("/api/org/entities", json={"name": f"E-{request.node.name[:26]}"},
                            headers=headers).json()["data"]["id"]
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        entity = db.get(Entity, uuid.UUID(entity_id))
        db.expunge_all()
    finally:
        db.close()
    return entity, user, {**headers, "X-Entity-Id": entity_id}


def _dims(**kw) -> dict:
    base = {k: UNASSIGNED for k in
            ("business_unit", "department", "cost_center", "work_location",
             "work_state", "grade", "designation", "employment_type", "skill_category")}
    base.update(kw)
    return base


def _register(entity, user, period: date, rows: list[dict]) -> None:
    db = SessionLocal()
    try:
        reg = SalaryRegister(user_id=user.id, entity_id=entity.id, period_month=period,
                             filename="register.csv", employee_count=len(rows))
        db.add(reg); db.flush()
        for r in rows:
            db.add(SalaryRegisterRow(
                register_id=reg.id, user_id=user.id, entity_id=entity.id,
                period_month=period, employee_id=r["employee_id"],
                employee_name=r.get("name"),
                components=r.get("components", {"basic": 30000.0}),
                arrears={}, deductions={}, dimensions=r.get("dimensions") or _dims(),
                increment_arrear_total=Decimal("0"),
            ))
        db.commit()
    finally:
        db.close()


def _sheets(payload: bytes) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(payload))
    return wb.sheetnames


def _cells(payload: bytes, sheet: str) -> list[list]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(payload))
    return [list(row) for row in wb[sheet].iter_rows(values_only=True)]


def _seed(entity, user) -> None:
    for month in (4, 5):
        _register(entity, user, date(2026, month, 1), [
            {"employee_id": "E1", "name": "Priya Sharma",
             "dimensions": _dims(department="Engineering")},
            {"employee_id": "E2", "name": "Rahul Verma",
             "dimensions": _dims(department="Sales")},
        ])


# ── the catalogue ───────────────────────────────────────────────────────────

def test_every_report_in_the_catalogue_can_be_generated(client, workspace):
    """A listed report that 500s is worse than one that is not listed."""
    entity, user, headers = workspace
    _seed(entity, user)

    listed = client.get("/api/reports", headers=headers).json()["data"]["reports"]
    # Restricted reports are absent until the entity authorises them, so the
    # catalogue is the unrestricted set rather than everything that exists.
    assert {r["key"] for r in listed} == set(reporting.REPORTS) - reporting.RESTRICTED_REPORTS

    for report in listed:
        r = client.get(f"/api/reports/{report['key']}.xlsx", headers=headers)
        assert r.status_code == 200, f"{report['key']}: {r.text[:200]}"
        assert r.headers["content-type"].startswith(XLSX)
        assert zipfile.is_zipfile(io.BytesIO(r.content)), report["key"]


def test_an_unknown_report_is_a_404(client, workspace):
    _, _, headers = workspace
    assert client.get("/api/reports/vibes.xlsx", headers=headers).status_code == 404


def test_a_restricted_report_is_not_listed_until_it_is_authorised(client, workspace):
    """Offering a card that always 403s teaches people to ignore errors."""
    entity, user, headers = workspace
    _seed(entity, user)

    keys = {r["key"] for r in client.get("/api/reports", headers=headers).json()["data"]["reports"]}
    assert "pay-equity" not in keys
    assert client.get("/api/reports/pay-equity.xlsx", headers=headers).status_code == 403

    client.put("/api/bi/pay-equity/settings", json={"enabled": True}, headers=headers)

    keys = {r["key"] for r in client.get("/api/reports", headers=headers).json()["data"]["reports"]}
    assert "pay-equity" in keys
    assert client.get("/api/reports/pay-equity.xlsx", headers=headers).status_code == 200


# ── provenance ──────────────────────────────────────────────────────────────

def test_every_workbook_opens_with_where_it_came_from(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)

    payload = client.get("/api/reports/management-summary.xlsx", headers=headers).content
    assert _sheets(payload)[0] == "About this report"

    fields = {row[0]: row[1] for row in _cells(payload, "About this report") if row and row[0]}
    assert fields["Report"] == "Management summary"
    assert fields["Generated by"] == user.email
    assert fields["Latest register"] == "May 2026"
    assert fields["Registers stored"] == 2
    assert fields["Generated at"]


def test_the_filters_applied_are_recorded_on_the_report(client, workspace):
    """A department report that does not say it is one gets read as the company."""
    entity, user, headers = workspace
    _seed(entity, user)

    payload = client.get(
        "/api/reports/department-cost.xlsx?department=Engineering", headers=headers
    ).content
    fields = {row[0]: row[1] for row in _cells(payload, "About this report") if row and row[0]}
    assert "department=Engineering" in fields["Filters applied"]


def test_an_unfiltered_report_says_none_rather_than_leaving_it_blank(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)
    payload = client.get("/api/reports/management-summary.xlsx", headers=headers).content
    fields = {row[0]: row[1] for row in _cells(payload, "About this report") if row and row[0]}
    assert fields["Filters applied"] == "none"


# ── the figures match the dashboard ─────────────────────────────────────────

def test_a_report_agrees_with_the_dashboard_it_was_downloaded_from(client, workspace):
    """A report that disagrees with the screen destroys confidence in both."""
    entity, user, headers = workspace
    _seed(entity, user)

    screen = client.get("/api/bi/cost-analysis?group_by=department",
                        headers=headers).json()["data"]["totals"]
    payload = client.get("/api/reports/management-summary.xlsx", headers=headers).content
    sheet = {row[0]: row[1] for row in _cells(payload, "Summary") if row and row[0]}

    assert sheet["Total CTC"] == screen["ctc"]
    assert sheet["Gross pay"] == screen["gross"]
    assert sheet["Employer contributions"] == screen["employer_cost"]
    assert sheet["Headcount (distinct employees)"] == screen["headcount"]


def test_the_employee_report_has_one_row_per_employee_per_month(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)

    payload = client.get("/api/reports/employee-cost.xlsx", headers=headers).content
    rows = _cells(payload, "Employee cost")
    assert len(rows) == 1 + 4      # header, two employees over two months


# ── masking reaches the exports ─────────────────────────────────────────────

def test_a_masked_session_exports_masked_names(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)

    payload = client.get(
        "/api/reports/employee-cost.xlsx", headers={**headers, "X-Mask-Identity": "on"}
    ).content
    rows = _cells(payload, "Employee cost")
    names = {row[2] for row in rows[1:]}
    ids = {row[1] for row in rows[1:]}
    assert "Priya Sharma" not in names
    assert names <= {"P. S.", "R. V."}
    assert all(str(i).startswith("EMP-") for i in ids)

    fields = {row[0]: row[1] for row in _cells(payload, "About this report") if row and row[0]}
    assert fields["Employee identity"] == "masked"


# ── the audit trail ─────────────────────────────────────────────────────────

def test_downloading_a_report_is_recorded(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)
    client.get("/api/reports/management-summary.xlsx", headers=headers)

    events = client.get("/api/audit", headers=headers).json()["data"]["events"]
    downloads = [e for e in events if e["action"] == "report.downloaded"]
    assert downloads and downloads[0]["object_id"] == "management-summary"
    assert downloads[0]["user_email"] == user.email


def test_the_trail_can_be_filtered_by_action(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)
    client.get("/api/reports/headcount.xlsx", headers=headers)

    events = client.get("/api/audit?action=report.downloaded",
                        headers=headers).json()["data"]["events"]
    assert events and {e["action"] for e in events} == {"report.downloaded"}


def test_the_trail_is_newest_first(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)
    client.get("/api/reports/headcount.xlsx", headers=headers)
    client.get("/api/reports/compensation.xlsx", headers=headers)

    events = client.get("/api/audit?action=report.downloaded",
                        headers=headers).json()["data"]["events"]
    assert events[0]["object_id"] == "compensation"


def test_there_is_no_endpoint_that_edits_or_removes_an_event(client, workspace):
    """A trail that can be edited is not one."""
    from app.main import app

    audit_paths = {
        path: set(operations)
        for path, operations in app.openapi()["paths"].items()
        if path.startswith("/api/audit")
    }
    assert audit_paths
    assert all(methods == {"get"} for methods in audit_paths.values()), audit_paths


def test_an_audit_entry_and_the_thing_it_describes_share_a_transaction(workspace):
    """A trail that survives a rolled-back change records events that never happened."""
    from app.services import audit

    entity, user, _ = workspace
    db = SessionLocal()
    try:
        audit.record(db, entity_id=entity.id, user=user, action="test.rollback",
                     object_type="test", summary="should not survive")
        db.rollback()
    finally:
        db.close()

    db = SessionLocal()
    try:
        assert db.query(AuditEvent).filter(AuditEvent.action == "test.rollback").count() == 0
    finally:
        db.close()


def test_the_trail_does_not_leak_between_entities(client, workspace):
    entity, user, headers = workspace
    _seed(entity, user)
    client.get("/api/reports/headcount.xlsx", headers=headers)
    other = client.post("/api/org/entities", json={"name": "Other"},
                        headers=headers).json()["data"]["id"]

    events = client.get("/api/audit", headers={**headers, "X-Entity-Id": other})
    assert events.json()["data"]["events"] == []
