"""The configuration file is a spreadsheet, and may replace one entity's setup, never another's."""
import io

from openpyxl import Workbook, load_workbook

PASSWORD = "Passw0rd!x"
COMPONENT_HEADER = ["component_name", "pf_applicable", "esic_applicable", "pt_applicable", "lwf_applicable",
                    "bonus_applicable", "included_in_wages", "taxable", "tax_exemption_type"]


def _workspace(client, email):
    result = client.post("/api/auth/signup", json={"email": email, "password": PASSWORD, "company_name": "Bundle test"})
    assert result.status_code == 200, result.text
    headers = {"Authorization": f"Bearer {result.json()['data']['access_token']}"}
    entity = client.post("/api/org/entities", headers=headers, json={"name": email}).json()["data"]
    headers["X-Entity-Id"] = entity["id"]
    return headers


def _send(client, headers, name, content, dry_run=True):
    return client.post(f"/api/config/bundle/import?dry_run={str(dry_run).lower()}", headers=headers,
                       files={"file": (name, io.BytesIO(content), "application/octet-stream")})


def _csv(*lines):
    return ("section,record,field,value\r\n" + "\r\n".join(lines) + "\r\n").encode()


BASIC = [f"components,1,{f},{v}" for f, v in zip(
    COMPONENT_HEADER, ["Basic", "Yes", "Yes", "Yes", "No", "No", "Yes", "Yes", "none"], strict=True)]


def test_csv_preview_apply_and_entity_isolation(client):
    first = _workspace(client, "config-bundle-first@example.com")
    second = _workspace(client, "config-bundle-second@example.com")
    preview = _send(client, first, "configuration.csv", _csv(*BASIC))
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["replace_sections"] == {"components": 1}
    assert client.get("/api/components", headers=first).json()["data"] == []
    assert _send(client, first, "configuration.csv", _csv(*BASIC), dry_run=False).status_code == 200
    rows = client.get("/api/components", headers=first).json()["data"]
    assert rows[0]["component_name"] == "Basic" and rows[0]["pf_applicable"] is True
    assert client.get("/api/components", headers=second).json()["data"] == []
    # A file cannot name an id or another company's record: unknown columns are refused.
    refused = _send(client, first, "configuration.csv", _csv(*BASIC, "components,1,entity_id," + second["X-Entity-Id"]))
    assert refused.status_code == 422


def test_the_download_has_no_code_in_it(client):
    headers = _workspace(client, "config-bundle-nojson@example.com")
    text = client.get("/api/config/bundle/export.csv", headers=headers).text
    assert text.lstrip("﻿").startswith("section,record,field,value")
    assert "{" not in text and "[" not in text
    assert client.get("/api/config/bundle/export", headers=headers).status_code in (404, 405)


def test_excel_download_and_upload_round_trip(client):
    headers = _workspace(client, "config-bundle-excel@example.com")
    assert _send(client, headers, "c.csv", _csv(*BASIC), dry_run=False).status_code == 200
    download = client.get("/api/config/bundle/export.xlsx", headers=headers)
    assert download.status_code == 200
    wb = load_workbook(io.BytesIO(download.content))
    assert "Read me" in wb.sheetnames and "components" in wb.sheetnames and "statutory" in wb.sheetnames
    assert [c.value for c in wb["components"][1]] == COMPONENT_HEADER
    assert wb["components"]["A2"].value == "Basic"

    edited = Workbook()
    edited.active.title = "components"
    edited["components"].append(COMPONENT_HEADER)
    edited["components"].append(["Basic", "Yes", "Yes", "Yes", "No", "No", "Yes", "Yes", "none"])
    edited["components"].append(["HRA", "No", "Yes", "Yes", "No", "No", "Yes", "Yes", "none"])
    buf = io.BytesIO()
    edited.save(buf)
    applied = _send(client, headers, "configuration.xlsx", buf.getvalue(), dry_run=False)
    assert applied.status_code == 200, applied.text
    names = [c["component_name"] for c in client.get("/api/components", headers=headers).json()["data"]]
    assert sorted(names) == ["Basic", "HRA"]


def test_statutory_row_sets_the_engine_and_keeps_the_rest(client):
    headers = _workspace(client, "config-bundle-statutory@example.com")
    before = client.get("/api/config/statutory", headers=headers).json()["data"]
    lines = [
        "statutory,1,pf_employee_rate,0.12", "statutory,1,pf_wage_ceiling,25000",
        "statutory,1,pf_restrict_to_ceiling,Yes", "statutory,1,pf_employee_formula,\"MIN(pf_wage, 25000) * 12%\"",
        "statutory,1,esic_wage_ceiling,21000", "statutory,1,esic_rounding,up",
        "statutory,1,pt_states,\"Karnataka, Maharashtra\"",
    ]
    assert _send(client, headers, "c.csv", _csv(*lines), dry_run=False).status_code == 200
    after = client.get("/api/config/statutory", headers=headers).json()["data"]
    assert after["pf"]["wage"]["wage_ceiling"] in ("25000", "25000.00", 25000)
    assert after["pf"]["employee_formula"] == "MIN(pf_wage, 25000) * 12%"
    # What the file does not carry is untouched.
    assert after["pf"]["rates"]["eps_rate"] == before["pf"]["rates"]["eps_rate"]
    assert after["esic"]["rates"]["employer_rate"] == before["esic"]["rates"]["employer_rate"]
    csv_text = client.get("/api/config/bundle/export.csv", headers=headers).text
    assert "pt_states,\"Karnataka, Maharashtra\"" in csv_text


def test_mappings_and_slabs_round_trip_as_plain_rows(client):
    headers = _workspace(client, "config-bundle-mappings@example.com")
    lines = [
        "import_mappings,1,profile,Client export", "import_mappings,1,file_column,Emp Code",
        "import_mappings,1,product_field,employee_id",
        "import_mappings,2,profile,Client export", "import_mappings,2,file_column,Base Pay",
        "import_mappings,2,product_field,basic",
        "pt_lwf_slabs,1,state,Karnataka", "pt_lwf_slabs,1,rule_type,PT", "pt_lwf_slabs,1,min_salary,25000",
        "pt_lwf_slabs,1,max_salary,99999999", "pt_lwf_slabs,1,deduction_amount,200",
        "pt_lwf_slabs,1,frequency,monthly", "pt_lwf_slabs,1,gender,ALL", "pt_lwf_slabs,1,applicable_months,\"2, 3\"",
        "pt_lwf_slabs,1,sort_order,1",
    ]
    assert _send(client, headers, "c.csv", _csv(*lines), dry_run=False).status_code == 200
    profiles = client.get("/api/payroll/import-profiles", headers=headers).json()["data"]
    assert profiles[0]["name"] == "Client export" and profiles[0]["column_mapping"]["Base Pay"] == "basic"
    again = client.get("/api/config/bundle/export.csv", headers=headers).content
    assert b"import_mappings,2,file_column,Base Pay" in again
    assert _send(client, headers, "round.csv", again, dry_run=False).status_code == 200


def test_unusable_files_are_refused_plainly(client):
    headers = _workspace(client, "config-bundle-refuse@example.com")
    legacy = b"section,record,field,type,value\r\ncomponents,1,component_name,text,Basic\r\n"
    r = _send(client, headers, "old.csv", legacy)
    assert r.status_code == 422 and "first line must be" in r.text
    assert _send(client, headers, "config.json", b"{}").status_code == 422
    blank = _send(client, headers, "blank.csv", _csv("components,1,component_name,"))
    assert blank.status_code == 422 and "no filled rows" in blank.text
    bad = _send(client, headers, "bad.csv", _csv("components,1,component_name,Basic", "components,1,pf_applicable,maybe"))
    assert bad.status_code == 422 and "Yes or No" in bad.text
