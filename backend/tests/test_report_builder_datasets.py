"""Report Builder: three datasets, units, pivots, charts, PDF and schedules.

Each dataset is the same arithmetic the product's pages use; these tests pin
the parts a report adds on top: that movement adds up, that findings count a
person's exposure once, that a pivot totals only what can be totalled, that
the PDF states what the Excel states, and that a schedule makes last month's
file for its owner — and stops, saying why, when the owner loses access.
"""
from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest

from app.database import SessionLocal
from app.models import FindingRecord, ReportJob, ReportSchedule, ValidationRun
from app.services import report_builder, report_exports, report_jobs
from tests.test_reports_and_audit import _dims, _register, workspace  # noqa: F401

APR, MAY, JUN = date(2026, 4, 1), date(2026, 5, 1), date(2026, 6, 1)


def _spec(**overrides):
    return {"dataset": "payroll_cost", "dimension": "department",
            "fields": ["period", "dimension", "headcount", "gross", "ctc"],
            "date_from": "2026-04-01", "date_to": "2026-06-01", **overrides}


def _preview(client, headers, spec):
    r = client.post("/api/reports/builder/preview", json={"specification": spec}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["data"]


def _staff(ids_by_dept: dict[str, list[str]]) -> list[dict]:
    return [{"employee_id": e, "dimensions": _dims(department=d), "components": {"basic": 20000}}
            for d, ids in ids_by_dept.items() for e in ids]


# ── the catalogue and the definition ────────────────────────────────────────

def test_the_catalogue_offers_three_datasets_each_with_units_and_additivity():
    catalogue = {d["key"]: d for d in report_builder.dataset_catalogue()}
    assert set(catalogue) == {"payroll_cost", "workforce_movement", "validation_findings"}
    cost = {f["key"]: f for f in catalogue["payroll_cost"]["fields"]}
    assert cost["ctc"]["unit"] == "inr" and cost["ctc"]["across_periods"] and cost["ctc"]["across_groups"]
    # People paid adds up across departments in a month, never across months.
    assert cost["headcount"]["across_groups"] and not cost["headcount"]["across_periods"]
    findings = {d["key"] for d in catalogue["validation_findings"]["dimensions"]}
    assert findings == {"severity", "check", "component"}
    assert catalogue["validation_findings"]["filters"] is False


def test_a_calculation_states_its_unit_and_names_only_its_datasets_metrics():
    spec = report_builder.validate(_spec(
        calculations=[{"key": "calc_share", "label": "Employer share", "expression": "employer_cost / ctc * 100", "unit": "pct"}],
        fields=["period", "dimension", "calc_share"]))
    assert spec["calculations"][0]["unit"] == "pct"
    meta = {c["key"]: c for c in report_builder.column_meta(spec)}
    assert meta["calc_share"]["label"] == "Employer share" and meta["calc_share"]["unit"] == "pct"
    with pytest.raises(ValueError, match="unit"):
        report_builder.validate(_spec(calculations=[{"key": "calc_x", "expression": "gross", "unit": "euros"}]))
    # "joiners" belongs to movement, not cost.
    with pytest.raises(ValueError, match="Unknown calculation metric"):
        report_builder.validate(_spec(calculations=[{"key": "calc_x", "expression": "joiners * 2"}]))
    # A definition saved before units existed still reads, as a plain number.
    old = report_builder.validate(_spec(calculations=[{"key": "calc_half", "expression": "gross / 2"}]))
    assert old["calculations"][0]["unit"] == "number" and old["layout"] == "table"


def test_a_pivot_or_chart_needs_a_number_column_and_both_axes():
    with pytest.raises(ValueError, match="both the period and the breakdown"):
        report_builder.validate(_spec(layout="pivot", fields=["dimension", "ctc"]))
    with pytest.raises(ValueError, match="number columns"):
        report_builder.validate(_spec(chart="bar", pivot_value="dimension"))
    with pytest.raises(ValueError, match="no company-dimension filters"):
        report_builder.validate({"dataset": "validation_findings", "dimension": "severity",
                                 "fields": ["period", "dimension", "findings"],
                                 "filters": {"department": ["A"]}})


# ── workforce movement ──────────────────────────────────────────────────────

def test_movement_adds_up_and_an_unmeasurable_month_is_absent_not_zero(client, workspace):
    entity, user, headers = workspace
    _register(entity, user, APR, _staff({"Store": ["E1", "E2", "E3"], "Office": ["E4"]}))
    # May: E3 left, E5 joined Store, E2 moved to Office.
    _register(entity, user, MAY, _staff({"Store": ["E1", "E5"], "Office": ["E2", "E4"]}))
    data = _preview(client, headers, {
        "dataset": "workforce_movement", "dimension": "department",
        "fields": ["period", "dimension", "opening", "joiners", "moved_in", "exits", "moved_out", "closing", "attrition_pct"],
        "date_from": "2026-04-01", "date_to": "2026-05-01"})
    rows = {(r["period"], r["dimension"]): r for r in data["rows"]}
    # April has no March register to measure against: closing only.
    assert rows[("2026-04-01", "Store")]["opening"] is None and rows[("2026-04-01", "Store")]["closing"] == 3
    store, office = rows[("2026-05-01", "Store")], rows[("2026-05-01", "Office")]
    assert (store["opening"], store["joiners"], store["exits"], store["moved_out"], store["closing"]) == (3, 1, 1, 1, 2)
    assert (office["opening"], office["moved_in"], office["closing"]) == (1, 1, 2)
    for r in (store, office):
        assert r["opening"] + r["joiners"] + r["moved_in"] - r["exits"] - r["moved_out"] == r["closing"]
    assert store["attrition_pct"] == 40.0  # 1 exit over an average of 2.5
    assert data["control_totals"] == {"joiners": 1, "exits": 1}


# ── validation findings ─────────────────────────────────────────────────────

def _run_with_findings(entity, user, period, findings):
    db = SessionLocal()
    try:
        run = ValidationRun(user_id=user.id, entity_id=entity.id, period_month=period, status="current",
                            employee_count=3, total_findings=len(findings))
        db.add(run)
        db.flush()
        for employee, rule, severity, impact in findings:
            db.add(FindingRecord(run_id=run.id, entity_id=entity.id, period_month=period,
                                 fingerprint=uuid.uuid4().hex, employee_id=employee, rule_id=rule,
                                 rule_name=f"Rule {rule}", severity=severity, status="FAIL",
                                 financial_impact=Decimal(str(impact))))
        db.commit()
    finally:
        db.close()


def test_findings_count_a_persons_exposure_once_and_never_price_the_unpriced(client, workspace):
    entity, user, headers = workspace
    _run_with_findings(entity, user, JUN, [
        ("E1", "PF-001", "CRITICAL", 1800), ("E1", "PF-002", "CRITICAL", 900),   # overlapping, same person
        ("E2", "PF-001", "CRITICAL", 500),
        ("E3", "ID-001", "WARNING", 0),                                        # a check that does not price
    ])
    data = _preview(client, headers, {
        "dataset": "validation_findings", "dimension": "severity",
        "fields": ["period", "dimension", "findings", "critical", "employees_affected", "priced_exposure", "unpriced"],
        "date_from": "2026-06-01", "date_to": "2026-06-01"})
    rows = {r["dimension"]: r for r in data["rows"]}
    from app.services.explain import MONETARY_RULES
    critical = rows["Critical"]
    assert critical["findings"] == 3 and critical["employees_affected"] == 2
    if {"PF-001", "PF-002"} <= set(MONETARY_RULES):
        assert critical["priced_exposure"] == 2300.0  # E1's largest, 1,800, plus E2's 500
    assert rows["Warning"]["unpriced"] == 1 and rows["Warning"]["priced_exposure"] is None


# ── pivot and chart ─────────────────────────────────────────────────────────

def test_a_pivot_totals_only_what_can_be_totalled(client, workspace):
    entity, user, headers = workspace
    _register(entity, user, APR, _staff({"Store": ["E1", "E2"], "Office": ["E3"]}))
    _register(entity, user, MAY, _staff({"Store": ["E1", "E2"], "Office": ["E3"]}))
    money = _preview(client, headers, _spec(date_to="2026-05-01", layout="pivot", pivot_value="gross", chart="line"))
    p = money["pivot"]
    assert p["periods"] == ["2026-04-01", "2026-05-01"] and p["row_totals"]
    store = next(r for r in p["rows"] if r["dimension"] == "Store")
    assert store["total"] == store["cells"]["2026-04-01"] + store["cells"]["2026-05-01"]
    assert p["grand_total"] == money["control_totals"]["gross"]
    assert money["chart"]["type"] == "line" and {s["group"] for s in money["chart"]["series"]} == {"Store", "Office"}

    heads = _preview(client, headers, _spec(date_to="2026-05-01", layout="pivot", pivot_value="headcount"))["pivot"]
    assert heads["row_totals"] is False and all(r["total"] is None for r in heads["rows"])
    assert heads["column_totals"] == {"2026-04-01": 3, "2026-05-01": 3}
    assert "not totalled across months" in heads["note"]


# ── output ──────────────────────────────────────────────────────────────────

def _saved(client, headers, spec, name="Report"):
    r = client.post("/api/reports/builder/saved", json={"name": name, "specification": spec,
                                                        "visibility": "private"}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["data"]["id"]


def _generate(client, headers, report_id, fmt):
    job = client.post(f"/api/reports/builder/saved/{report_id}/jobs", json={"format": fmt}, headers=headers)
    assert job.status_code == 200, job.text
    job_id = job.json()["data"]["id"]
    for _ in range(20):
        report_jobs.run_once()
        state = client.get(f"/api/reports/builder/jobs/{job_id}", headers=headers).json()["data"]
        if state["state"] in ("succeeded", "failed"):
            break
    return state


def test_excel_carries_the_pivot_a_native_chart_and_unit_formats(client, workspace):
    from openpyxl import load_workbook

    entity, user, headers = workspace
    _register(entity, user, APR, _staff({"Store": ["E1"], "Office": ["E2"]}))
    report_id = _saved(client, headers, _spec(date_to="2026-04-01", layout="pivot", pivot_value="ctc", chart="bar"))
    state = _generate(client, headers, report_id, "xlsx")
    assert state["state"] == "succeeded", state
    assert state["format"] == "xlsx" and state["origin"] == "person"
    download = client.get(f"/api/reports/builder/jobs/{state['id']}/download", headers=headers)
    assert download.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(download.content))
    assert "Pivot - Total CTC" in wb.sheetnames and "Chart" in wb.sheetnames
    assert wb["Chart"]._charts, "the chart sheet has no chart"
    details = wb["Details"]
    assert [c.value for c in details[1]] == ["Payroll period", "Department", "People paid", "Gross pay", "Total CTC"]
    assert details.cell(row=2, column=4).number_format.startswith('"₹"')


def test_pdf_states_what_the_report_covers(client, workspace):
    pytest.importorskip("fpdf")
    try:
        report_exports._font_dir()
    except ValueError:
        pytest.skip("DejaVu font not installed here; the production image installs fonts-dejavu-core")
    entity, user, headers = workspace
    _register(entity, user, APR, _staff({"Store": ["E1", "E2"], "Office": ["E3"]}))
    _register(entity, user, MAY, _staff({"Store": ["E1"], "Office": ["E3"]}))
    report_id = _saved(client, headers, _spec(date_to="2026-05-01", layout="pivot", pivot_value="gross", chart="line"),
                       name="Store cost")
    state = _generate(client, headers, report_id, "pdf")
    assert state["state"] == "succeeded", state
    download = client.get(f"/api/reports/builder/jobs/{state['id']}/download", headers=headers)
    assert download.headers["content-type"] == "application/pdf"
    assert download.headers["content-disposition"].endswith('.pdf"')
    assert download.content.startswith(b"%PDF")
    from pypdf import PdfReader  # noqa: E402

    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(download.content)).pages)
    assert "Store cost" in text and "Period from: 2026-04-01" in text and "Gross pay by breakdown and month" in text


def test_indian_grouping():
    assert report_exports.fmt_value(12345678.5, "inr") == "₹1,23,45,678.50"
    assert report_exports.fmt_value(950, "number") == "950"
    assert report_exports.fmt_value(None, "inr") == "—"
    assert report_exports.fmt_value(-100000, "inr") == "₹-1,00,000.00"


# ── schedules ───────────────────────────────────────────────────────────────

def test_the_next_run_and_the_window_are_worked_out_in_india_time():
    # 5th of the month at 09:00 IST is 03:30 UTC.
    after = datetime(2026, 6, 10, 12, 0, tzinfo=UTC)
    assert report_jobs.next_run("monthly", 5, 9, after) == datetime(2026, 7, 5, 3, 30, tzinfo=UTC)
    assert report_jobs.next_run("monthly", 15, 9, after) == datetime(2026, 6, 15, 3, 30, tzinfo=UTC)
    # Wednesday 10 June 2026; Monday (0) next is 15 June.
    assert report_jobs.next_run("weekly", 0, 9, after) == datetime(2026, 6, 15, 3, 30, tzinfo=UTC)
    assert report_jobs.next_run("monthly", 5, 9, datetime(2026, 12, 20, tzinfo=UTC)) == datetime(2027, 1, 5, 3, 30, tzinfo=UTC)
    assert report_jobs.rolling_window(1, datetime(2026, 7, 5, 3, 30, tzinfo=UTC)) == ("2026-06-01", "2026-06-01")
    assert report_jobs.rolling_window(3, datetime(2026, 1, 5, 3, 30, tzinfo=UTC)) == ("2025-10-01", "2025-12-01")


def test_a_schedule_makes_last_months_file_for_its_owner_and_stops_when_access_goes(client, workspace):
    entity, user, headers = workspace
    _register(entity, user, JUN, _staff({"Store": ["E1"]}))
    report_id = _saved(client, headers, _spec(date_from="2026-01-01", date_to="2026-01-01"), name="Scheduled")
    put = client.put(f"/api/reports/builder/saved/{report_id}/schedule", headers=headers,
                     json={"frequency": "monthly", "day": 5, "hour": 9, "months": 1, "format": "xlsx"})
    assert put.status_code == 200, put.text
    schedule = put.json()["data"]["schedule"]
    assert schedule["next_run_at"] and "no email" in schedule["delivery"]
    assert client.put(f"/api/reports/builder/saved/{report_id}/schedule", headers=headers,
                      json={"frequency": "monthly", "day": 31}).status_code == 422

    db = SessionLocal()
    try:
        row = db.get(ReportSchedule, uuid.UUID(schedule["id"]))
        row.next_run_at = datetime(2026, 7, 5, 3, 30, tzinfo=UTC)  # as if saved in June
        db.commit()
        assert report_jobs.schedule_due(db, now=datetime(2026, 7, 5, 4, 0, tzinfo=UTC)) == 1
        db.refresh(row)
        job = db.get(ReportJob, row.last_job_id)
        # The saved definition says January; the schedule makes June — the month just closed.
        assert (job.specification["date_from"], job.specification["date_to"]) == ("2026-06-01", "2026-06-01")
        assert job.origin == "schedule" and job.requester_id == user.id
        assert row.next_run_at.replace(tzinfo=UTC) == datetime(2026, 8, 5, 3, 30, tzinfo=UTC)
    finally:
        db.close()
    report_jobs.run_once()
    jobs = client.get("/api/reports/builder/jobs", headers=headers).json()["data"]["jobs"]
    made = next(j for j in jobs if j["origin"] == "schedule")
    assert made["state"] == "succeeded" and made["record_count"] == 1

    # The owner loses access: the next run does not happen, and the schedule says why.
    from app.models import OrgMembership
    db = SessionLocal()
    try:
        membership = db.query(OrgMembership).filter(OrgMembership.user_id == user.id).one()
        membership.role = "viewer"
        db.commit()
        assert report_jobs.schedule_due(db, now=datetime(2026, 8, 5, 4, 0, tzinfo=UTC)) == 0
        row = db.get(ReportSchedule, uuid.UUID(schedule["id"]))
        assert row.enabled is False and "revoked" in row.last_error
    finally:
        db.close()
