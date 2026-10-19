"""Templates follow the company's own components and come as CSV or Excel."""
from __future__ import annotations

import io

from openpyxl import load_workbook

PASSWORD = "Passw0rd!x"


def _workspace(client, email):
    r = client.post("/api/auth/signup", json={"email": email, "password": PASSWORD, "company_name": "Template Co"})
    h = {"Authorization": f"Bearer {r.json()['data']['access_token']}"}
    h["X-Entity-Id"] = client.post("/api/org/entities", json={"name": "T"}, headers=h).json()["data"]["id"]
    return h


def _add(client, h, name):
    assert client.post("/api/components", json={"component_name": name}, headers=h).status_code == 201


def test_a_new_component_appears_in_the_register_and_ctc_templates(client):
    h = _workspace(client, "templates-one@template-example.com")
    first = client.get("/api/templates/register.csv", headers=h).text.splitlines()[0]
    assert "Employee ID" in first and "Basic" not in first
    _add(client, h, "Basic")
    _add(client, h, "Travel Allowance")
    for kind in ("register", "ctc"):
        header = client.get(f"/api/templates/{kind}.csv", headers=h).text.splitlines()[0]
        assert "Basic" in header and "Travel Allowance" in header, kind


def test_every_template_is_available_as_excel_with_a_guide(client):
    h = _workspace(client, "templates-two@template-example.com")
    _add(client, h, "Basic")
    for kind in ("register", "ctc", "attendance", "master"):
        r = client.get(f"/api/templates/{kind}.xlsx", headers=h)
        assert r.status_code == 200, kind
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Template", "How to fill"]
        assert wb["Template"]["A1"].value == "Employee ID"
        if kind in ("register", "ctc"):
            assert "Basic" in [c.value for c in wb["Template"][1]]


def test_other_formats_and_kinds_are_refused(client):
    h = _workspace(client, "templates-three@template-example.com")
    assert client.get("/api/templates/register.json", headers=h).status_code == 404
    assert client.get("/api/templates/nonsense.csv", headers=h).status_code == 404


def test_the_ctc_template_is_accepted_by_the_ctc_importer(client):
    h = _workspace(client, "templates-four@template-example.com")
    _add(client, h, "Basic")
    csv_bytes = client.get("/api/templates/ctc.csv", headers=h).content
    r = client.post("/api/ctc/upload", headers=h, files={"file": ("ctc.csv", csv_bytes, "text/csv")},
                    data={"meta": '{"default_effective_from": "2026-04-01"}'})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["records"][0]["employee_id"] == "E001"
