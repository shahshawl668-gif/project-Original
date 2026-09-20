"""
Sign-off and the evidence pack.

The properties under test are the ones a due diligence turns on: the record
says who approved what, it does not change afterwards, and a draft can never be
mistaken for an approved period.
"""
from __future__ import annotations

import io
import uuid
from datetime import date

import openpyxl
import pytest

from app.database import SessionLocal
from app.models import Entity, OrgMembership, User
from app.services import finding_store

PASSWORD = "Passw0rd!x"


def _signup(client, email: str) -> dict:
    r = client.post(
        "/api/auth/signup",
        json={"email": email, "password": PASSWORD, "company_name": "Signoff Tests"},
    )
    if r.status_code != 200:
        r = client.post("/api/auth/login", json={"email": email, "password": PASSWORD})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['data']['access_token']}"}


@pytest.fixture()
def workspace(client, request):
    email = f"{request.node.name[:40]}@signoff-example.com".replace("_", "-")
    headers = _signup(client, email)
    created = client.post(
        "/api/org/entities",
        json={"name": f"E-{request.node.name[:26]}", "pf_establishment_code": "KN/BNG/12345"},
        headers=headers,
    )
    entity_id = created.json()["data"]["id"]
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        entity = db.get(Entity, uuid.UUID(entity_id))
        db.expunge_all()
    finally:
        db.close()
    return entity, user, {**headers, "X-Entity-Id": entity_id}


def _finding(employee_id: str, rule_id: str, impact: float) -> dict:
    return {
        "employee_id": employee_id,
        "employee_name": f"Employee {employee_id}",
        "rule_id": rule_id,
        "rule_name": f"Rule {rule_id}",
        "component": "pf_employee",
        "severity": "CRITICAL",
        "status": "FAIL",
        "expected_value": "1800",
        "actual_value": "0",
        "difference": "1800",
        "reason": "PF not deducted",
        "financial_impact": impact,
    }


def _record(entity, user, period: date, findings: list[dict]) -> None:
    db = SessionLocal()
    try:
        finding_store.record_run(
            db,
            entity_id=entity.id,
            user_id=user.id,
            period_month=period,
            findings=findings,
            employee_count=len(findings),
        )
        db.commit()
    finally:
        db.close()


PERIOD = "2025-04-01"


def test_preview_shows_what_signing_would_commit_to(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1),
            [_finding("1001", "STAT-001", 1800.0), _finding("2002", "STAT-004", 900.0)])

    snapshot = client.get(f"/api/signoff/{PERIOD}/preview", headers=h).json()["data"]

    assert snapshot["period_month"] == PERIOD
    assert snapshot["entity"]["pf_establishment_code"] == "KN/BNG/12345"
    assert len(snapshot["outstanding_findings"]) == 2
    assert snapshot["validation_run"]["total_findings"] == 2
    # The rules and rates in force, so a finding can be re-derived later.
    assert "statutory" in snapshot["configuration"]
    assert "exposure" in snapshot["configuration"]


def test_accepted_findings_are_listed_apart_but_still_included(client, workspace):
    """The point of the record is that accepted items were accepted knowingly."""
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1),
            [_finding("1001", "STAT-001", 1800.0), _finding("2002", "STAT-004", 900.0)])

    fp = [
        f["fingerprint"]
        for f in client.get("/api/findings", headers=h).json()["data"]
        if f["employee_id"] == "1001"
    ][0]
    client.post(
        f"/api/findings/{fp}/decision",
        json={"state": "waived", "reason": "Employee exempt under para 26(6)"},
        headers=h,
    )

    snapshot = client.get(f"/api/signoff/{PERIOD}/preview", headers=h).json()["data"]
    assert [f["employee_id"] for f in snapshot["outstanding_findings"]] == ["2002"]
    assert len(snapshot["accepted_findings"]) == 1
    assert snapshot["accepted_findings"][0]["waiver_reason"] == "Employee exempt under para 26(6)"


def test_submit_then_sign_records_who_approved_it(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])

    submitted = client.post(
        "/api/signoff/submit",
        json={"period_month": PERIOD, "notes": "March corrections applied"},
        headers=h,
    ).json()["data"]
    assert submitted["state"] == "pending_approval"

    signed = client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h).json()["data"]
    assert signed["state"] == "signed"
    assert signed["signed_by_email"] == user.email
    assert signed["signed_at"] is not None
    assert signed["snapshot_digest"]


def test_a_signed_period_cannot_be_signed_twice(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)
    client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h)

    again = client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h)
    assert again.status_code == 400


def test_the_snapshot_does_not_move_when_the_data_does(client, workspace):
    """A record that changes with the underlying data is not a record."""
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])

    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)
    signed = client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h).json()["data"]
    digest_at_signing = signed["snapshot_digest"]

    # More problems turn up afterwards.
    _record(entity, user, date(2025, 4, 1),
            [_finding("1001", "STAT-001", 1800.0), _finding("9999", "STAT-004", 5000.0)])

    stored = client.get(f"/api/signoff/{PERIOD}", headers=h).json()["data"]
    assert stored["signoff"]["snapshot_digest"] == digest_at_signing
    assert len(stored["snapshot"]["outstanding_findings"]) == 1


def test_reopening_preserves_what_was_signed(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)
    client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h)

    reopened = client.post(
        "/api/signoff/reopen",
        json={"period_month": PERIOD, "reason": "Vendor re-issued the register"},
        headers=h,
    )
    assert reopened.status_code == 200
    assert reopened.json()["data"]["state"] == "reopened"

    history = client.get(f"/api/signoff/{PERIOD}", headers=h).json()["data"]["history"]
    reopen_event = next(e for e in history if e["to_state"] == "reopened")
    assert reopen_event["reason"] == "Vendor re-issued the register"
    assert reopen_event["actor_email"] == user.email
    # The signing itself is still on the record.
    assert any(e["to_state"] == "signed" for e in history)


def test_only_a_signed_period_can_be_reopened(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)

    r = client.post("/api/signoff/reopen", json={"period_month": PERIOD, "reason": "x"}, headers=h)
    assert r.status_code == 400


def test_signing_requires_more_than_an_analyst_role(client, workspace):
    """Whoever ran the payroll should not be the only person who looked at it."""
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)

    db = SessionLocal()
    try:
        membership = db.query(OrgMembership).filter(OrgMembership.user_id == user.id).one()
        membership.role = "analyst"
        db.add(membership)
        db.commit()
    finally:
        db.close()

    assert client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h).status_code == 403


def _workbook(response) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_evidence_pack_is_a_workbook_with_the_expected_sheets(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)
    client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h)

    r = client.get(f"/api/signoff/{PERIOD}/evidence-pack", headers=h)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    wb = _workbook(r)
    assert {"Cover", "Outstanding", "Accepted", "Exposure", "Minimum wage rates"} <= set(wb.sheetnames)

    cover = {row[0]: row[1] for row in wb["Cover"].iter_rows(values_only=True) if row[0]}
    assert cover["Source"] == "Signed snapshot"
    assert cover["Signed by"] == user.email
    assert cover["PF establishment code"] == "KN/BNG/12345"


def test_an_unsigned_pack_says_so_on_its_face(client, workspace):
    """A draft must never be mistakable for an approved record."""
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])

    wb = _workbook(client.get(f"/api/signoff/{PERIOD}/evidence-pack", headers=h))
    cover = {row[0]: row[1] for row in wb["Cover"].iter_rows(values_only=True) if row[0]}

    assert "NOT signed" in cover["Source"]
    assert cover["Status"] == "not started"


def test_the_pack_lists_accepted_findings_with_their_reasons(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [_finding("1001", "STAT-001", 1800.0)])
    fp = client.get("/api/findings", headers=h).json()["data"][0]["fingerprint"]
    client.post(
        f"/api/findings/{fp}/decision",
        json={"state": "waived", "reason": "Director confirmed exemption"},
        headers=h,
    )
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)
    client.post("/api/signoff/sign", json={"period_month": PERIOD}, headers=h)

    wb = _workbook(client.get(f"/api/signoff/{PERIOD}/evidence-pack", headers=h))
    rows = list(wb["Accepted"].iter_rows(values_only=True))

    assert rows[0][0] == "Employee ID"
    assert rows[1][0] == "1001"
    assert "Director confirmed exemption" in rows[1]


def test_sign_offs_do_not_leak_between_entities(client, workspace):
    entity, user, h = workspace
    _record(entity, user, date(2025, 4, 1), [])
    client.post("/api/signoff/submit", json={"period_month": PERIOD}, headers=h)

    other = client.post(
        "/api/org/entities", json={"name": "Another Client"}, headers=h
    ).json()["data"]["id"]

    assert client.get("/api/signoff", headers=h).json()["data"]
    assert client.get("/api/signoff", headers={**h, "X-Entity-Id": other}).json()["data"] == []
    assert client.get(f"/api/signoff/{PERIOD}", headers={**h, "X-Entity-Id": other}).status_code == 404
