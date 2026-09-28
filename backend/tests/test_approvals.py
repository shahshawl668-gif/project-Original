"""
Maker–checker and readiness.

A month cannot be prepared or approved while there is nothing trustworthy to
approve (no run, a stale run, or material checks that could not run without a
stated reason), and — where the owner requires it — the person who prepared a
month cannot approve it, nor publish a rule they drafted.
"""
from __future__ import annotations

import pytest

from tests.test_coverage import BASE_ROWS, _register, _run
from tests.test_run_history import PERIOD, PASSWORD, _company, _data, _empty_queue  # noqa: F401


def _second_manager(client, owner_headers) -> dict:
    """A second person in the same organisation, as a manager."""
    email = f"checker-{owner_headers['X-Entity-Id'][:8]}@approvals-example.com"
    inv = _data(client.post("/api/org/invitations", headers=owner_headers,
                            json={"email": email, "role": "manager"}))
    joined = client.post("/api/org/invitations/register", json={"token": inv["token"], "password": PASSWORD})
    assert joined.status_code == 200, joined.text
    token = joined.json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}", "X-Entity-Id": owner_headers["X-Entity-Id"]}


@pytest.fixture()
def company(client):
    return _company(client, "appr")


def test_an_unvalidated_month_cannot_be_prepared(client, company):
    r = client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD})
    assert r.status_code == 409
    assert "not been validated" in r.json()["error"]["detail"]
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["readiness"]["ready"] is False


def test_a_stale_month_cannot_be_prepared(client, company):
    _run(client, company, _register(BASE_ROWS))
    comps = _data(client.get("/api/components", headers=company))
    hra = next(c for c in comps if c["component_name"] == "HRA")
    assert client.patch(f"/api/components/{hra['id']}", headers=company,
                        json={"pf_applicable": True}).status_code == 200
    r = client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD})
    assert r.status_code == 409
    assert "Revalidate" in r.json()["error"]["detail"]


def test_checks_that_could_not_run_need_a_stated_reason(client, company):
    _run(client, company, _register(BASE_ROWS, drop={"pf_employee"}))
    refused = client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD})
    assert refused.status_code == 409
    assert "could not be performed" in refused.json()["error"]["detail"]
    reason = "PF is verified from the ECR this month, not the register"
    ok = client.post("/api/signoff/submit", headers=company,
                     json={"period_month": PERIOD, "accept_incomplete_reason": reason})
    assert ok.status_code == 200, ok.text
    snap = _data(client.get(f"/api/signoff/{PERIOD}", headers=company))["snapshot"]
    assert snap["readiness"]["accepted_gaps"]["reason"] == reason
    assert snap["validation_run"]["coverage"]["material_cannot_validate"] > 0


def test_self_approval_is_refused_when_independence_is_required(client, company):
    _run(client, company, _register(BASE_ROWS))
    reason = {"accept_incomplete_reason": "Minimum wage not yet configured"}
    policy = _data(client.put("/api/org/approval-policy", headers=company,
                              json={"signoff_requires_independent_approver": True}))
    assert policy["signoff_requires_independent_approver"] is True
    assert client.post("/api/signoff/submit", headers=company,
                       json={"period_month": PERIOD, **reason}).status_code == 200
    own = client.post("/api/signoff/sign", headers=company, json={"period_month": PERIOD, **reason})
    assert own.status_code == 409
    assert "someone other than" in own.json()["error"]["detail"]

    checker = _second_manager(client, company)
    signed = client.post("/api/signoff/sign", headers=checker, json={"period_month": PERIOD, **reason})
    assert signed.status_code == 200, signed.text
    snap = _data(client.get(f"/api/signoff/{PERIOD}", headers=company))["snapshot"]
    assert snap["approval"]["independent"] is True
    assert snap["approval"]["policy"]["signoff_requires_independent_approver"] is True


def test_without_the_policy_self_approval_is_allowed_but_recorded(client, company):
    _run(client, company, _register(BASE_ROWS))
    reason = {"accept_incomplete_reason": "Minimum wage not yet configured"}
    client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD, **reason})
    signed = client.post("/api/signoff/sign", headers=company, json={"period_month": PERIOD, **reason})
    assert signed.status_code == 200, signed.text
    snap = _data(client.get(f"/api/signoff/{PERIOD}", headers=company))["snapshot"]
    assert snap["approval"]["independent"] is False


def test_only_an_owner_changes_approval_controls(client, company):
    checker = _second_manager(client, company)
    r = client.put("/api/org/approval-policy", headers=checker,
                   json={"signoff_requires_independent_approver": False})
    assert r.status_code == 403
    events = _data(client.get("/api/audit?limit=50", headers=company))["events"]
    client.put("/api/org/approval-policy", headers=company, json={"signoff_requires_independent_approver": True})
    events = _data(client.get("/api/audit?limit=50", headers=company))["events"]
    assert any(e["action"] == "approval_policy.changed" for e in events)


def test_a_changed_input_after_signing_is_reported_not_hidden(client, company):
    _run(client, company, _register(BASE_ROWS))
    reason = {"accept_incomplete_reason": "Minimum wage not yet configured"}
    client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD, **reason})
    assert client.post("/api/signoff/sign", headers=company,
                       json={"period_month": PERIOD, **reason}).status_code == 200
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "signed_off" and status["changed_since_signoff"] == []
    comps = _data(client.get("/api/components", headers=company))
    hra = next(c for c in comps if c["component_name"] == "HRA")
    client.patch(f"/api/components/{hra['id']}", headers=company, json={"pf_applicable": True})
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "signed_off"
    assert [c["input"] for c in status["changed_since_signoff"]] == ["configuration"]


def test_a_rule_cannot_be_published_by_its_author_when_independence_is_required(client, company):
    from tests.test_validation_matrix import rule_body

    client.put("/api/org/approval-policy", headers=company,
               json={"matrix_publish_requires_independent_approver": True})
    draft = _data(client.post("/api/validation-matrix", headers=company, json=rule_body(key="CUST-IND-01")))
    _data(client.post(f"/api/validation-matrix/{draft['id']}/submit", headers=company))
    own = client.post(f"/api/validation-matrix/{draft['id']}/publish", headers=company)
    assert own.status_code == 409
    checker = _second_manager(client, company)
    other = client.post(f"/api/validation-matrix/{draft['id']}/publish", headers=checker)
    assert other.status_code == 200, other.text


def test_the_evidence_pack_carries_the_run_its_inputs_and_its_coverage(client, company):
    import io

    from openpyxl import load_workbook

    run_id = _run(client, company, _register(BASE_ROWS))
    reason = {"accept_incomplete_reason": "Minimum wage not yet configured"}
    client.post("/api/signoff/submit", headers=company, json={"period_month": PERIOD, **reason})
    client.post("/api/signoff/sign", headers=company, json={"period_month": PERIOD, **reason})
    pack = client.get(f"/api/signoff/{PERIOD}/evidence-pack", headers=company)
    assert pack.status_code == 200
    book = load_workbook(io.BytesIO(pack.content))
    for sheet in ("Run and inputs", "Coverage", "Findings in the run", "PT and LWF schedules"):
        assert sheet in book.sheetnames
    items = {r[0].value: r[1].value for r in book["Run and inputs"].iter_rows(min_row=2)}
    assert items["Validation run ID"] == run_id
    assert len(items["Source file SHA-256"]) == 64
    assert items["Gaps accepted at approval"] == "Minimum wage not yet configured"
    findings = list(book["Findings in the run"].iter_rows(min_row=2, values_only=True))
    e002 = next(r for r in findings if r[0] == "E002" and r[4] == "STAT-001")
    assert e002[2] == 3 and e002[10] == 100.0          # file row 3, ₹100 short
    assert any(r[10] == "Impact not calculated" for r in findings)
