"""
Durable validation and preserved run history.

The journey under test is the product's own: upload → queue → worker → read
the run by id. The properties that matter are the ones a payroll team relies on
without seeing: a re-run never erases the earlier one, a failed or cancelled
job leaves nothing half-written, a stale run says so, and nothing crosses from
one company to another.

Expected results are computed here, independently of the engine: PF is 12% of
PF wages (Basic + Special Allowance) capped at the ₹15,000 ceiling. E003's PF
is planted at zero in the first upload and corrected in the second.
"""
from __future__ import annotations

import hashlib
import io
import json
import uuid

import pytest
from sqlalchemy import text

from app.database import SessionLocal
from app.models import RegisterUpload, ValidationJob, ValidationRun, ValidationRunEmployee
from app.services import validation_jobs as jobs
from app.services import validation_worker as worker

PASSWORD = "Passw0rd!x"
PERIOD = "2026-06-01"

COMPONENTS = [
    {"component_name": "Basic", "pf_applicable": True, "esic_applicable": True,
     "pt_applicable": True, "lwf_applicable": True, "included_in_wages": True, "taxable": True},
    {"component_name": "HRA", "esic_applicable": True, "pt_applicable": True, "taxable": True},
    {"component_name": "Special Allowance", "pf_applicable": True, "esic_applicable": True,
     "pt_applicable": True, "taxable": True},
]

# employee, basic, pan
STAFF = [
    ("E001", 40000, "ABCPE1234F"),
    ("E002", 30000, "ABCPE1235F"),
    ("E003", 25000, "ABCPE1236F"),
    ("E004", 22000, "BAD-PAN"),   # planted: invalid PAN, only visible from the upload itself
    ("E005", 50000, "ABCPE1238F"),
]


def expected_pf(basic: int) -> int:
    special = round(basic * 0.25)
    return round(min(basic + special, 15000) * 0.12)


def register_csv(*, pf_zero_for: set[str]) -> bytes:
    lines = ["employee_id,employee_name,state,paid_days,lop_days,basic,hra,special allowance,"
             "pf_employee,pf_employer,pt,pan"]
    for eid, basic, pan in STAFF:
        hra, special = round(basic * 0.4), round(basic * 0.25)
        pf = 0 if eid in pf_zero_for else expected_pf(basic)
        lines.append(f"{eid},Person {eid},Karnataka,30,0,{basic},{hra},{special},{pf},{pf},200,{pan}")
    return ("\n".join(lines) + "\n").encode()


def _data(r):
    assert r.status_code == 200, r.text
    return r.json()["data"]


def _company(client, tag: str) -> dict:
    email = f"history-{tag}-{uuid.uuid4().hex[:6]}@history-example.com"
    r = client.post("/api/auth/signup", json={
        "email": email, "password": PASSWORD, "company_name": f"History {tag} Pvt Ltd",
    })
    headers = {"Authorization": f"Bearer {_data(r)['access_token']}"}
    entity_id = _data(client.post(
        "/api/org/entities", json={"name": f"History {tag}", "primary_state": "Karnataka"},
        headers=headers,
    ))["id"]
    headers["X-Entity-Id"] = entity_id
    for component in COMPONENTS:
        assert client.post("/api/components", json=component, headers=headers).status_code == 201
    for rule_type in ("PT", "LWF"):
        client.post(f"/api/rule-engine/slabs/import-defaults/all?rule_type={rule_type}&overwrite=true",
                    headers=headers)
    return headers


def _upload(client, headers, content: bytes, name="register.csv"):
    return _data(client.post(
        "/api/payroll/upload", headers=headers,
        files={"file": (name, io.BytesIO(content), "text/csv")},
        data={"meta": json.dumps({
            "period_month": PERIOD, "strict_header_check": False, "return_employees": False,
        })},
    ))


def _enqueue(client, headers, **extra):
    return _data(client.post("/api/validation/jobs", headers=headers,
                             json={"period_month": PERIOD, **extra}))


def _drain():
    db = SessionLocal()
    try:
        return worker.drain(db, "test-worker")
    finally:
        db.close()


@pytest.fixture(autouse=True)
def _empty_queue():
    """Other tests leave jobs behind; each test here starts with none live."""
    db = SessionLocal()
    try:
        db.query(ValidationJob).filter(ValidationJob.state.in_(("queued", "running"))).update(
            {"state": "cancelled"}, synchronize_session=False)
        db.commit()
    finally:
        db.close()
    yield


@pytest.fixture()
def company(client):
    return _company(client, "a")


def _validated(client, headers, content: bytes) -> dict:
    _upload(client, headers, content)
    job = _enqueue(client, headers)["job"]
    assert _drain() >= 1
    job = _data(client.get(f"/api/validation/jobs/{job['id']}", headers=headers))
    assert job["state"] == "succeeded", job
    return job


# ---------------------------------------------------------------------------
# Uploads are kept, with their hashes
# ---------------------------------------------------------------------------
def test_every_upload_is_kept_with_the_file_hash(client, company):
    first = register_csv(pf_zero_for={"E003"})
    second = register_csv(pf_zero_for=set())
    a = _upload(client, company, first)["upload"]
    b = _upload(client, company, second)["upload"]

    assert a["file_sha256"] == hashlib.sha256(first).hexdigest()
    assert b["file_sha256"] == hashlib.sha256(second).hexdigest()
    assert (a["revision"], b["revision"]) == (1, 2)
    listed = _data(client.get(f"/api/validation/uploads?period_month={PERIOD}", headers=company))
    assert [u["revision"] for u in listed] == [2, 1]


def test_the_upload_screen_can_ask_for_a_summary_instead_of_every_row(client, company):
    out = _upload(client, company, register_csv(pf_zero_for=set()))
    assert out["employees"] == []
    assert out["employee_count"] == len(STAFF)


# ---------------------------------------------------------------------------
# The job journey
# ---------------------------------------------------------------------------
def test_a_queued_validation_produces_a_run_readable_by_id(client, company):
    job = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    assert job["percent"] == 100 and job["stage"] == "succeeded"

    run = _data(client.get(f"/api/validation/runs/{job['run_id']}", headers=company))
    assert run["status"] == "current" and run["run_number"] == 1
    assert run["source"] == "job"
    assert run["upload"]["revision"] == 1
    assert run["inputs_recorded"] is True
    assert set(run["input_digests"]) >= {"register", "master", "attendance", "configuration", "engine", "file"}

    employees = _data(client.get(
        f"/api/validation/runs/{job['run_id']}/employees?page_size=2&sort=employee_id&order=asc",
        headers=company))
    assert employees["total"] == len(STAFF)
    assert employees["pages"] == 3
    assert [e["employee_id"] for e in employees["items"]] == ["E001", "E002"]

    pf = _data(client.get(
        f"/api/validation/runs/{job['run_id']}/findings?rule_id=STAT-001", headers=company))
    assert {f["employee_id"] for f in pf["items"]} == {"E003"}
    # Independently: E003's shortfall is the whole expected PF.
    assert pf["items"][0]["expected_value"] == f"{expected_pf(25000):.2f}"


def test_validating_the_frozen_upload_keeps_checks_the_stored_register_would_lose(client, company):
    """PAN is not a column the decomposed register keeps. The job still checks it."""
    job = _validated(client, company, register_csv(pf_zero_for=set()))
    ids = _data(client.get(
        f"/api/validation/runs/{job['run_id']}/findings?rule_id=ID-001", headers=company))
    assert {f["employee_id"] for f in ids["items"]} == {"E004"}


def test_a_second_request_joins_the_live_job_instead_of_starting_another(client, company):
    _upload(client, company, register_csv(pf_zero_for=set()))
    first = _enqueue(client, company)
    second = _enqueue(client, company)
    assert first["already_queued"] is False
    assert second["already_queued"] is True
    assert second["job"]["id"] == first["job"]["id"]


def test_a_returning_user_finds_their_active_job(client, company):
    _upload(client, company, register_csv(pf_zero_for=set()))
    job = _enqueue(client, company)["job"]
    active = _data(client.get("/api/validation/jobs?active=true", headers=company))
    assert [j["id"] for j in active] == [job["id"]]
    assert active[0]["queue_position"] == 0


def test_validation_without_an_upload_is_refused_with_what_to_do(client, company):
    r = client.post("/api/validation/jobs", headers=company, json={"period_month": "2026-01-01"})
    assert r.status_code == 400
    assert "Upload it first" in r.json()["error"]["detail"]


# ---------------------------------------------------------------------------
# Cancellation and retry leave nothing half-written
# ---------------------------------------------------------------------------
def test_cancelling_a_queued_job_is_immediate(client, company):
    _upload(client, company, register_csv(pf_zero_for=set()))
    job = _enqueue(client, company)["job"]
    out = _data(client.post(f"/api/validation/jobs/{job['id']}/cancel", headers=company))
    assert out["state"] == "cancelled"
    assert _drain() == 0


def test_cancelling_a_running_job_stops_it_and_writes_no_run(client, company, monkeypatch):
    _upload(client, company, register_csv(pf_zero_for=set()))
    job = _enqueue(client, company)["job"]
    real = worker.validate_employees

    def slow(*args, on_progress=None, **kwargs):
        # A person presses Cancel while the worker is mid-register.
        db = SessionLocal()
        try:
            j = db.get(ValidationJob, uuid.UUID(job["id"]))
            jobs.request_cancel(db, j, None)
            db.commit()
        finally:
            db.close()
        on_progress(1)
        return real(*args, on_progress=on_progress, **kwargs)

    monkeypatch.setattr(worker, "validate_employees", slow)
    _drain()
    out = _data(client.get(f"/api/validation/jobs/{job['id']}", headers=company))
    assert out["state"] == "cancelled"
    assert out["run_id"] is None
    runs = _data(client.get(f"/api/validation/runs?period_month={PERIOD}", headers=company))
    assert runs == []


def test_a_failed_job_can_be_retried_as_a_new_job_and_produces_one_run(client, company, monkeypatch):
    _upload(client, company, register_csv(pf_zero_for=set()))
    job = _enqueue(client, company)["job"]
    db = SessionLocal()
    try:
        j = db.get(ValidationJob, uuid.UUID(job["id"]))
        j.max_attempts = 1
        db.commit()
    finally:
        db.close()

    def explode(*_a, **_k):
        raise ConnectionError("lost the database")

    monkeypatch.setattr(worker, "validate_employees", explode)
    _drain()
    failed = _data(client.get(f"/api/validation/jobs/{job['id']}", headers=company))
    assert failed["state"] == "failed" and failed["can_retry"] is True
    assert "lost the database" not in (failed["error"] or "")
    monkeypatch.undo()

    retry = _data(client.post(f"/api/validation/jobs/{job['id']}/retry", headers=company))["job"]
    assert retry["retry_of_job_id"] == job["id"]
    _drain()
    done = _data(client.get(f"/api/validation/jobs/{retry['id']}", headers=company))
    assert done["state"] == "succeeded"
    runs = _data(client.get(f"/api/validation/runs?period_month={PERIOD}", headers=company))
    assert len(runs) == 1


# ---------------------------------------------------------------------------
# History is preserved and comparable
# ---------------------------------------------------------------------------
def test_a_rerun_supersedes_rather_than_deletes_and_the_two_can_be_compared(client, company):
    first = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    second = _validated(client, company, register_csv(pf_zero_for=set()))

    runs = _data(client.get(f"/api/validation/runs?period_month={PERIOD}", headers=company))
    assert [(r["run_number"], r["status"]) for r in runs] == [(2, "current"), (1, "superseded")]
    assert runs[1]["superseded_by_run_id"] == second["run_id"]

    # The first run still reads exactly as it was reported.
    old = _data(client.get(
        f"/api/validation/runs/{first['run_id']}/findings?rule_id=STAT-001", headers=company))
    assert old["total"] == 1

    diff = _data(client.get(
        f"/api/validation/runs/compare?base={first['run_id']}&target={second['run_id']}",
        headers=company))
    resolved = {(i["employee_id"], i["rule_id"]) for i in diff["items"]["resolved"]}
    assert ("E003", "STAT-001") in resolved
    # Deducting E003's PF makes "UAN missing with PF deduction" apply to them,
    # as it already did to everyone else in this fixture (no UAN column). The
    # comparison reports that consequence, and only that, as new.
    assert {(i["employee_id"], i["rule_id"]) for i in diff["items"]["new"]} == {("E003", "ID-003")}
    # The PAN problem was in both uploads and did not change.
    unchanged = {(i["employee_id"], i["rule_id"]) for i in diff["items"]["unchanged"]}
    assert ("E004", "ID-001") in unchanged
    # The lifecycle agrees with the comparison: fixed in a same-month re-run is resolved.
    state = _data(client.get("/api/findings?rule_id=STAT-001&employee_id=E003", headers=company))
    assert state and state[0]["state"] == "resolved"


def test_only_one_run_per_period_is_ever_current(client, company):
    _validated(client, company, register_csv(pf_zero_for=set()))
    _validated(client, company, register_csv(pf_zero_for=set()))
    _validated(client, company, register_csv(pf_zero_for=set()))
    db = SessionLocal()
    try:
        entity_id = uuid.UUID(company["X-Entity-Id"])
        statuses = [r.status for r in db.query(ValidationRun).filter(ValidationRun.entity_id == entity_id)]
        assert sorted(statuses) == ["current", "superseded", "superseded"]
    finally:
        db.close()


def test_the_synchronous_endpoint_also_keeps_history(client, company):
    content = register_csv(pf_zero_for=set())
    full = _data(client.post(
        "/api/payroll/upload", headers=company,
        files={"file": ("r.csv", io.BytesIO(content), "text/csv")},
        data={"meta": json.dumps({"period_month": PERIOD, "strict_header_check": False})},
    ))
    body = {"employees": full["employees"], "run_type": "regular", "period_month": PERIOD}
    one = _data(client.post("/api/payroll/validate", headers=company, json=body))["lifecycle"]
    two = _data(client.post("/api/payroll/validate", headers=company, json=body))["lifecycle"]
    assert (one["run_number"], two["run_number"]) == (1, 2)
    first = _data(client.get(f"/api/validation/runs/{one['run_id']}", headers=company))
    assert first["status"] == "superseded"


# ---------------------------------------------------------------------------
# Staleness
# ---------------------------------------------------------------------------
def test_a_changed_input_makes_the_run_stale_and_says_which(client, company):
    job = _validated(client, company, register_csv(pf_zero_for=set()))
    fresh = _data(client.get(f"/api/validation/runs/{job['run_id']}", headers=company))["freshness"]
    assert fresh["revalidation_required"] is False, fresh

    comps = _data(client.get("/api/components", headers=company))
    hra = next(c for c in comps if c["component_name"] == "HRA")
    r = client.patch(f"/api/components/{hra['id']}", headers=company, json={"pf_applicable": True})
    assert r.status_code == 200, r.text

    stale = _data(client.get(f"/api/validation/runs/{job['run_id']}", headers=company))["freshness"]
    assert stale["revalidation_required"] is True
    assert [c["input"] for c in stale["changes"]] == ["configuration"]
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "revalidation_required"


def test_a_reupload_makes_the_run_stale(client, company):
    job = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    _upload(client, company, register_csv(pf_zero_for=set()))
    stale = _data(client.get(f"/api/validation/runs/{job['run_id']}", headers=company))["freshness"]
    assert "register" in [c["input"] for c in stale["changes"]]


def test_the_month_is_never_clear_by_default(client, company):
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "not_uploaded"
    _upload(client, company, register_csv(pf_zero_for=set()))
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "uploaded"
    _enqueue(client, company)
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "validation_pending"
    _drain()
    status = _data(client.get(f"/api/validation/periods/{PERIOD}/status", headers=company))
    assert status["stage"] == "issues_found"  # the PAN finding is open


# ---------------------------------------------------------------------------
# Isolation
# ---------------------------------------------------------------------------
def test_another_company_cannot_see_any_of_it(client, company):
    job = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    stranger = _company(client, "b")
    run_id, job_id = job["run_id"], job["id"]
    for path in (
        f"/api/validation/jobs/{job_id}",
        f"/api/validation/runs/{run_id}",
        f"/api/validation/runs/{run_id}/employees",
        f"/api/validation/runs/{run_id}/employees/E001",
        f"/api/validation/runs/{run_id}/findings",
        f"/api/validation/runs/{run_id}/configuration",
        f"/api/validation/runs/{run_id}/export.xlsx",
        f"/api/validation/runs/compare?base={run_id}&target={run_id}",
    ):
        assert client.get(path, headers=stranger).status_code == 404, path
    for path in (f"/api/validation/jobs/{job_id}/cancel", f"/api/validation/jobs/{job_id}/retry"):
        assert client.post(path, headers=stranger).status_code == 404, path
    assert _data(client.get("/api/validation/uploads", headers=stranger)) == []


def test_employee_detail_shows_the_row_as_uploaded(client, company):
    job = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    detail = _data(client.get(f"/api/validation/runs/{job['run_id']}/employees/E003", headers=company))
    assert detail["source_row"]["pf_employee"] == 0
    assert detail["source_row"]["_source_row"] == 4  # header is row 1; E003 is the third data row
    assert any(f["rule_id"] == "STAT-001" for f in detail["result"]["findings"])


def test_the_run_exports_as_recorded(client, company):
    job = _validated(client, company, register_csv(pf_zero_for={"E003"}))
    r = client.get(f"/api/validation/runs/{job['run_id']}/export.xlsx", headers=company)
    assert r.status_code == 200
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Run", "Findings", "Employees"]
    assert book["Employees"].max_row == len(STAFF) + 1


# ---------------------------------------------------------------------------
# Migration of existing data
# ---------------------------------------------------------------------------
def test_the_history_migration_upgrades_an_old_runs_table_in_place(tmp_path):
    """An installation with runs from before history keeps them, as run 1, current."""
    from sqlalchemy import create_engine

    from app.migrations import preserve_run_history

    engine = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with engine.begin() as conn:
        conn.execute(text(
            "CREATE TABLE validation_runs (id CHAR(32) PRIMARY KEY, entity_id CHAR(32), "
            "period_month DATE, total_findings INTEGER)"
        ))
        conn.execute(text(
            "INSERT INTO validation_runs VALUES ('a'||hex(randomblob(15)), 'e1', '2026-04-01', 3)"
        ))
        conn.execute(text(
            "CREATE TABLE validation_jobs (id CHAR(32) PRIMARY KEY, state VARCHAR(16))"
        ))
    for _ in range(2):  # idempotent
        with engine.begin() as conn:
            preserve_run_history(conn)
    with engine.begin() as conn:
        row = conn.execute(text("SELECT status, run_number, total_findings FROM validation_runs")).one()
        assert tuple(row) == ("current", 1, 3)
        cols = {r[1] for r in conn.execute(text("PRAGMA table_info(validation_jobs)"))}
        assert {"stage", "upload_id", "error_message", "cancel_requested_at"} <= cols


def test_old_runs_without_recorded_inputs_are_reported_as_unknown_not_fresh(client, company):
    job = _validated(client, company, register_csv(pf_zero_for=set()))
    db = SessionLocal()
    try:
        run = db.get(ValidationRun, uuid.UUID(job["run_id"]))
        run.input_digests = None
        db.commit()
    finally:
        db.close()
    fresh = _data(client.get(f"/api/validation/runs/{job['run_id']}", headers=company))["freshness"]
    assert fresh["revalidation_required"] is True
    assert fresh["changes"][0]["input"] == "unrecorded"


def test_run_employee_rows_are_compact(client, company):
    job = _validated(client, company, register_csv(pf_zero_for=set()))
    db = SessionLocal()
    try:
        rows = db.query(ValidationRunEmployee).filter(
            ValidationRunEmployee.run_id == uuid.UUID(job["run_id"])).all()
        assert len(rows) == len(STAFF)
        assert all(len(r.detail_gz) < 8000 for r in rows)
        upload = db.query(RegisterUpload).filter(
            RegisterUpload.entity_id == uuid.UUID(company["X-Entity-Id"])).first()
        assert upload.rows_gz[:2] == b"\x1f\x8b"  # gzip
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Every field the engine reads must be importable
# ---------------------------------------------------------------------------
#: Canonical names the engine reads from a register row (rule_engine_v2,
#: validation, identity_checks, row_composition, pf_basis). When a new check
#: reads a new column, it goes here and into the importer together.
ENGINE_READS = {
    "aadhaar", "adolescent_permit", "bank_account", "death_or_disablement", "disability",
    "dob", "doj", "dol", "employment_type", "esi_number", "esic_employee", "esic_employer",
    "gender", "ifsc", "international_worker", "lop_days", "lwf_employee", "lwf_employer",
    "net", "paid_days", "pan", "payment_mode", "pf_employee", "pf_employer", "pf_eps", "pt",
    "state", "tax_regime", "tds", "total_days", "total_deductions", "uan", "gross",
}


def test_every_field_the_engine_reads_can_be_imported():
    """A check whose input the importer drops can never run — or reports it missing."""
    from app.services.payroll_parse import allowed_destinations

    assert ENGINE_READS - allowed_destinations(set()) == set()


def test_an_esi_number_in_the_file_is_not_reported_missing(client, company):
    lines = register_csv(pf_zero_for=set()).decode().splitlines()
    lines[0] += ",esi_number,esic_employee,esic_employer"
    # Make E004 ESIC-eligible-looking by deducting ESIC, and give it an IP number.
    rows = [lines[0]]
    for line in lines[1:]:
        eid = line.split(",")[0]
        rows.append(line + (",1234567890,100,400" if eid == "E004" else ",,0,0"))
    job = _validated(client, company, ("\n".join(rows) + "\n").encode())
    missing = _data(client.get(
        f"/api/validation/runs/{job['run_id']}/findings?rule_id=ID-004", headers=company))
    assert "E004" not in {f["employee_id"] for f in missing["items"]}
    detail = _data(client.get(f"/api/validation/runs/{job['run_id']}/employees/E004", headers=company))
    assert detail["source_row"]["esi_number"] == "1234567890"
