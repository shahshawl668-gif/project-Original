"""Generate the statutory verification worksheet a reviewer signs.

GO_LIVE.md D6 makes a payroll professional's sign-off a gate on real client
data. A verbal "looks fine" does not survive an audit, so this emits the
artefact instead: every statutory value the product ships as a default, where
it came from, where it lives in the code, and a column for the reviewer to
accept or correct it.

Values are read from the application at run time rather than retyped, so the
worksheet cannot drift from what the software actually computes with. Re-run it
whenever a Budget lands or a state amends a schedule:

    python tools/statutory_worksheet.py [output.xlsx]

Reviewer columns are left empty on purpose. A pre-filled "Y" is not a review.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.branding import PRODUCT_NAME
from app.schemas.statutory_config import ESICConfig, PFConfig
from app.services import lwf_defaults, pt_defaults
from app.services.tax_year_defaults import DEFAULT_TAX_YEAR, default_tax_years

FONT = "Arial"

# Reviewer-facing colours. Yellow marks what the reviewer fills in; nothing
# else in the workbook is yellow, so "fill every yellow cell" is the whole
# instruction.
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="0C4A6E")
BAND_FILL = PatternFill("solid", fgColor="E0F2FE")
RISK_FILL = PatternFill("solid", fgColor="FEE2E2")

THIN = Side(style="thin", color="BFDBFE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

REVIEW_COLS = [
    ("Correct?", 11),
    ("Corrected value", 18),
    ("Effective from", 15),
    ("Source cited by reviewer", 34),
    ("Reviewer initials", 16),
    ("Date checked", 14),
    ("Notes", 40),
]


def _title(ws, text: str, subtitle: str, width: int) -> int:
    ws["A1"] = text
    ws["A1"].font = Font(FONT, size=14, bold=True, color="0C4A6E")
    ws["A2"] = subtitle
    ws["A2"].font = Font(FONT, size=10, italic=True, color="475569")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    return 4


def _header(ws, row: int, headers: list[tuple[str, int]]) -> int:
    for i, (label, width) in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=label)
        cell.font = Font(FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return row + 1


def _data_row(ws, row: int, values: list, review_start: int) -> None:
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.font = Font(FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BOX
    for i in range(review_start, review_start + len(REVIEW_COLS)):
        cell = ws.cell(row=row, column=i)
        cell.fill = INPUT_FILL
        cell.border = BOX
        cell.font = Font(FONT, size=10)


def _sheet(wb, name: str, title: str, subtitle: str,
           own_headers: list[tuple[str, int]]):
    ws = wb.create_sheet(name)
    headers = own_headers + REVIEW_COLS
    start = _title(ws, title, subtitle, len(headers))
    row = _header(ws, start, headers)
    return ws, row, len(own_headers) + 1


# --------------------------------------------------------------------------
def build(path: Path) -> dict[str, tuple[int, int, int, int]]:
    wb = Workbook()
    wb.remove(wb.active)
    # tab -> (value count, 'Correct?' column index, first data row, last data row)
    counts: dict[str, tuple[int, int, int, int]] = {}

    # ---------------------------------------------------------------- guide
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 110
    lines = [
        (f"{PRODUCT_NAME} — statutory verification worksheet", 14, True, "0C4A6E"),
        ("", 10, False, None),
        ("Why this exists", 11, True, "0C4A6E"),
        ("GO_LIVE.md D6 makes a payroll professional's sign-off a gate on real client data. "
         "This worksheet is that sign-off in a form an auditor can read a year from now.", 10, False, None),
        ("", 10, False, None),
        ("What to do", 11, True, "0C4A6E"),
        ("1. Fill every YELLOW cell. Nothing else is yellow, so that is the whole instruction.", 10, False, None),
        ("2. 'Correct?' takes Y or N. Leave it blank if you could not verify it — blank is an "
         "honest answer and the Sign-off tab counts it separately.", 10, False, None),
        ("3. On any N, put the right figure in 'Corrected value' and the date it took effect "
         "in 'Effective from'.", 10, False, None),
        ("4. Cite your source. 'Finance Act 2026 s.2' or a gazette number — not 'checked online'.", 10, False, None),
        ("5. Sign the Sign-off tab last. It counts the other tabs automatically.", 10, False, None),
        ("", 10, False, None),
        ("Order of work — highest value first", 11, True, "0C4A6E"),
        ("Income tax FY 2026-27 covers the year currently being billed and is one Finance Act. "
         "Do it first; it is an afternoon. PT and LWF are 22 and 15 separate state schedules and "
         "can follow.", 10, False, None),
        ("", 10, False, None),
        ("What is NOT in this worksheet", 11, True, "0C4A6E"),
        ("Minimum wages. They are per state, per skill category, revised roughly twice a year, and "
         "are held as tenant data rather than shipped defaults. They need their own review. "
         "The product reports a missing rate as a finding rather than as a pass, so an unmaintained "
         "table shows up as unverified instead of clean.", 10, False, None),
        ("", 10, False, None),
        ("The Labour Codes", 11, True, "0C4A6E"),
        ("Everything here is built on the Acts currently in force. If the four Labour Codes have "
         "commenced, the definition of 'wages' changes and PF basis, gratuity and bonus all move "
         "with it. That is a redesign, not a rate correction — flag it on the Sign-off tab and stop.",
         10, False, None),
        ("", 10, False, None),
        (f"Generated {date.today().isoformat()} from the application's own defaults. "
         "Re-runnable: python tools/statutory_worksheet.py", 9, False, "475569"),
    ]
    for i, (text, size, bold, colour) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(FONT, size=size, bold=bold, color=colour or "000000")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if len(text) > 95:
            ws.row_dimensions[i].height = 30
    ws.cell(row=len(lines) + 2, column=1, value="EXAMPLE of a completed row (do not leave this on a data tab):").font = Font(FONT, size=10, bold=True)
    ws.cell(row=len(lines) + 3, column=1,
            value="Correct? = N  |  Corrected value = 20000  |  Effective from = 2026-04-01  |  "
                  "Source = Karnataka PT (Amendment) Act 2026, Gazette 14-Mar-2026  |  "
                  "Reviewer = RKS  |  Date = 2026-09-30").font = Font(FONT, size=10, italic=True)

    # ------------------------------------------------------------ income tax
    years = default_tax_years()
    ws, row, rc = _sheet(
        wb, "Income tax", f"Income tax — FY {DEFAULT_TAX_YEAR}",
        "HIGHEST PRIORITY. This is the financial year currently being billed. The code asserts "
        "Budget 2026 carried FY 2025-26 forward unchanged — that assertion is the first thing to test.",
        [("Parameter", 30), ("Regime", 14), ("Value in code", 18), ("Unit", 12), ("Where in code", 42)],
    )
    first_row = row
    src = "app/services/tax_year_defaults.py"
    n = 0
    cfg = years[DEFAULT_TAX_YEAR]
    for regime_name, regime in (("New Regime", cfg.new_regime), ("Old Regime", cfg.old_regime)):
        for slab in regime.slabs:
            upper = "no limit" if slab.up_to is None else f"{slab.up_to:,.0f}"
            _data_row(ws, row, [f"Slab up to {upper}", regime_name,
                                float(slab.rate), "rate (fraction)", src], rc)
            row += 1
            n += 1
        _data_row(ws, row, ["Standard deduction", regime_name,
                            float(regime.standard_deduction), "INR", src], rc); row += 1; n += 1
        _data_row(ws, row, ["87A rebate — income limit", regime_name,
                            float(regime.rebate.taxable_income_limit), "INR", src], rc); row += 1; n += 1
        _data_row(ws, row, ["87A rebate — maximum", regime_name,
                            float(regime.rebate.max_rebate), "INR", src], rc); row += 1; n += 1
        _data_row(ws, row, ["87A marginal relief applies", regime_name,
                            "Yes" if regime.rebate.marginal_relief else "No", "flag", src], rc); row += 1; n += 1
        for band in regime.surcharge_brackets:
            upper = "above top band" if band.up_to is None else f"{band.up_to:,.0f}"
            _data_row(ws, row, [f"Surcharge up to {upper}", regime_name,
                                float(band.rate), "rate (fraction)", src], rc); row += 1; n += 1
        _data_row(ws, row, ["Chapter VI-A deductions allowed", regime_name,
                            "Yes" if regime.allow_chapter_via else "No", "flag", src], rc); row += 1; n += 1
    _data_row(ws, row, ["Health & education cess", "Both",
                        float(cfg.cess_rate), "rate (fraction)", src], rc); row += 1; n += 1
    counts["Income tax"] = (n, rc, first_row, row - 1)

    # ----------------------------------------------------------- PF and ESIC
    pf, esic = PFConfig(), ESICConfig()
    ws, row, rc = _sheet(
        wb, "PF and ESIC", "Provident Fund and ESIC",
        "Central rates. These change rarely, but the PF wage ceiling has been under review for years "
        "and a change here moves every payslip in the country.",
        [("Parameter", 34), ("Value in code", 18), ("Unit", 16), ("Where in code", 42)],
    )
    first_row = row
    s = "app/schemas/statutory_config.py"
    pf_rows = [
        ("PF — employee contribution", float(pf.rates.employee_rate), "rate (fraction)"),
        ("PF — employer contribution", float(pf.rates.employer_rate), "rate (fraction)"),
        ("PF — EPS share of employer", float(pf.rates.eps_rate), "rate (fraction)"),
        ("PF — EDLI", float(pf.rates.edli_rate), "rate (fraction)"),
        ("PF — admin charges", float(pf.rates.admin_rate), "rate (fraction)"),
        ("PF — wage ceiling", float(pf.wage.wage_ceiling), "INR per month"),
        ("PF — cap wage at ceiling by default", "Yes" if pf.wage.restrict_to_ceiling else "No", "flag"),
        ("ESIC — employee contribution", float(esic.rates.employee_rate), "rate (fraction)"),
        ("ESIC — employer contribution", float(esic.rates.employer_rate), "rate (fraction)"),
        ("ESIC — eligibility ceiling", float(esic.wage.wage_ceiling), "INR per month"),
        ("Gratuity — accrual fraction", "15/26 per year", "fraction of Basic+DA"),
    ]
    for label, value, unit in pf_rows:
        where = "app/services/cost_model.py" if label.startswith("Gratuity") else s
        _data_row(ws, row, [label, value, unit, where], rc); row += 1
    counts["PF and ESIC"] = (len(pf_rows), rc, first_row, row - 1)

    # -------------------------------------------------------------------- PT
    ws, row, rc = _sheet(
        wb, "PT by state", "Professional Tax — shipped state defaults",
        "22 states, each amending on its own timetable. Source stated in code: Simpliance e-Library, "
        "verified late-2025 / early-2026. Check the states your clients actually operate in first.",
        [("State", 20), ("Wage from", 14), ("Wage to", 14), ("Amount", 12),
         ("Frequency", 13), ("Gender", 10), ("Months", 14), ("Where in code", 34)],
    )
    first_row = row
    n = 0
    for state in sorted(pt_defaults.PT_DEFAULTS):
        for slab in pt_defaults.PT_DEFAULTS[state]:
            months = slab.get("applicable_months")
            _data_row(ws, row, [
                state, float(slab["min_salary"]), float(slab["max_salary"]),
                float(slab["deduction_amount"]), slab["frequency"],
                slab.get("gender", "ALL"),
                "all" if not months else ", ".join(str(m) for m in months),
                "app/services/pt_defaults.py",
            ], rc)
            row += 1
            n += 1
    counts["PT by state"] = (n, rc, first_row, row - 1)

    # ------------------------------------------------------------------- LWF
    ws, row, rc = _sheet(
        wb, "LWF by state", "Labour Welfare Fund — shipped state defaults",
        "15 states. LWF amounts change quietly and with little notice, so treat an unchanged figure "
        "as unverified rather than confirmed.",
        [("State", 20), ("Wage from", 14), ("Wage to", 14), ("Employee", 12),
         ("Employer", 12), ("Frequency", 13), ("Months", 14), ("Where in code", 34)],
    )
    first_row = row
    n = 0
    for state in sorted(lwf_defaults.LWF_DEFAULTS):
        for slab in lwf_defaults.LWF_DEFAULTS[state]:
            months = slab.get("applicable_months")
            _data_row(ws, row, [
                state, float(slab["min_salary"]), float(slab["max_salary"]),
                float(slab["deduction_amount"]), float(slab["employer_amount"]),
                slab["frequency"],
                "all" if not months else ", ".join(str(m) for m in months),
                "app/services/lwf_defaults.py",
            ], rc)
            row += 1
            n += 1
    counts["LWF by state"] = (n, rc, first_row, row - 1)

    # -------------------------------------------------------------- sign-off
    ws = wb.create_sheet("Sign-off")
    for col, width in (("A", 26), ("B", 14), ("C", 14), ("D", 14), ("E", 16), ("F", 18)):
        ws.column_dimensions[col].width = width
    r = _title(ws, "Sign-off", "Counts update themselves from the tabs. Sign only when "
                               "'Not checked' is zero on every row you are accepting.", 6)
    for i, label in enumerate(["Tab", "Rows", "Correct (Y)", "Wrong (N)", "Not checked", "Status"], start=1):
        c = ws.cell(row=r, column=i, value=label)
        c.font = Font(FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.border = BOX
    r += 1
    first = r
    for tab, (count, review_col, top, bottom) in counts.items():
        col = get_column_letter(review_col)
        rng = f"'{tab}'!${col}${top}:${col}${bottom}"
        ws.cell(row=r, column=1, value=tab).font = Font(FONT, size=10)
        ws.cell(row=r, column=2, value=count).font = Font(FONT, size=10)
        ws.cell(row=r, column=3, value=f'=COUNTIF({rng},"Y")').font = Font(FONT, size=10)
        ws.cell(row=r, column=4, value=f'=COUNTIF({rng},"N")').font = Font(FONT, size=10)
        ws.cell(row=r, column=5, value=f"=B{r}-C{r}-D{r}").font = Font(FONT, size=10)
        ws.cell(row=r, column=6,
                value=f'=IF(E{r}>0,"INCOMPLETE",IF(D{r}>0,"CORRECTIONS NEEDED","ACCEPTED"))')
        ws.cell(row=r, column=6).font = Font(FONT, size=10, bold=True)
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = BOX
            if col == 5:
                ws.cell(row=r, column=col).fill = BAND_FILL
        r += 1
    last = r - 1

    r += 1
    ws.cell(row=r, column=1, value="Total values").font = Font(FONT, size=10, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{last})").font = Font(FONT, size=10, bold=True)
    ws.cell(row=r, column=5, value=f"=SUM(E{first}:E{last})").font = Font(FONT, size=10, bold=True)
    ws.cell(row=r, column=6,
            value=f'=IF(E{r}>0,"NOT READY FOR CLIENT DATA","D6 SATISFIED")')
    ws.cell(row=r, column=6).font = Font(FONT, size=11, bold=True)
    ws.cell(row=r, column=6).fill = RISK_FILL

    r += 3
    ws.cell(row=r, column=1, value="Reviewer declaration").font = Font(FONT, size=11, bold=True, color="0C4A6E")
    r += 1
    for label in ["Name", "Qualification / membership no.", "Firm", "Date",
                  "Signature", "Labour Codes commenced? (Y/N)"]:
        ws.cell(row=r, column=1, value=label).font = Font(FONT, size=10)
        cell = ws.cell(row=r, column=2)
        cell.fill = INPUT_FILL
        cell.border = BOX
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        r += 1
    r += 1
    ws.cell(row=r, column=1,
            value="I have checked the values marked Y against the sources I cited, and they are "
                  "correct as at the date above.").font = Font(FONT, size=10, italic=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)

    wb.save(path)
    return counts


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "statutory_verification_worksheet.xlsx")
    totals = build(out)
    print(f"wrote {out}")
    for tab, (count, _col, top, bottom) in totals.items():
        print(f"  {tab}: {count} values to verify (rows {top}-{bottom})")
    print(f"  TOTAL: {sum(c for c, *_ in totals.values())}")
