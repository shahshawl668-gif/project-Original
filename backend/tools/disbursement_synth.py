"""
Synthetic disbursement test data: fictional employees, planted errors, an answer key.

    python tools/disbursement_synth.py --out /tmp/dsb --employees 500 --seed 7 --scenario all --layout both

Everything here is invented. Names come from short generic lists; IFSC codes use
the made-up bank prefixes ZZZA0 … ZZZH0, so none can match a real branch;
account numbers are random digits. Never point this at, or mix it with, real data.

Scenarios
---------
A  a clean month                       → CLEAR_TO_RELEASE, zero findings
B  row-level errors only               → RELEASE_WITH_HOLDS (every hold/flag check planted at least twice)
C  file-level breaks                   → DO_NOT_RELEASE
D  optional inputs missing             → those checks NOT_RUN, the rest still run

The bank pays **total salary**, not net pay: net pay plus reimbursements and any
held salary released, less salary held this month. About a fifth of employees
claim reimbursements; a few have part or all of their salary held, or an earlier
hold released.

The clean population is deliberately untidy the way real exports are — trailing
spaces, lower-case IFSC, amounts written as text with commas, salutations and
initials in beneficiary names, account numbers whose leading zeros a spreadsheet
dropped from the master — and must still produce no findings.

Each pack writes its files, ``expected_findings.csv`` (one row per expected
rule and employee; NOT_RUN rules with a blank employee), and ``scenario.json``
(profile, template, settings and expected verdict).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import random
import sys
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

# Run as a script from anywhere: make the backend package importable.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FIRST = ["Aarav", "Vivaan", "Aditya", "Vihaan", "Arjun", "Sai", "Reyansh", "Krishna", "Ishaan", "Rohan",
         "Ananya", "Diya", "Aadhya", "Saanvi", "Kavya", "Priya", "Meera", "Lakshmi", "Pooja", "Sneha",
         "Rahul", "Amit", "Suresh", "Ramesh", "Vikram", "Nikhil", "Deepak", "Manoj", "Sunita", "Anjali",
         "Farhan", "Imran", "Ayesha", "Zoya", "Harpreet", "Gurpreet", "Joseph", "Mary", "Thomas", "Anitha"]
MIDDLE = ["Kumar", "Kumari", "Prasad", "Lal", "Devi", "Rani", "Chandra", "Mohan"]
LAST = ["Sharma", "Verma", "Gupta", "Iyer", "Nair", "Reddy", "Rao", "Patel", "Shah", "Mehta",
        "Singh", "Kaur", "Das", "Bose", "Banerjee", "Mukherjee", "Pillai", "Menon", "Joshi", "Kulkarni",
        "Desai", "Naidu", "Yadav", "Mishra", "Pandey", "Khan", "Ansari", "Fernandes", "D'Souza", "Thakur"]
STATES = ["Karnataka", "Maharashtra", "Tamil Nadu", "Telangana", "Delhi", "Gujarat", "West Bengal",
          "Uttar Pradesh", "Kerala", "Haryana", "Rajasthan", "Punjab"]
BANKS = "ABCDEFGH"
ALNUM = "ABCDEFGHJKLMNPQRSTUVWXYZ0123456789"

PERIOD = "2026-09"
P_START, P_END = date(2026, 9, 1), date(2026, 9, 30)
PREV_START = date(2026, 8, 1)
VALUE_DATE = date(2026, 9, 30)
DEBIT_ACCOUNT = "ZZ0000123456"     # synthetic salary account

LAYOUTS = {
    "generic": {"profile": "generic", "template": "generic_csv", "ext": "csv"},
    "darwinbox_style": {"profile": "darwinbox_style", "template": "batch_pipe_hdt", "ext": "xlsx"},
}

HEADERS = {
    "generic": {
        "register": ["Employee ID", "Employee Name", "Work State", "Net Pay", "Status", "Date of Joining",
                     "Date of Exit", "FnF Processed", "On Hold", "Arrears", "Increment", "Reimbursement",
                     "Salary Hold", "Hold Release", "Total Salary"],
        "bank_master": ["Employee ID", "Account Number", "IFSC", "Beneficiary Name", "Verification Status",
                        "Last Changed"],
        "change_log": ["Employee ID", "Field", "Old Value", "New Value", "Changed On", "Verification Status"],
        "previous": ["Employee ID", "Net Pay", "Total Salary", "Account Number", "IFSC"],
        "hold_list": ["Employee ID", "Category", "Reason", "Effective Date"],
        "offcycle": ["Employee ID", "Amount", "Payment Date", "Reference"],
    },
    "darwinbox_style": {
        "register": ["Employee Code", "Employee Name", "Work Location State", "Net Salary", "Employment Status",
                     "Date Of Joining", "Date Of Exit", "FnF Status", "Payroll Hold", "Arrear Amount",
                     "Increment Applied", "Reimbursements", "Salary On Hold", "Hold Released",
                     "Total Salary Payable"],
        "bank_master": ["Employee Code", "Bank Account Number", "IFSC Code", "Name As Per Bank",
                        "Bank Details Status", "Bank Details Updated On"],
        "change_log": ["Employee Code", "Changed Field", "Previous Value", "New Value", "Changed On",
                       "Approval Status"],
        "previous": ["Employee Code", "Net Salary", "Total Salary Payable", "Bank Account Number", "IFSC Code"],
        "hold_list": ["Employee Code", "Hold Type", "Hold Reason", "Hold From"],
        "offcycle": ["Employee Code", "Off Cycle Amount", "Paid On", "Payment Reference"],
    },
}


@dataclass
class Emp:
    eid: str
    name: str
    state: str
    doj: date
    net: Decimal
    account: str
    ifsc: str
    status: str = "active"           # active | separated
    exit: date | None = None
    ff: bool = False
    hold: bool = False
    hold_entry: tuple[str, str] | None = None
    prev_net: Decimal | None = None
    prev_account: str | None = None
    prev_ifsc: str | None = None
    verified: bool = True
    last_changed: date = date(2022, 1, 1)
    arrears: Decimal = Decimal("0")
    increment: bool = False
    reimbursement: Decimal = Decimal("0")
    salary_hold: Decimal = Decimal("0")
    hold_release: Decimal = Decimal("0")
    total_override: Decimal | None = None     # a register whose total salary does not add up
    prev_reimbursement: Decimal = Decimal("0")
    offcycle: list[tuple[Decimal, date, str]] = field(default_factory=list)
    changes: list[tuple[str, str, str, date, bool | None]] = field(default_factory=list)
    in_file: bool = True
    copies: int = 1
    file_amount: Decimal | None = None
    file_name: str | None = None
    file_account: str | None = None
    file_ifsc: str | None = None
    file_id: str | None = None
    master_account: str | None = None   # what the master export shows, when it differs in form
    in_register: bool = True
    kind: str = "plain"
    messy: set[str] = field(default_factory=set)

    @property
    def total(self) -> Decimal:
        """What the bank pays: net + reimbursements + released hold − salary held."""
        if self.total_override is not None:
            return self.total_override
        return self.net + self.reimbursement + self.hold_release - self.salary_hold

    @property
    def prev_total(self) -> Decimal | None:
        return None if self.prev_net is None else self.prev_net + self.prev_reimbursement


@dataclass
class Pack:
    scenario: str
    layout: str
    files: dict[str, tuple[str, bytes]]
    expected: list[dict[str, str]]
    settings: dict[str, Any]
    value_date: date | None
    verdict: str
    profile: str
    template: str


class Builder:
    def __init__(self, n: int, seed: int, scenario: str, layout: str):
        self.rng = random.Random(seed)
        self.n, self.scenario, self.layout = n, scenario, layout
        self.emps: list[Emp] = []
        self.expected: list[dict[str, str]] = []
        self.extra_rows: list[dict[str, Any]] = []
        self.header_overrides: dict[str, Any] = {}
        self.footer_overrides: dict[str, Any] = {}

    # -- population ---------------------------------------------------------------
    def _ifsc(self) -> str:
        return "ZZZ" + self.rng.choice(BANKS) + "0" + "".join(self.rng.choice(ALNUM) for _ in range(6))

    def _account(self, zeros: bool = False) -> str:
        length = self.rng.randint(11, 16) if zeros else self.rng.randint(9, 18)
        body = "".join(self.rng.choice("0123456789") for _ in range(length))
        if zeros:
            return "00" + "1" + body[3:]
        return str(self.rng.randint(1, 9)) + body[1:]

    def _name(self) -> str:
        parts = [self.rng.choice(FIRST)]
        if self.rng.random() < 0.2:
            parts.append(self.rng.choice(MIDDLE))
        parts.append(self.rng.choice(LAST))
        return " ".join(parts)

    def _money(self, lo: int, hi: int) -> Decimal:
        return Decimal(self.rng.randint(lo * 100, hi * 100)) / 100

    def population(self) -> None:
        rng = self.rng
        for i in range(1, self.n + 1):
            net = self._money(18000, 250000)
            e = Emp(eid=f"E{i:05d}", name=self._name(), state=rng.choice(STATES),
                    doj=date(2012, 1, 1) + timedelta(days=rng.randint(0, 5200)), net=net,
                    account=self._account(zeros=rng.random() < 0.1), ifsc=self._ifsc())
            e.prev_net = (net * Decimal(str(round(rng.uniform(0.95, 1.05), 4)))).quantize(Decimal("0.01"))
            e.prev_account, e.prev_ifsc = e.account, e.ifsc
            self.emps.append(e)
        pool = list(range(len(self.emps)))
        rng.shuffle(pool)
        share = max(2, self.n // 50)          # about 2% each, never fewer than 2
        take = lambda k: [self.emps[pool.pop()] for _ in range(min(k, len(pool)))]   # noqa: E731
        for e in take(share):                 # joined this month: not paid last month
            e.kind, e.doj, e.prev_net, e.prev_account, e.prev_ifsc = "joiner", P_START + timedelta(days=rng.randint(0, 20)), None, None, None
            e.last_changed = e.doj
        for e in take(share):                 # joined last month: part month then
            e.kind, e.doj = "joiner_prev", PREV_START + timedelta(days=rng.randint(5, 25))
            e.prev_net = (e.net * Decimal("0.4")).quantize(Decimal("0.01"))
            e.last_changed = e.doj
        for e in take(share):                 # leaving this month: last salary paid
            e.kind, e.status, e.exit = "leaver", "separated", P_START + timedelta(days=rng.randint(5, 28))
            e.net = (e.net * Decimal("0.6")).quantize(Decimal("0.01"))
        for e in take(max(2, share // 2)):    # F&F processed: paid separately, not in the file
            e.kind, e.status, e.exit, e.ff, e.in_file = "ff", "separated", PREV_START + timedelta(days=10), True, False
            e.hold_entry = ("F&F processed", "Full and final settled separately")
        for e in take(max(2, share // 2)):    # on hold: not in the file
            e.kind, e.hold, e.in_file = "hold", True, False
            e.hold_entry = ("On hold", "Pending documents")
        for e in take(max(2, share // 2)):    # paid off-cycle already: not in the file
            e.kind, e.in_file = "offcycle", False
            e.offcycle.append((e.total, P_START + timedelta(days=12), f"OFC{e.eid[1:]}"))
        for e in take(share):                 # arrears this month: big change, explained
            e.kind, e.arrears = "arrears", self._money(5000, 40000)
            e.net = (e.prev_net * Decimal("1.4")).quantize(Decimal("0.01"))
        for e in take(max(2, share // 2)):    # increment this month
            e.kind, e.increment = "increment", True
            e.net = (e.prev_net * Decimal("1.3")).quantize(Decimal("0.01"))
        for e in take(max(2, share // 2)):    # changed bank details last month, verified
            e.kind = "changed_verified"
            old = e.account
            e.account = self._account()
            e.prev_account = old
            e.last_changed = PREV_START + timedelta(days=20)
            e.changes.append(("Account Number", old, e.account, e.last_changed, True))
        for e in take(max(2, share // 2)):    # changed years ago: outside the window
            e.changes.append(("IFSC", self._ifsc(), e.ifsc, date(2023, 3, 14), True))
        for e in take(max(2, share // 2)):    # part of this month's salary held
            e.kind, e.salary_hold = "partial_hold", (e.net * Decimal("0.3")).quantize(Decimal("0.01"))
        for e in take(max(2, share // 2)):    # all of it held: nothing payable, not in the file
            e.kind, e.salary_hold, e.in_file = "held_full", e.net, False
        for e in take(max(2, share // 2)):    # an earlier hold released this month
            e.kind, e.hold_release = "released", self._money(5000, 40000)
        self.plain = [self.emps[i] for i in pool]
        for e in self.emps:                   # about a fifth claim reimbursements
            if e.kind in ("plain", "partial_hold", "released", "leaver") and self.rng.random() < 0.2:
                e.reimbursement = self._money(500, 15000)
            if e.prev_net is not None and self.rng.random() < 0.2:
                e.prev_reimbursement = self._money(500, 15000)
        # Untidy but correct: must produce no findings.
        for e in self.plain[: max(4, self.n // 25)]:
            e.messy.add(self.rng.choice(["space_id", "lower_ifsc", "salutation", "initials", "amount_text"]))
        for e in self.emps:
            if e.account.startswith("00") and self.rng.random() < 0.5:
                e.master_account = e.account.lstrip("0")     # the master export lost the zeros
        self.plain = self.plain[max(4, self.n // 25):]

    # -- planting ------------------------------------------------------------------
    def pick(self, k: int = 2) -> list[Emp]:
        out = [self.plain.pop() for _ in range(k)]
        for e in out:
            e.kind = "planted"
        return out

    def expect(self, rule: str, sev: str, eid: str = "", note: str = "") -> None:
        self.expected.append({"scenario": self.scenario, "layout": self.layout, "rule_id": rule,
                              "severity": sev, "employee_id": eid, "note": note})

    def plant_rows(self) -> None:
        a, b, c, d = self.pick(4)
        for x, y in ((a, b), (c, d)):         # DSB-05 two employees, one account
            y.account, y.ifsc, y.prev_account, y.prev_ifsc = x.account, x.ifsc, x.account, x.ifsc
            y.master_account = x.master_account
            for e in (x, y):
                self.expect("DSB-05", "HOLD_ROW", e.eid, "shared account")
        for e in self.pick(2):                # DSB-06 changed this month, not verified
            old = e.account
            e.account = self._account()
            e.prev_account, e.verified, e.last_changed = old, False, P_START + timedelta(days=3)
            e.changes.append(("Account Number", old, e.account, e.last_changed, None))
            self.expect("DSB-06", "HOLD_ROW", e.eid, "account changed, pending verification")
        z, neg = self.pick(2)                 # DSB-07 zero and negative
        z.net = z.prev_net = Decimal("0.00")
        neg.net = neg.prev_net = Decimal("-1250.00")
        for e in (z, neg):
            e.reimbursement = e.prev_reimbursement = Decimal("0")
        for e in (z, neg):
            self.expect("DSB-07", "HOLD_ROW", e.eid, "zero or negative amount")
        up, down = self.pick(2)               # DSB-08 file amount differs
        up.file_amount, down.file_amount = up.total + Decimal("1500.00"), down.total - Decimal("0.50")
        for e in (up, down):
            self.expect("DSB-08", "HOLD_ROW", e.eid, "amount differs from total salary")
        for e in self.pick(2):                # DSB-17 the register's total does not add up; the bank pays it
            e.reimbursement = self._money(1000, 8000)
            e.total_override = e.net + e.reimbursement + Decimal("2000.00")
            self.expect("DSB-17", "FLAG", e.eid, "total salary is not net + reimbursements − held")
        h, s, f = self.pick(3)                # DSB-09 should not be in the file
        h.hold, h.hold_entry = True, ("On hold", "Disciplinary review")
        s.status, s.exit = "separated", PREV_START + timedelta(days=15)
        f.status, f.exit, f.ff = "separated", PREV_START + timedelta(days=8), True
        for e in (h, s, f):
            self.expect("DSB-09", "HOLD_ROW", e.eid, "on hold / separated before period / F&F processed")
        for e in self.pick(2):                # DSB-10 twice in the file
            e.copies = 2
            self.expect("DSB-10", "HOLD_ROW", e.eid, "duplicate line")
        for e in self.pick(2):                # DSB-11 already paid off-cycle
            e.offcycle.append((e.total, P_START + timedelta(days=18), f"OFC{e.eid[1:]}"))
            self.expect("DSB-11", "HOLD_ROW", e.eid, "paid off-cycle this period")
        bad_ifsc, letters, short = self.pick(3)  # DSB-12 (the master carries the same bad value)
        bad_ifsc.ifsc = bad_ifsc.prev_ifsc = "ZZZB1A2B3C4"
        letters.account = letters.prev_account = "12345AB7890"
        short.account = short.prev_account = "123456"
        for e in (bad_ifsc, letters, short):
            e.master_account = None
            self.expect("DSB-12", "HOLD_ROW", e.eid, "invalid IFSC or account")
        rise, fall = self.pick(2)             # DSB-13 big change, unexplained
        rise.prev_net = (rise.net / Decimal("1.6")).quantize(Decimal("0.01"))
        fall.prev_net = (fall.net / Decimal("0.55")).quantize(Decimal("0.01"))
        for e in (rise, fall):
            self.expect("DSB-13", "FLAG", e.eid, "net pay change beyond 25%")
        for e in self.pick(2):                # DSB-14 someone else's name
            others = [n for n in (self._name() for _ in range(5)) if not set(n.lower().split()) & set(e.name.lower().split())]
            e.file_name = others[0] if others else "Zzz Unrelated Person"
            self.expect("DSB-14", "FLAG", e.eid, "beneficiary name differs")
        for e in self.pick(2):                # DSB-15 due but missing
            e.in_file = False
            self.expect("DSB-15", "FLAG", e.eid, "missing from the bank file")
        other, zeros, ifsc = self.pick(3)     # DSB-16 file is not the master
        other.file_account = self._account()
        zeros.account = zeros.prev_account = "00" + "".join(self.rng.choice("0123456789") for _ in range(12))
        zeros.master_account = None
        zeros.file_account = zeros.account.lstrip("0")
        ifsc.file_ifsc = self._ifsc()
        for e in (other, zeros, ifsc):
            self.expect("DSB-16", "HOLD_ROW", e.eid, "file account/IFSC differs from master")

    def plant_file_breaks(self) -> None:
        for i in (1, 2):                      # DSB-03 paid, not in the register
            eid = f"X9000{i}"
            self.extra_rows.append({"eid": eid, "name": self._name(), "account": self._account(),
                                    "ifsc": self._ifsc(), "amount": self._money(20000, 90000)})
            self.expect("DSB-03", "STOP_FILE", eid, "not in the register")
            # Someone the register has never heard of is not on the bank master either.
            self.expect("DSB-16", "HOLD_ROW", eid, "no bank master record (follows from DSB-03)")
        if self.layout == "darwinbox_style":
            self.header_overrides = {"total_delta": Decimal("1000.00"), "debit": "ZZ0000999999",
                                     "value_date": "29092026"}
            self.footer_overrides = {"count_delta": 1}
            self.expect("DSB-01", "STOP_FILE", "", "header total does not match the lines")
            self.expect("DSB-02", "STOP_FILE", "", "footer count does not match the lines")
            self.expect("DSB-04", "STOP_FILE", "", "debit account is not the configured one")
            self.expect("DSB-04", "STOP_FILE", "", "value date is not the one given for the run")
        else:
            wrong = self.plain[: max(2, int(len(self.emps) * 0.3))]
            self.plain = self.plain[len(wrong):]
            for e in wrong:                   # a file from the wrong run: amounts all off
                e.kind = "planted"
                e.file_amount = (e.total * Decimal("1.07")).quantize(Decimal("0.01"))
                self.expect("DSB-08", "HOLD_ROW", e.eid, "amount differs (wrong run)")
            self.expect("DSB-01", "STOP_FILE", "", "held amount share above the cap")
            self.expect("DSB-02", "STOP_FILE", "", "held row share above the cap")

    def plant_missing_inputs(self) -> None:
        up, down = self.pick(2)
        up.file_amount, down.file_amount = up.total + Decimal("250.00"), down.total - Decimal("99.00")
        for e in (up, down):
            self.expect("DSB-08", "HOLD_ROW", e.eid, "amount differs from register")
        bad, short = self.pick(2)
        bad.ifsc = bad.prev_ifsc = "zzzc0abc12"       # ten characters
        short.account = short.prev_account = "98765"
        for e in (bad, short):
            e.master_account = None
            self.expect("DSB-12", "HOLD_ROW", e.eid, "invalid IFSC or account")
        for e in self.emps:                   # nobody said they were paid off-cycle
            if e.kind == "offcycle":
                self.expect("DSB-15", "FLAG", e.eid, "paid off-cycle, but no off-cycle file was given")
        for rule in ("DSB-06", "DSB-11", "DSB-13"):
            self.expect(rule, "NOT_RUN", "", "optional input not supplied")
        if self.layout == "darwinbox_style":
            self.expect("DSB-04", "NOT_RUN", "", "no debit account or value date configured")

    # -- writing ---------------------------------------------------------------------
    def _status_text(self, e: Emp) -> str:
        if self.layout == "generic":
            if e.hold:
                return "On Hold"
            return "Separated" if e.status == "separated" else "Active"
        if e.status == "separated":
            return "Resigned"
        return "Notice Period" if e.kind == "leaver" else "Active"

    def _amount_out(self, value: Decimal, text: bool) -> Any:
        if self.layout == "generic":
            return f"{value:,.2f}" if text else f"{value:.2f}"
        return f"{value:,.2f}" if text else float(value)

    def _date_out(self, d: date | None) -> Any:
        if d is None:
            return None if self.layout != "generic" else ""
        return d.isoformat() if self.layout == "generic" else d

    def tables(self, missing: set[str]) -> dict[str, list[list[Any]]]:
        g = self.layout == "generic"
        yes, no = ("Yes", "No") if g else ("Yes", None)
        out: dict[str, list[list[Any]]] = {k: [] for k in HEADERS[self.layout]}
        for e in self.emps:
            if not e.in_register:
                continue
            net_text = "amount_text" in e.messy
            out["register"].append([
                e.eid, e.name, e.state, self._amount_out(e.net, net_text), self._status_text(e),
                self._date_out(e.doj), self._date_out(e.exit),
                (yes if e.ff else no) if g else ("FnF Processed" if e.ff else ("Pending" if e.status == "separated" else None)),
                (yes if e.hold else no) if g else ("Hold" if e.hold else None),
                self._amount_out(e.arrears, False) if e.arrears else ("" if g else None),
                (yes if e.increment else no) if g else ("Yes" if e.increment else None),
                self._amount_out(e.reimbursement, False) if e.reimbursement else ("" if g else None),
                self._amount_out(e.salary_hold, False) if e.salary_hold else ("" if g else None),
                self._amount_out(e.hold_release, False) if e.hold_release else ("" if g else None),
                self._amount_out(e.total, "amount_text" in e.messy),
            ])
            account = e.master_account or e.account
            if not g and e.master_account and e.master_account.isdigit():
                account = int(e.master_account)               # a spreadsheet number: zeros gone
            verified = ("Verified" if e.verified else "Pending") if g else ("Approved" if e.verified else "Pending Approval")
            out["bank_master"].append([e.eid, account, e.ifsc, e.name, verified, self._date_out(e.last_changed)])
            for what, old, new, when, ok in e.changes:
                status = None if ok is None else (("Verified" if ok else "Rejected") if g else ("Approved" if ok else "Rejected"))
                if ok is None:
                    status = "Pending" if g else "Pending Approval"
                out["change_log"].append([e.eid, what, old, new, self._date_out(when), status])
            if e.prev_net is not None:
                out["previous"].append([e.eid, self._amount_out(e.prev_net, False),
                                        self._amount_out(e.prev_total, False), e.prev_account, e.prev_ifsc])
            if e.hold_entry:
                cat, why = e.hold_entry
                out["hold_list"].append([e.eid, cat, why, self._date_out(P_START)])
            for amount, when, ref in e.offcycle:
                out["offcycle"].append([e.eid, self._amount_out(amount, False), self._date_out(when), ref])
        for slot in missing:
            out.pop(slot, None)
        return out

    def bank_lines(self) -> list[dict[str, Any]]:
        lines = []
        for e in self.emps:
            if not e.in_file:
                continue
            name = e.file_name or e.name
            if "salutation" in e.messy:
                name = "Mr. " + name if e.name.split()[0] not in ("Ananya", "Diya", "Aadhya", "Saanvi", "Kavya", "Priya",
                                                                "Meera", "Lakshmi", "Pooja", "Sneha", "Sunita",
                                                                "Anjali", "Ayesha", "Zoya", "Mary", "Anitha") else "Ms. " + name
            if "initials" in e.messy:
                parts = e.name.split()
                name = " ".join(p[0] + "." for p in parts[:-1]) + " " + parts[-1].upper()
            ifsc = e.file_ifsc or e.ifsc
            if "lower_ifsc" in e.messy:
                ifsc = ifsc.lower()
            eid = e.eid + ("  " if "space_id" in e.messy else "")
            amount = e.file_amount if e.file_amount is not None else e.total
            for _ in range(e.copies):
                lines.append({"eid": eid, "name": name, "account": e.file_account or e.account, "ifsc": ifsc,
                              "amount": amount})
        lines.extend(self.extra_rows)
        return lines

    def bank_file(self) -> tuple[str, bytes]:
        lines = self.bank_lines()
        total = sum((ln["amount"] for ln in lines), Decimal("0"))
        narration = "SAL SEP 2026"
        if self.layout == "generic":
            buf = io.StringIO()
            w = csv.writer(buf, lineterminator="\r\n")
            w.writerow(["Employee ID", "Beneficiary Name", "Account Number", "IFSC", "Amount", "Narration"])
            for ln in lines:
                w.writerow([ln["eid"], ln["name"], ln["account"], ln["ifsc"], f"{ln['amount']:.2f}", narration])
            return "bank_file.csv", buf.getvalue().encode()
        h = self.header_overrides
        f = self.footer_overrides
        out = [f"H|{h.get('debit', DEBIT_ACCOUNT)}|{h.get('value_date', VALUE_DATE.strftime('%d%m%Y'))}|"
               f"{len(lines)}|{total + h.get('total_delta', Decimal('0')):.2f}"]
        for ln in lines:
            out.append(f"D|{ln['eid']}|{ln['name']}|{ln['account']}|{ln['ifsc']}|{ln['amount']:.2f}|{narration}")
        out.append(f"T|{len(lines) + f.get('count_delta', 0)}|{total:.2f}")
        return "bank_file.txt", ("\n".join(out) + "\n").encode()


def _csv_bytes(header: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(["" if v is None else v for v in r] for r in rows)
    return buf.getvalue().encode()


def _xlsx_bytes(header: list[str], rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            if isinstance(cell.value, date):
                cell.number_format = "DD-MMM-YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate(n: int = 500, seed: int = 20260901, scenario: str = "B", layout: str = "generic") -> Pack:
    if layout not in LAYOUTS:
        raise ValueError(f"layout is one of {', '.join(LAYOUTS)}")
    if not 10 <= n <= 10000:
        raise ValueError("employees must be between 10 and 10,000")
    b = Builder(n, seed, scenario, layout)
    b.population()
    missing: set[str] = set()
    settings: dict[str, Any] = {"debit_account": DEBIT_ACCOUNT}
    value_date: date | None = VALUE_DATE
    verdict = "CLEAR_TO_RELEASE"
    if scenario == "B":
        b.plant_rows()
        verdict = "RELEASE_WITH_HOLDS"
    elif scenario == "C":
        b.plant_file_breaks()
        verdict = "DO_NOT_RELEASE"
    elif scenario == "D":
        b.plant_missing_inputs()
        missing = {"change_log", "previous", "hold_list", "offcycle"}
        settings, value_date = {}, None
        verdict = "RELEASE_WITH_HOLDS"
    elif scenario != "A":
        raise ValueError("scenario is A, B, C or D")
    ext = LAYOUTS[layout]["ext"]
    files: dict[str, tuple[str, bytes]] = {"bank_file": b.bank_file()}
    for slot, rows in b.tables(missing).items():
        header = HEADERS[layout][slot]
        data = _csv_bytes(header, rows) if ext == "csv" else _xlsx_bytes(header, rows)
        files[slot] = (f"{slot}.{ext}", data)
    return Pack(scenario, layout, files, b.expected, settings, value_date, verdict,
                LAYOUTS[layout]["profile"], LAYOUTS[layout]["template"])


def write(pack: Pack, out: Path) -> Path:
    where = out / f"scenario_{pack.scenario}_{pack.layout}"
    where.mkdir(parents=True, exist_ok=True)
    for _slot, (name, data) in pack.files.items():
        (where / name).write_bytes(data)
    with (where / "expected_findings.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["scenario", "layout", "rule_id", "severity", "employee_id", "note"])
        w.writeheader()
        w.writerows(pack.expected)
    (where / "scenario.json").write_text(json.dumps({
        "scenario": pack.scenario, "layout": pack.layout, "profile": pack.profile, "template": pack.template,
        "period": PERIOD, "value_date": pack.value_date.isoformat() if pack.value_date else None,
        "settings": pack.settings, "expected_verdict": pack.verdict,
        "files": {slot: name for slot, (name, _) in pack.files.items()},
    }, indent=2))
    return where


def run_pack(pack: Pack):
    """Load a pack's files through the product's own readers and run the checks."""
    from app.services.disbursement import config, engine, inputs, template

    profile = inputs.builtin_profiles()[pack.profile]
    tmpl = template.builtin_templates()[pack.template]
    loaded, parsed = inputs.load(pack.files, profile, tmpl)
    result = engine.run(loaded, config.build_settings(PERIOD, pack.settings, pack.value_date))
    return result, parsed, tmpl


def score(packs_and_results) -> dict[str, dict[str, int]]:
    """Per rule: cases planted, caught, and findings on rows nobody planted (false positives)."""
    table: dict[str, dict[str, int]] = {}
    for pack, result in packs_and_results:
        expected = {(e["rule_id"], e["employee_id"]) for e in pack.expected}
        got = {(f.rule_id, f.employee_id or "") for f in result.findings}
        planted: dict[str, int] = {}
        for e in pack.expected:
            planted[e["rule_id"]] = planted.get(e["rule_id"], 0) + 1
        for rule, n in planted.items():
            row = table.setdefault(rule, {"planted": 0, "caught": 0, "false_positives": 0})
            row["planted"] += n
            missed = {k for k in expected if k[0] == rule} - got
            row["caught"] += n - sum(1 for e in pack.expected if (e["rule_id"], e["employee_id"]) in missed)
        for rule, _emp in got - expected:
            table.setdefault(rule, {"planted": 0, "caught": 0, "false_positives": 0})["false_positives"] += 1
    return dict(sorted(table.items()))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--employees", type=int, default=500)
    ap.add_argument("--seed", type=int, default=20260901)
    ap.add_argument("--scenario", default="all", choices=["A", "B", "C", "D", "all"])
    ap.add_argument("--layout", default="both", choices=["generic", "darwinbox_style", "both"])
    ap.add_argument("--check", action="store_true",
                    help="also run the checks on each pack and print planted / caught / false positives per rule")
    args = ap.parse_args(argv)
    scenarios = ["A", "B", "C", "D"] if args.scenario == "all" else [args.scenario]
    layouts = list(LAYOUTS) if args.layout == "both" else [args.layout]
    ran = []
    for sc in scenarios:
        for lay in layouts:
            pack = generate(args.employees, args.seed, sc, lay)
            where = write(pack, args.out)
            print(f"wrote {where}")
            if args.check:
                result, _, _ = run_pack(pack)
                ran.append((pack, result))
                print(f"   verdict {result.verdict} (expected {pack.verdict})")
    if args.check:
        print(f"\n{'Rule':8} {'Planted':>8} {'Caught':>8} {'False +':>8}")
        for rule, row in score(ran).items():
            print(f"{rule:8} {row['planted']:>8} {row['caught']:>8} {row['false_positives']:>8}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
